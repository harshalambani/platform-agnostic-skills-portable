"""The run result text and the Excel workbook for the card-spend booking skill.

Block order in the text (decision: payments without a statement first):
  1. RED FLAG: bank payments to Credit Card Payment that match no statement
  2. fees booked
  3. large spends (possible assets)
  4. statements refused / not settled / partly settled
  5. EMI rows: not booked, awaiting decision
  6. issues, then the headline, the CCP tie-out and the counts
The headline never says "completed successfully" while anything above is open.
"""
from __future__ import annotations

from dataclasses import dataclass, field

from . import journal as J
from .matcher import BookPayment


def money(v) -> str:
    return f"{v:,.2f}"


@dataclass
class RunReport:
    label: str = ""
    threshold: float = 0.0
    ccp_path: str = ""
    entity: str = ""
    journals: list = field(default_factory=list)
    settlements: list = field(default_factory=list)
    pairings: list = field(default_factory=list)
    unmatched_payments: list = field(default_factory=list)      # [(BookPayment, reason)]
    ambiguous: list = field(default_factory=list)
    lines_unmatched: list = field(default_factory=list)
    refused: list = field(default_factory=list)                 # [(Statement, TieOut)]
    awaiting: list = field(default_factory=list)                # Statements not yet paid (no later statement)
    unpaid: list = field(default_factory=list)                  # Statements with a later statement but no payment
    partly_booked: list = field(default_factory=list)           # [(Statement, reason)]
    emi: dict = field(default_factory=dict)
    issues: list = field(default_factory=list)
    skipped: list = field(default_factory=list)
    duplicates: list = field(default_factory=list)
    statements_read: int = 0
    tie: dict = field(default_factory=dict)
    output_csv: str = ""
    output_xlsx: str = ""
    errors: list = field(default_factory=list)

    @property
    def ready(self) -> list:
        return [j for j in self.journals if j.status == J.ST_READY]

    @property
    def withheld(self) -> list:
        return [j for j in self.journals if j.status not in (J.ST_READY, J.ST_BOOKED)]

    @property
    def unmatched_total(self) -> float:
        return round(sum(p.amount for p, _r in self.unmatched_payments), 2)

    @property
    def open_items(self) -> bool:
        """True while anything stops the run from being called a success."""
        return bool(self.issues or self.unmatched_payments or self.lines_unmatched or self.ambiguous
                    or self.refused or self.unpaid or self.partly_booked or self.withheld
                    or (self.emi and self.emi.get("rows"))
                    or any(s.status != "SETTLED" for s in self.settlements)
                    or abs(self.tie.get("other", 0.0)) > self.tie.get("tol", 0.5))


def red_flag_block(rep: RunReport) -> str:
    if not rep.unmatched_payments:
        return ""
    lines = [f"RED FLAG: {len(rep.unmatched_payments)} bank payment(s) to {rep.ccp_path} match no statement "
             f"(total {money(rep.unmatched_total)}). Money left the bank for a card and no spend is booked "
             f"against it. Nothing here was matched to the nearest statement."]
    for p, reason in sorted(rep.unmatched_payments, key=lambda x: (x[0].date, x[0].amount)):
        lines.append(f"  {p.date:%d %b %Y} | {p.bank_account or 'bank account not found'} | {money(p.amount)} | "
                     f"{'CRED' if p.cred else 'direct'} | {reason}")
    return "\n".join(lines)


def fee_block(rep: RunReport) -> str:
    fees = [j for j in rep.journals if j.kind in (J.K_FEE, J.K_FEE_REV) and j.status == J.ST_READY]
    if not fees:
        return ""
    t = J.fee_totals(rep.journals)
    lines = [f"FEES BOOKED ({t['count']}): charged {money(t['fees'])}, reversed {money(t['reversals'])}, "
             f"net {money(t['net'])} (EMI interest is not in this total)"]
    for j in sorted(fees, key=lambda x: (x.card, x.real_date)):
        tail = f" (reverses {j.paired_with})" if j.paired_with else ""
        lines.append(f"  {j.card} | {j.real_date:%d %b %Y} | {J.KIND_LABEL[j.kind]} | {j.raw_description[:50] or j.description[:50]} "
                     f"| {money(j.amount)}{tail}")
    return "\n".join(lines)


def large_spend_block(rep: RunReport) -> str:
    flagged = [j for j in rep.journals if j.flag]
    if not flagged:
        return ""
    lines = [f"LARGE SPENDS ({len(flagged)}) of {money(rep.threshold)} or more - {J.LARGE_SPEND_FLAG}. "
             f"Mapped as usual, never moved to an asset account:"]
    for j in sorted(flagged, key=lambda x: (x.real_date, x.card)):
        lines.append(f"  {j.card} | {j.real_date:%d %b %Y} | {j.raw_description[:50]} | {money(j.amount)} | "
                     f"{j.account or 'no account'} | {j.status}")
    return "\n".join(lines)


def statement_block(rep: RunReport) -> str:
    out = []
    if rep.refused:
        out.append("REFUSED - statements that do not tie out (nothing booked from them):\n" + "\n".join(
            f"  {st.card} {st.label}: {t.result}; {t.gap or t.note}" for st, t in rep.refused))
    partly = [s for s in rep.settlements if s.status == "PARTLY SETTLED"]
    if partly:
        out.append("PARTLY SETTLED - paid less than the statement total, nothing booked:\n" + "\n".join(
            f"  {s.target.card} {s.target.label}: {s.reason}" for s in partly))
    other = [s for s in rep.settlements if s.status not in ("SETTLED", "PARTLY SETTLED")]
    if other:
        out.append("PAYMENTS THAT DID NOT SETTLE A STATEMENT IN THIS RUN:\n" + "\n".join(
            f"  {(s.target or s.paid_on).card} {(s.target or s.paid_on).label}: {s.status} - {s.reason}"
            for s in other))
    if rep.unpaid:
        out.append("NOT SETTLED - a later statement exists but shows no payment for it, nothing booked:\n" + "\n".join(
            f"  {st.card} {st.label}" for st in rep.unpaid))
    if rep.awaiting:
        out.append("AWAITING PAYMENT (latest statement, not yet paid; nothing booked):\n" + "\n".join(
            f"  {st.card} {st.label}" for st in rep.awaiting))
    if rep.lines_unmatched:
        out.append("STATEMENT PAYMENT LINES WITH NO BANK PAYMENT IN THE BOOK:\n" + "\n".join(
            f"  {ln.stmt.card} {ln.date:%d %b %Y} {money(ln.amount)}" for ln in rep.lines_unmatched))
    if rep.ambiguous:
        out.append("AMBIGUOUS payments (not matched, never guessed):\n" + "\n".join(
            "  " + (f"{a['line'].stmt.card} line {a['line'].date:%d %b %Y} {money(a['line'].amount)}: "
                    if a["line"] is not None else "")
            + f"{a['reason']} ({', '.join(f'{p.date:%d %b} {money(p.amount)}' for p in a['payments'])})"
            for a in rep.ambiguous))
    if rep.withheld:
        out.append("WITHHELD journals (not in the CSV):\n" + "\n".join(
            f"  {j.card} {j.real_date:%d %b %Y} {j.raw_description[:40] or j.description[:40]} {money(j.amount)}: "
            f"{j.status}{' - ' + j.why if j.why else ''}" for j in rep.withheld))
    if rep.partly_booked:
        out.append("PARTLY BOOKED (never presented as settled):\n" + "\n".join(
            f"  {st.card} {st.label}: {why}" for st, why in rep.partly_booked))
    return "\n\n".join(out)


def emi_block(rep: RunReport) -> str:
    rows = (rep.emi or {}).get("rows") or []
    if not rows:
        return ""
    lines = [f"{J.EMI_UNBOOKED_TITLE.upper()} ({len(rows)}):"]
    for r in rows:
        when = r.date.strftime("%d %b %Y") if r.date else "undated"
        lines.append(f"  {r.card} | {r.period} | {when} | {r.kind} | {r.description[:50]} | "
                     f"{money(r.amount)} {r.direction}")
    interest = (rep.emi or {}).get("interest") or {}
    if interest:
        lines.append("  EMI interest (its own subtotal, not a fee):")
        for (card, fy), v in interest.items():
            lines.append(f"    {card} | FY {fy} | {money(v)}")
    return "\n".join(lines)


def tie_block(rep: RunReport) -> str:
    t = rep.tie
    if not t:
        return ""
    lines = ["Credit Card Payment (CCP) tie-out for the range:",
             f"  before (book, as it stands): {money(t['before'])}",
             f"  these journals: {money(t['effect'])}",
             f"  after: {money(t['after'])}  <- residual",
             f"    of which payments with no statement: {money(t['unmatched'])}",
             f"    of which payments on statements not booked: {money(t['not_booked'])}",
             f"    other: {money(t['other'])}"]
    return "\n".join(lines)


def build_message(rep: RunReport) -> str:
    parts = []
    for block in (red_flag_block(rep), fee_block(rep), large_spend_block(rep), statement_block(rep),
                  emi_block(rep)):
        if block:
            parts.append(block)
    if rep.issues:
        parts.append("Issues:\n" + "\n".join(f"  - {i}" for i in rep.issues))
    n_ready = len(rep.ready)
    if rep.open_items:
        n = len(rep.issues) + len(rep.unmatched_payments) + len(rep.refused) + len(rep.withheld) \
            + len(rep.lines_unmatched) + len(rep.ambiguous) + len(rep.unpaid) + len(rep.partly_booked)
        n += sum(1 for s in rep.settlements if s.status != "SETTLED")
        head = f"Credit card spend booking NOT complete: {max(n, 1)} open item(s) above."
    else:
        head = "Credit card spend booking completed successfully."
    parts.append(head)
    tb = tie_block(rep)
    if tb:
        parts.append(tb)
    parts.append(f"Range run: {rep.label}" + (f" | entity {rep.entity}" if rep.entity else ""))
    ready_n = n_ready
    booked_n = sum(1 for j in rep.journals if j.status == J.ST_BOOKED)
    parts.append(f"{ready_n} journal(s) ready to import" + (f", {booked_n} already booked" if booked_n else "")
                 + f" from {sum(1 for s in rep.settlements if s.status == 'SETTLED')} settled statement(s); "
                 f"{rep.statements_read} statement(s) read in the range")
    if rep.skipped:
        parts.append(f"Skipped: not a statement ({len(rep.skipped)}):\n" + "\n".join(f"  {d}" for d in rep.skipped))
    if rep.duplicates:
        parts.append("Duplicate statements:\n" + "\n".join(f"  {d}" for d in rep.duplicates))
    parts.append(f"Output: {rep.output_csv}" if rep.output_csv else "No import file written: nothing is ready to import.")
    if rep.output_xlsx:
        parts.append(f"Workbook: {rep.output_xlsx}")
    return "\n\n".join(parts)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def write_workbook(rep: RunReport, path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(list(r))
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 70)
        return ws

    sheet("Journals",
          ["Num", "Kind", "Date booked", "Statement date", "Description", "Debit account", "Credit account",
           "Amount", "Card", "Statement", "Status", "Why / note", "Large spend flag", "Confidence", "Match reason"],
          [(j.num, J.KIND_LABEL[j.kind], j.date, j.real_date, j.description,
            j.account if j.side == "W" else rep.ccp_path, rep.ccp_path if j.side == "W" else j.account,
            j.amount, j.card, j.statement, j.status, "; ".join(x for x in (j.why, j.note) if x), j.flag,
            j.confidence, j.match_reason) for j in rep.journals], first=True)
    sheet("Payments without statement",
          ["Date", "Bank account", "Amount", "CRED / direct", "Reason"],
          [(p.date, p.bank_account, p.amount, "CRED" if p.cred else "direct", why)
           for p, why in sorted(rep.unmatched_payments, key=lambda x: (x[0].date, x[0].amount))])
    sheet("Payments matched",
          ["Statement with the payment line", "Line date", "Line amount", "Bank date(s)", "Bank amount", "Match",
           "CRED overage", "Settles statement", "Result"],
          [(f"{s.paid_on.card} {s.paid_on.label}", pr.line.date, pr.line.amount,
            ", ".join(f"{p.date:%d %b %Y}" for p in pr.payments), sum(p.amount for p in pr.payments), pr.kind,
            pr.overage, f"{s.target.card} {s.target.label}" if s.target else "none read", s.status)
           for s in rep.settlements for pr in s.pairings])
    sheet("Large spends",
          ["Date", "Card", "Description", "Amount", "Account", "Flag", "Status"],
          [(j.real_date, j.card, j.raw_description, j.amount, j.account, j.flag, j.status)
           for j in rep.journals if j.flag])
    fees = [j for j in rep.journals if j.kind in (J.K_FEE, J.K_FEE_REV)]
    sheet("Fees", ["Date", "Card", "Kind", "Description", "Amount", "Paired with", "Status"],
          [(j.real_date, j.card, J.KIND_LABEL[j.kind], j.raw_description or j.description, j.amount, j.paired_with,
            j.status) for j in fees])
    emi_rows = (rep.emi or {}).get("rows") or []
    sheet("EMI rows not booked",
          ["Card", "Statement", "Date", "Kind", "Description", "Amount", "Direction"],
          [(r.card, r.period, r.date, r.kind, r.description, r.amount, r.direction) for r in emi_rows])
    sheet("EMI interest",
          ["Card", "Financial year", "EMI interest"],
          [(card, fy, v) for (card, fy), v in ((rep.emi or {}).get("interest") or {}).items()])
    sheet("Statements",
          ["Card", "Statement", "Tie-out", "Settlement", "Detail"],
          [(st.card, st.label, t.result, "", t.gap or t.note) for st, t in rep.refused]
          + [(s.target.card, s.target.label, "PASS", s.status, s.reason) for s in rep.settlements if s.target]
          + [(st.card, st.label, "PASS", "AWAITING PAYMENT", "") for st in rep.awaiting]
          + [(st.card, st.label, "PASS", "NOT SETTLED", "later statement shows no payment") for st in rep.unpaid])
    t = rep.tie or {}
    sheet("CCP tie-out", ["Line", "Amount"],
          [("Before (book as it stands)", t.get("before")), ("These journals", t.get("effect")),
           ("After (residual)", t.get("after")), ("Payments with no statement", t.get("unmatched")),
           ("Payments on statements not booked", t.get("not_booked")), ("Other", t.get("other"))])
    sheet("Issues", ["Issue"], [(i,) for i in rep.issues])
    wb.save(path)
    return path
