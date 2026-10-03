#!/usr/bin/env python3
"""
SBM Bank India statement (OLE-encrypted .xlsx) -> canonical 8-column CSV.

Deterministic, no LLM. The workbook is decrypted in memory with msoffcrypto
(the password is never logged, echoed or persisted). One sheet,
"Transaction-Statement", header in row 1:

  Transaction Timestamp | Narration | RRN | Withdrawal(Dr) | Deposit(Cr) | Balance

There is no opening-balance row and no account-number field. Rows within a day
can be out of time order, so each day is ordered by the balance chain
(prior balance +/- amount == balance). Opening = first chained row's balance
- deposit + withdrawal. The account number is read from interest narrations
("<acct>:Int.Pd:..."). If no consistent chain exists the parse FAILS loudly.
"""
import io
import re
from datetime import date, datetime
from pathlib import Path

from agents.bank_common import password as _password
from agents.bank_contract import BankResult, BankStatementMeta
from agents.canonical_io import run_balance_check, write_canonical_csv, write_sidecar

BANK_KEY = "sbm"
SHEET_NAME = "Transaction-Statement"
HEADERS = ("Transaction Timestamp", "Narration", "RRN", "Withdrawal(Dr)", "Deposit(Cr)", "Balance")
_ACCT_RE = re.compile(r"(\d{6,})\s*:\s*Int\.Pd", re.IGNORECASE)
_EPS = 0.005


class SBMParseError(ValueError):
    """A statement problem (broken chain, bad shape); never carries narration
    text or the password."""


def formats() -> tuple[str, ...]:
    return (".xlsx",)


def _open_workbook(path: str, password):
    """The only place that touches msoffcrypto / the password."""
    import msoffcrypto
    import openpyxl

    with open(path, "rb") as f:
        try:
            office = msoffcrypto.OfficeFile(f)
            if office.is_encrypted():
                office.load_key(password=password or "", verify_password=True)
                buf = io.BytesIO()
                office.decrypt(buf)
            else:
                f.seek(0)
                buf = io.BytesIO(f.read())
        except Exception as e:
            if _password.is_password_error(e):
                raise ValueError(_password.password_error_message(
                    "the statement workbook's own open password", doc_type="xlsx")) from None
            raise ValueError("Could not open this file as an encrypted xlsx workbook: %s"
                             % type(e).__name__) from None
    buf.seek(0)
    try:
        return openpyxl.load_workbook(buf, data_only=True)
    except Exception as e:
        raise ValueError("Decryption succeeded but the result is not a readable xlsx workbook: %s"
                         % type(e).__name__) from None


def _num(v, what, rown):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        raise SBMParseError("row %d: %s is not a number" % (rown, what))
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        raise SBMParseError("row %d: %s is not a number" % (rown, what)) from None


def _iso(v, rown):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or "").strip()[:10]
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        raise SBMParseError("row %d: date is not yyyy-mm-dd" % rown)
    return s


def _read_rows(wb) -> list[dict]:
    if SHEET_NAME not in wb.sheetnames:
        raise SBMParseError("Expected a sheet named %r; found %s" % (SHEET_NAME, wb.sheetnames))
    ws = wb[SHEET_NAME]
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if header is None:
        raise SBMParseError("The %r sheet is empty" % SHEET_NAME)
    cells = [str(c).strip() if c is not None else "" for c in header]
    missing = [h for h in HEADERS if h not in cells]
    if missing:
        raise SBMParseError("Header row is missing column(s): %s" % ", ".join(missing))
    ix = {h: cells.index(h) for h in HEADERS}
    out = []
    for n, row in enumerate(it, start=2):
        row = list(row) + [None] * (len(cells) - len(row))
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            continue
        wd = _num(row[ix["Withdrawal(Dr)"]], "Withdrawal(Dr)", n)
        dp = _num(row[ix["Deposit(Cr)"]], "Deposit(Cr)", n)
        bal = _num(row[ix["Balance"]], "Balance", n)
        if bal is None:
            raise SBMParseError("row %d: Balance is empty" % n)
        if wd and dp:
            raise SBMParseError("row %d: both Withdrawal and Deposit are set" % n)
        rrn = row[ix["RRN"]]
        out.append({
            "src_row": n,
            "date": _iso(row[ix["Transaction Timestamp"]], n),
            "narration": str(row[ix["Narration"]] or "").strip(),
            "rrn": "" if rrn is None else str(rrn).strip(),
            "wd": wd or 0.0, "dp": dp or 0.0, "bal": bal,
        })
    return out


def _delta(r):
    return r["dp"] - r["wd"]


def _order_day(rows, prev):
    """Order one day's rows into a balance chain starting after ``prev``
    (None = unknown opening: the first chained row defines it). Candidates are
    tried in file order, so an already-consistent file is never reshuffled.
    Returns the ordered list or None."""
    failed = set()

    def dfs(remaining, bal, acc):
        if not remaining:
            return acc
        key = (remaining, round(bal, 2) if bal is not None else None)
        if key in failed:
            return None
        for i in sorted(remaining):
            r = rows[i]
            if bal is None or abs(bal + _delta(r) - r["bal"]) < _EPS:
                res = dfs(remaining - {i}, r["bal"], acc + [r])
                if res is not None:
                    return res
        failed.add(key)
        return None

    return dfs(frozenset(range(len(rows))), prev, [])


def _chain(rows):
    by_day = {}
    for r in rows:
        by_day.setdefault(r["date"], []).append(r)
    ordered, prev = [], None
    for day in sorted(by_day):
        got = _order_day(by_day[day], prev)
        if got is None:
            raise SBMParseError(
                "Balance chain is broken on %s (%d row(s)): no ordering makes "
                "prior balance +/- amount equal the balance. Statement rows are "
                "missing or altered." % (day, len(by_day[day])))
        ordered.extend(got)
        prev = got[-1]["bal"]
    return ordered


def _fmt(x):
    return "" if not x else "%.2f" % x


def _extract(path, password):
    wb = _open_workbook(path, password)
    raw = _read_rows(wb)
    if not raw:
        raise SBMParseError("No transactions found in the %r sheet." % SHEET_NAME)
    ordered = _chain(raw)
    first = ordered[0]
    opening = round(first["bal"] - first["dp"] + first["wd"], 2)
    acct = None
    for r in ordered:
        m = _ACCT_RE.search(r["narration"])
        if m:
            acct = m.group(1)
            break
    rows = [{
        "Date": r["date"], "Transaction ID": r["rrn"], "Description": r["narration"],
        "Account": "", "Deposit": _fmt(r["dp"]), "Withdrawal": _fmt(r["wd"]),
        "Balance": "%.2f" % r["bal"], "Currency": "INR",
    } for r in ordered]
    reordered = sum(1 for a, b in zip(ordered, raw) if a is not b)
    return rows, opening, ordered[-1]["bal"], acct, reordered, ordered[0]["date"], ordered[-1]["date"]


def detect(path) -> float:
    p = Path(str(path))
    if p.suffix.lower() not in formats():
        return 0.0
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        if SHEET_NAME in wb.sheetnames:
            hdr = next(wb[SHEET_NAME].iter_rows(values_only=True), ())
            if all(h in [str(c).strip() for c in hdr if c is not None] for h in HEADERS):
                return 0.9
        return 0.1
    except Exception:
        # An OLE-encrypted workbook cannot be peeked into without the password.
        return 0.2


def parse(path, password=None) -> BankResult:
    path = str(path)
    if not Path(path).is_file():
        raise FileNotFoundError(path)
    if Path(path).suffix.lower() not in formats():
        raise ValueError("Unsupported file type: " + Path(path).suffix.lower())
    rows, opening, closing, acct, reordered, p_from, p_to = _extract(path, password)
    check = run_balance_check(rows)
    warnings = []
    if reordered:
        warnings.append("%d row(s) were re-ordered by the balance chain (not dropped)." % reordered)
    meta = BankStatementMeta(
        bank_key=BANK_KEY, account_number=acct, period_from=p_from, period_to=p_to,
        source_format="xlsx", fidelity="exact", password_used=bool(password))
    return BankResult(
        rows=rows, bank_key=BANK_KEY, currency="INR",
        opening_balance=opening, closing_balance=closing,
        balance_check=check, sidecar_path=None, warnings=warnings, meta=meta)


def run(xlsx_path, output_path, config_path=None, model_override=None, pdf_password=None):
    """UI entry point: SBM statement -> canonical CSV (+ summary sidecar)."""
    if not Path(str(xlsx_path)).is_file():
        return "File not found: " + str(xlsx_path)
    try:
        res = parse(xlsx_path, password=pdf_password)
    except Exception as e:
        return "Error processing " + Path(str(xlsx_path)).name + ": " + str(e)
    write_canonical_csv(res.rows, output_path)
    write_sidecar(output_path, "SBM", "derived", res.opening_balance, res.closing_balance,
                  len(res.rows), account_number=res.meta.account_number)
    msg = "SBM (xlsx): extracted %d transactions -> canonical CSV (opening %.2f, closing %.2f, balance check %s)" % (
        len(res.rows), res.opening_balance, res.closing_balance, "ok" if res.balance_check.ok else "FAILED")
    return msg + "".join("\n" + w for w in res.warnings)


class SBMBankSkill:
    bank_key = BANK_KEY

    def formats(self) -> tuple[str, ...]:
        return formats()

    def detect(self, path) -> float:
        return detect(path)

    def parse(self, path, password=None) -> BankResult:
        return parse(path, password=password)


bank_skill = SBMBankSkill()
