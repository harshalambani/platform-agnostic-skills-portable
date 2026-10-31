"""CC-12: a withheld row keeps its account and confidence; the already-booked check also compares
payee words; the CCP tie-out lists what "Other" is made of. Synthetic only."""
from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from agents.skill_cc_spend_booking import journal as J
from agents.skill_cc_spend_booking.bookio import BookTxn, CCP_PATH
from tests.test_cc_spend_booking import (BANK, FOOD, april_rows, basic, book_pay, go, may_rows, stmt, tie, APR, MAY)


def _hand(desc, amount=1500.0, d=date(2025, 4, 5)):
    return BookTxn("hand1", d, desc, "", ((FOOD, amount), (BANK, -amount)))


def _jrows(out):
    ws = load_workbook(out.with_name(out.stem + "_report.xlsx"))["Journals"]
    hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]


def _run(tmp_path, monkeypatch, desc):
    return basic(tmp_path, monkeypatch, txns=[book_pay("p1", 1900, date(2025, 5, 12)), _hand(desc)])


def test_a_same_payee_hand_entry_is_written_and_flagged_and_keeps_its_account(tmp_path, monkeypatch):
    _r, rows, _c, out = _run(tmp_path, monkeypatch, "Swiggy dinner with friends")
    assert "SWIGGY ORDER" in {r["Description"] for r in rows}                # written, never withheld
    j = next(x for x in _jrows(out) if "SWIGGY" in str(x["Description"]).upper())
    assert j["Status"] == J.ST_POSSIBLY
    assert j["Debit account"] == FOOD and j["Confidence"] == "high"         # kept, not blanked


def test_an_unrelated_payee_hand_entry_does_not_withhold(tmp_path, monkeypatch):          # NEGATIVE
    _r, rows, _c, out = _run(tmp_path, monkeypatch, "Electricity bill April")
    assert "SWIGGY ORDER" in {r["Description"] for r in rows}
    j = next(x for x in _jrows(out) if "SWIGGY" in str(x["Description"]).upper())
    assert j["Status"] != J.ST_POSSIBLY


def test_payee_check_is_for_spends_only_a_fee_is_still_withheld_on_amount_and_date():     # NEGATIVE
    f = J.Journal(num="CCSB-f", kind=J.K_FEE, date=date(2025, 4, 10), real_date=date(2025, 4, 10),
                  description="x", amount=100.0, side="W", card="C", statement="s", account="",
                  raw_description="ANNUAL FEE")
    J.mark_booked([f], [BookTxn("t", date(2025, 4, 10), "bank charges", "", ((FOOD, 100.0), (BANK, -100.0)))], CCP_PATH)
    assert f.status == J.ST_POSSIBLY


def test_no_payee_words_on_either_side_still_withholds():                                  # NEGATIVE (safe side)
    j = J.Journal(num="CCSB-a", kind=J.K_SPEND, date=date(2025, 4, 5), real_date=date(2025, 4, 5),
                  description="d", amount=100.0, side="W", card="C", statement="s", account=FOOD,
                  raw_description="12 34")
    J.mark_booked([j], [BookTxn("t", date(2025, 4, 5), "Electricity bill", "", ((FOOD, 100.0), (BANK, -100.0)))], CCP_PATH)
    assert j.status == J.ST_POSSIBLY


def test_other_equals_the_sum_of_its_listed_lines(tmp_path, monkeypatch):
    _r, _rows, _c, out = _run(tmp_path, monkeypatch, "Electricity bill April")
    ws = load_workbook(out.with_name(out.stem + "_report.xlsx"))["CCP tie-out"]
    cells = [(r[0].value, r[1].value) for r in ws.iter_rows(min_row=2)]
    other = next(v for k, v in cells if k == "Other")
    listed = [v for k, v in cells if str(k).startswith("  Other:")]
    assert listed and abs(sum(listed) - other) < 0.005


def test_other_lines_never_list_an_unmatched_payment(tmp_path, monkeypatch):               # NEGATIVE
    s1, s2 = stmt(APR, april_rows()), stmt(MAY, may_rows())
    _r, _rows, _c, out = go(tmp_path, monkeypatch, [s1, s2], [tie(s1, printed=1900.0), tie(s2, printed=300.0)],
                            [book_pay("p1", 1900, date(2025, 5, 12)), book_pay("stray", 777, date(2025, 5, 20))])
    ws = load_workbook(out.with_name(out.stem + "_report.xlsx"))["CCP tie-out"]
    cells = [(r[0].value, r[1].value) for r in ws.iter_rows(min_row=2)]
    other = next(v for k, v in cells if k == "Other")
    listed = [(k, v) for k, v in cells if str(k).startswith("  Other:")]
    assert abs(sum(v for _k, v in listed) - other) < 0.005
    assert not any(abs(v - 777) < 0.005 for _k, v in listed)
