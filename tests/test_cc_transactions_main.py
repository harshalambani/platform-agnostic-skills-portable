"""
tests/test_cc_transactions_main.py -- CC-02: row-wise statement parsing,
per-layout row parsers, kinds, tie-out, duplicates, fees and the result text.

Everything is synthetic. Statement text is supplied by monkeypatching
``read_lines`` (the pdfplumber row-wise extractor), so no real PDF and no real
figures are involved. Amounts and names are made up.
"""
from __future__ import annotations

import importlib.util
import sys
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
SCRIPT = SRC / "agents" / "skill_cc_transactions" / "scripts" / "create_cc_transaction_list.py"
for _p in (str(SRC), str(ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)


def _load_module():
    spec = importlib.util.spec_from_file_location("create_cc_transaction_list", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


m = _load_module()


def test_script_exists():
    assert SCRIPT.exists(), f"Script not found: {SCRIPT}"


def test_usage_when_no_arguments(capsys):
    rc = m.main(["create_cc_transaction_list.py"])
    assert rc == 1


def test_script_no_longer_uses_pdftotext():
    text = SCRIPT.read_text(encoding="utf-8")
    assert "resolve_pdftotext" not in text
    assert "subprocess" not in text
    assert "extract_text" in text


# ---------------------------------------------------------------------------
# Synthetic statements, one per layout.
# ---------------------------------------------------------------------------

AXIS = [
    "01/04/2025 - 30/04/2025 15/05/2025 25/05/2025",
    "5,000.00 Dr 5,000.00 30.00 3,250.50 0.00 0.00 3,220.50 Dr",
    "05/04/2025 EXAMPLE MART MUMBAI DEPT STORES 1,500.50 Dr 30.00 Cr",
    "10/04/2025 BBPS PAYMENT RECEIVED - 123456 PAYMENT 5,000.00 Cr",
    "12/04/2025 EXAMPLE FUEL STATION FUEL 1,750.00 Dr",
    "15/04/2025 CASHBACK CREDIT CASHBACK 30.00 Cr",
]
AXIS_FEES = [
    "01/04/2025 - 30/04/2025 15/05/2025 25/05/2025",
    "5,000.00 Dr 5,000.00 30.00 3,250.50 0.00 590.00 3,810.50 Dr",
    "05/04/2025 EXAMPLE MART MUMBAI DEPT STORES 1,500.50 Dr 30.00 Cr",
    "10/04/2025 BBPS PAYMENT RECEIVED - 123456 PAYMENT 5,000.00 Cr",
    "12/04/2025 EXAMPLE FUEL STATION FUEL 1,750.00 Dr",
    "15/04/2025 CASHBACK CREDIT CASHBACK 30.00 Cr",
    "25/04/2025 ANNUAL FEE FEES 500.00 Dr",
    "25/04/2025 IGST @18% FEES 90.00 Dr",
]
HDFC_OLD = [
    "Statement Date:30/06/2025",
    "Opening Balance Payment/Credits Purchase/Debits Finance Charges Total Dues",
    "10,000.00 5,000.00 1,200.00 0.00 6,200.00",
    "05/06/2025 12:30:45 EXAMPLE FOODS BANGALORE 1,200.00",
    "10/06/2025 BPPY CC PAYMENT (Ref# 12345) 5,000.00Cr",
]
HDFC_NEW = [
    "Billing Period 01 Sep, 2025 - 30 Sep, 2025",
    "PREVIOUS STATEMENT DUES PAYMENTS/CREDITS RECEIVED PURCHASES/DEBIT FINANCE CHARGES TOTAL AMOUNT DUE",
    "C10,000.00 C10,000.00 + C4,850.00 + C0.00 =",
    "_ C4,850.00",
    "03/09/2025| 14:22 EXAMPLE ONLINE + 219 C 4,500.00 l",
    "05/09/2025| 09:10 EXAMPLE CABS C 350.00 l",
    "10/09/2025| 18:00 BPPY CC PAYMENT + C 10,000.00 l",
]
ICICI = [
    "Statement period : September 1, 2025 to September 30, 2025",
    "Previous Balance Purchases / Charges Cash Advances Payments / Credits",
    "`8,000.00 `3,850.20 `0.00 `8,000.00",
    "Total Amount due `3,850.20",
    "05/09/2025 1234567890 EXAMPLE PAY IN 12 2,500.00",
    "07/09/2025 1234567891 EXAMPLE STREAM US -5 15.99 USD 1,350.20",
    "10/09/2025 1234567892 BBPS Payment received 0 8,000.00 CR",
]
HSBC = [
    "01 SEP 2025 To 30 SEP 2025",
    "OPENING BALANCE 4,000.00",
    "05SEP EXAMPLE STORES MUMBAI IN 2,000.00",
    "12SEP BBPS PMT 4,000.00 CR",
    "30SEP NET OUTSTANDING BALANCE 2,000.00",
]
HSBC_DEC_JAN = [
    "05 DEC 2025 To 04 JAN 2026",
    "OPENING BALANCE 1,000.00",
    "28DEC EXAMPLE CAFE MUMBAI IN 300.00",
    "02JAN EXAMPLE BOOKS DELHI IN 200.00",
    "04JAN NET OUTSTANDING BALANCE 1,500.00",
]
SBM = [
    "Statement Date : 30-Sep-2025",
    "Opening Balance Purchase/Debits Finance Charges Payment/Credit Total Amount Due",
    "2,000.00 1,000.00 0.00 2,000.00 1,000.00",
    "5-Sep-2025 EXAMPLE SHOP 1,000.00 -",
    "20-Sep-2025 Repayment - 2,000.00",
]
YES = [
    "Statement Period: 01/09/2025 To 30/09/2025",
    "Previous Balance : Rs. 3,000.00",
    "Current Purchases / Cash Advance : Rs. 1,500.00",
    "Payment & Credits Received : Rs. 3,000.00",
    "Total Amount Due: Rs. 1,500.00",
    "05/09/2025 EXAMPLE MARKET BANGALORE 1,500.00 Dr",
    "15/09/2025 PAYMENT RECEIVED BBPS - Ref No: 12345 3,000.00 Cr",
]

CASES = [
    ("Axis", "Flipkart", AXIS, 4),
    ("HDFC", "Regalia", HDFC_OLD, 2),
    ("HDFC", "TataNeu", HDFC_NEW, 3),
    ("ICICI", "Sapphiro", ICICI, 3),
    ("HSBC", "Premier", HSBC, 2),
    ("SBM", "Global", SBM, 2),
    ("YES", "Bank", YES, 2),
]


def _parse(bank, card, lines):
    return m.parse_statement(lines, bank, card, "x.pdf")


@pytest.mark.parametrize("bank,card,lines,count", CASES)
def test_each_layout_parses_every_row_dated_and_ties_out(bank, card, lines, count):
    st = _parse(bank, card, lines)
    assert not st.parsed.errors, st.parsed.errors
    assert len(st.parsed.rows) == count
    assert all(isinstance(r.date, datetime) for r in st.parsed.rows)
    tie = m.tie_out(st)
    assert tie.result == "PASS", (tie.note, tie.difference)


def test_axis_first_amount_is_the_transaction_and_cashback_is_never_a_row():
    st = _parse("Axis", "Flipkart", AXIS)
    mart = next(r for r in st.parsed.rows if "EXAMPLE MART" in r.description)
    assert mart.amount == 1500.50 and mart.direction == "Dr" and mart.kind == "spend"
    # NEGATIVE: the second (informational) 30.00 on the mart line is not a row
    assert [r.amount for r in st.parsed.rows].count(30.0) == 1
    cb = next(r for r in st.parsed.rows if "CASHBACK CREDIT" in r.description)
    assert cb.kind == "cashback" and cb.direction == "Cr"


def test_hdfc_new_coins_on_a_debit_stay_a_debit_and_plus_c_is_a_credit():
    st = _parse("HDFC", "TataNeu", HDFC_NEW)
    by = {r.description: r for r in st.parsed.rows}
    online = next(r for d, r in by.items() if "EXAMPLE ONLINE" in d)
    assert online.direction == "Dr" and online.amount == 4500.0
    pay = next(r for d, r in by.items() if "PAYMENT" in d)
    assert pay.direction == "Cr" and pay.kind == "payment" and pay.amount == 10000.0


def test_icici_international_row_takes_the_inr_amount():
    st = _parse("ICICI", "Sapphiro", ICICI)
    intl = next(r for r in st.parsed.rows if "STREAM" in r.description)
    assert intl.amount == 1350.20      # NEGATIVE: not the 15.99 USD figure
    assert intl.direction == "Dr"


def test_hsbc_dec_row_on_a_january_statement_belongs_to_the_earlier_year():
    st = _parse("HSBC", "Premier", HSBC_DEC_JAN)
    dates = sorted(r.date for r in st.parsed.rows)
    assert dates == [datetime(2025, 12, 28), datetime(2026, 1, 2)]
    assert m.tie_out(st).result == "PASS"


def test_sbm_debit_and_credit_columns():
    st = _parse("SBM", "Global", SBM)
    got = sorted((r.direction, r.kind) for r in st.parsed.rows)
    assert got == [("Cr", "payment"), ("Dr", "spend")]


# ---------------------------------------------------------------------------
# Kinds
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("desc,direction,kind", [
    ("BBPS PAYMENT RECEIVED - 99", "Cr", "payment"),
    ("BPPY CC PAYMENT", "Cr", "payment"),
    ("CASHBACK CREDIT", "Cr", "cashback"),
    ("CREDIT BALANCE REFUND", "Dr", "payment"),
    ("ANNUAL FEE", "Dr", "fee"),
    ("LATE PAYMENT FEE", "Dr", "fee"),
    ("IGST @18%", "Dr", "fee"),
    ("EXAMPLE MERCHANT REFUND", "Cr", "refund"),
    ("EXAMPLE SHOP", "Dr", "spend"),
])
def test_kind_classification(desc, direction, kind):
    assert m.classify(desc, direction) == kind


def test_a_fee_is_never_classified_as_spend():
    for desc in ("ANNUAL FEE", "LATE PAYMENT FEE", "FINANCE CHARGES", "IGST @18%", "CGST", "SGST"):
        assert m.classify(desc, "Dr") == "fee", desc


def test_a_payment_is_never_a_refund_or_spend():
    assert m.classify("BBPS PAYMENT RECEIVED - 1", "Cr") not in ("refund", "spend")


# ---------------------------------------------------------------------------
# Tie-out
# ---------------------------------------------------------------------------

def test_a_dropped_row_makes_the_tie_out_fail_not_pass():
    st = _parse("Axis", "Flipkart", [ln for ln in AXIS if "EXAMPLE FUEL" not in ln])
    assert m.tie_out(st).result == "FAIL"


def test_a_doubled_row_fails():
    st = _parse("YES", "Bank", YES + [YES[5]])
    assert m.tie_out(st).result == "FAIL"


def test_yes_with_a_missing_label_is_not_available_never_pass():
    st = _parse("YES", "Bank", [ln for ln in YES if not ln.startswith("Total Amount Due")])
    assert m.tie_out(st).result == "NOT AVAILABLE"


def test_icici_without_a_total_is_not_available_never_pass():
    st = _parse("ICICI", "Sapphiro", [ln for ln in ICICI if not ln.startswith("Total Amount due")])
    assert m.tie_out(st).result == "NOT AVAILABLE"


def test_printed_summary_that_does_not_add_up_fails():
    bad = [ln.replace("3,220.50 Dr", "9,999.00 Dr") for ln in AXIS]
    assert m.tie_out(_parse("Axis", "Flipkart", bad)).result == "FAIL"


# ---------------------------------------------------------------------------
# Whole-folder runs
# ---------------------------------------------------------------------------

def _folder(tmp_path, monkeypatch, files):
    for (folder, name) in files:
        d = tmp_path / "in" / folder
        d.mkdir(parents=True, exist_ok=True)
        (d / name).write_bytes(b"%PDF-1.4 synthetic")
    table = {name: lines for (_f, name), lines in files.items()}
    monkeypatch.setattr(m, "read_lines", lambda p: table[Path(p).name])
    return tmp_path / "in"


def _run(tmp_path, monkeypatch, files, start=date(2025, 4, 1), end=date(2026, 3, 31),
         label="FY 2025-26"):
    src = _folder(tmp_path, monkeypatch, files)
    out = tmp_path / "out.xlsx"
    return m.run_extraction(src, out, start, end, label), out


def test_clean_run_says_successfully_and_names_the_range(tmp_path, monkeypatch):
    res, out = _run(tmp_path, monkeypatch, {("Axis-Flipkart", "a.pdf"): AXIS})
    assert res.issues == []
    assert "completed successfully" in res.message
    assert "FY 2025-26" in res.message
    assert out.exists()


def test_excel_has_real_dates_kind_and_tieout_sheet(tmp_path, monkeypatch):
    res, out = _run(tmp_path, monkeypatch, {("Axis-Flipkart", "a.pdf"): AXIS,
                                            ("YES-Bank", "y.pdf"): YES})
    wb = load_workbook(out)
    assert {"Transactions", "Summary", "Tie-out", "Fees"} <= set(wb.sheetnames)
    rows = list(wb["Transactions"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 6
    header = [c.value for c in wb["Transactions"][1]]
    di, ki = header.index("Date"), header.index("Kind")
    assert all(isinstance(r[di], datetime) for r in rows)   # NEGATIVE: never text / blank
    assert {r[ki] for r in rows} <= {"spend", "payment", "refund", "cashback", "fee"}
    cells = [c.value for r in wb["Summary"].iter_rows() for c in r]
    assert any("FY 2025-26" in str(v) for v in cells)


def test_the_same_statement_twice_is_counted_once_and_listed(tmp_path, monkeypatch):
    res, _ = _run(tmp_path, monkeypatch, {("Axis-Flipkart", "a.pdf"): AXIS,
                                          ("Axis-Flipkart", "a copy.pdf"): AXIS})
    assert len(res.rows) == 4            # NEGATIVE: not 8
    assert len(res.duplicates) == 1
    assert len(res.ties) == 1


def test_range_keeps_overlapping_statements_only(tmp_path, monkeypatch):
    res, _ = _run(tmp_path, monkeypatch, {("Axis-Flipkart", "a.pdf"): AXIS},
                  start=date(2025, 4, 30), end=date(2025, 5, 31), label="x")
    assert len(res.rows) == 4            # April statement ends 30 Apr: overlaps
    res2, _ = _run(tmp_path / "b", monkeypatch, {("Axis-Flipkart", "a.pdf"): AXIS},
                   start=date(2025, 5, 1), end=date(2025, 5, 31), label="x")
    assert res2.rows == []               # NEGATIVE: outside the range


def test_a_failing_tie_out_is_never_called_successful(tmp_path, monkeypatch):
    broken = [ln for ln in AXIS if "EXAMPLE FUEL" not in ln]
    res, _ = _run(tmp_path, monkeypatch, {("Axis-Flipkart", "a.pdf"): broken})
    assert "completed successfully" not in res.message
    assert "completed with" in res.message and "issue" in res.message
    assert any("FAIL" in i for i in res.issues)


def test_not_available_is_an_issue_too(tmp_path, monkeypatch):
    res, _ = _run(tmp_path, monkeypatch,
                  {("YES-Bank", "y.pdf"): [ln for ln in YES if not ln.startswith("Total Amount Due")]})
    assert "completed successfully" not in res.message
    assert any("NOT AVAILABLE" in i for i in res.issues)


def test_a_pdf_with_zero_rows_is_an_issue(tmp_path, monkeypatch):
    res, _ = _run(tmp_path, monkeypatch,
                  {("Axis-Flipkart", "a.pdf"): ["01/04/2025 - 30/04/2025 15/05/2025 25/05/2025", "nothing"]})
    assert "completed successfully" not in res.message
    assert any("no transactions parsed" in i for i in res.issues)


def test_an_undated_row_is_an_issue_and_never_a_blank_date(tmp_path, monkeypatch):
    bad = AXIS + ["31/02/2025 EXAMPLE IMPOSSIBLE DATE SHOP 10.00 Dr"]
    res, _ = _run(tmp_path, monkeypatch, {("Axis-Flipkart", "a.pdf"): bad})
    assert "completed successfully" not in res.message
    assert res.issues
    assert all(r.date is not None for r in res.rows)


def test_unsupported_bank_folder_is_reported(tmp_path, monkeypatch):
    res, _ = _run(tmp_path, monkeypatch, {("SBI-Elite", "s.pdf"): ["whatever"]})
    assert res.issues
    assert "completed successfully" not in res.message


# --- fees (8b) --------------------------------------------------------------

def test_fees_are_listed_first_and_each_counts_as_an_issue(tmp_path, monkeypatch):
    res, out = _run(tmp_path, monkeypatch, {("Axis-Flipkart", "a.pdf"): AXIS_FEES})
    assert res.message.startswith("CARD FEES CHARGED")
    assert "ANNUAL FEE" in res.message and "IGST @18%" in res.message   # GST listed with the fee
    assert "completed successfully" not in res.message
    assert sum("fee charged" in i for i in res.issues) == 2
    assert len(res.fees) == 2
    wb = load_workbook(out)
    assert len(list(wb["Fees"].iter_rows(min_row=2, values_only=True))) == 2


def test_exactly_one_fee_never_says_successfully_and_shows_the_block_first(tmp_path, monkeypatch):
    one = [ln for ln in AXIS_FEES if "IGST" not in ln]
    one[1] = one[1].replace("590.00 3,810.50", "500.00 3,720.50")
    res, _ = _run(tmp_path, monkeypatch, {("Axis-Flipkart", "a.pdf"): one})
    assert res.message.startswith("CARD FEES CHARGED")
    assert "completed successfully" not in res.message
    assert "completed with 1 issue" in res.message


def _fee_row(desc, amt, direction, kind, day, bank="Axis", card="Flipkart"):
    return m.Row(bank, card, datetime(2025, 4, day), desc, amt, direction, kind)


def _stmt(bank, card, rows):
    return m.Statement(bank, card, "a.pdf", None, "", m.Parsed(rows=rows))


def test_a_later_same_card_same_amount_fee_credit_is_shown_as_reversed():
    fee = _fee_row("ANNUAL FEE", 500.0, "Dr", "fee", 25)
    rev = _fee_row("ANNUAL FEE REVERSAL", 500.0, "Cr", "refund", 30)
    fees = m.collect_fees([_stmt("Axis", "Flipkart", [fee, rev])])
    assert len(fees) == 1 and fees[0].reversed_by is rev
    assert "reversed" in m.fee_block(fees)


def test_a_merchant_refund_of_the_same_amount_is_not_a_fee_reversal():
    fee = _fee_row("ANNUAL FEE", 500.0, "Dr", "fee", 25)
    refund = _fee_row("EXAMPLE MART REFUND", 500.0, "Cr", "refund", 30)
    fees = m.collect_fees([_stmt("Axis", "Flipkart", [fee, refund])])
    assert fees[0].reversed_by is None
    assert "reversed" not in m.fee_block(fees)


# ---------------------------------------------------------------------------
# agent wiring and the file name
# ---------------------------------------------------------------------------

def test_agent_returns_the_script_message_and_rejects_a_bad_range(tmp_path, monkeypatch):
    from agents.skill_cc_transactions import agent
    monkeypatch.setattr(m, "read_lines", lambda p: AXIS)
    monkeypatch.setattr(agent, "_load_script", lambda: m)
    d = tmp_path / "in" / "Axis-Flipkart"
    d.mkdir(parents=True)
    (d / "a.pdf").write_bytes(b"%PDF")
    msg = agent.run(str(tmp_path / "in"), str(tmp_path / "o.xlsx"), period="FY 2025-26")
    assert "completed successfully" in msg
    bad = agent.run(str(tmp_path / "in"), str(tmp_path / "o2.xlsx"),
                    period="Custom date range (use dates below)", custom_start="2025-05-01")
    assert bad.startswith("ERROR")
    missing = agent.run(str(tmp_path / "nope"), str(tmp_path / "o3.xlsx"), period="FY 2025-26")
    assert missing.startswith("ERROR") and not (tmp_path / "o3.xlsx").exists()


def test_period_slug_forms():
    from agents.period_picker import period_slug
    assert period_slug("FY 2025-26") == "FY2025-26"
    assert period_slug("FY 2025-26 Q1") == "FY2025-26-Q1"
    assert period_slug("2025-04-01 to 2025-06-30") == "2025-04-01_to_2025-06-30"


def test_excel_file_name_carries_the_period(monkeypatch, tmp_path):
    from agents import registry
    from ui.tabs import _generic as generic_mod
    skill = registry.get("CC Transactions")
    got = {}
    monkeypatch.setattr("agents.registry.load_run_function",
                        lambda s: (lambda **kw: got.update(kw) or "done"))
    monkeypatch.setattr("ui._config.output_dir", lambda: tmp_path / "out")
    (tmp_path / "out").mkdir()
    folder = tmp_path / "pdfs"
    folder.mkdir()
    handler = generic_mod._make_run_handler(skill)
    list(handler(str(folder), "FY 2025-26 Q1", "", "", "m"))
    assert "FY2025-26-Q1" in Path(got["output_excel"]).name
    got.clear()
    list(handler(str(folder), "Custom date range (use dates below)", "2025-05-01", "2025-05-31", "m"))
    assert "2025-05-01_to_2025-05-31" in Path(got["output_excel"]).name
