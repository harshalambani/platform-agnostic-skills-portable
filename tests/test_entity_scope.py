"""
tests/test_entity_scope.py -- UI-16: an output-file picker offers only the
selected entity's own files.

Regression: on the Individual's partner form the 26AS dropdown pre-selected
the HUF's workbook, because "newest file wins". Everything here is synthetic
(made-up names and PANs).
"""
import os
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

ROOT = Path(__file__).resolve().parent.parent
for _p in (ROOT, ROOT / "src"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

from agents import entity_scope  # noqa: E402
from agents.registry import discover  # noqa: E402
from ui.tabs import _generic  # noqa: E402

IND = {"name": "Test Individual", "pan": "AAAAA0000A"}
HUF = {"name": "Test Family HUF", "pan": "BBBBB1111B"}


def _make_26as(path: Path, name, pan, fy="2025-26", mtime=None, with_pan=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Part I"
    ws.cell(row=1, column=1, value="PART I")
    parts = [f"Assessee Name: {name}"]
    if with_pan:
        parts.append(f"PAN: {pan}")
    parts.append(f"Financial Year: {fy}")
    ws.cell(row=2, column=1, value="  |  ".join(parts))
    wb.save(path)
    if mtime:
        os.utime(path, (mtime, mtime))
    return path


@pytest.fixture
def world(tmp_path, monkeypatch):
    data = tmp_path / "Data"
    (data / "itr").mkdir(parents=True)
    (data / "itr" / "entities.yaml").write_text(yaml.safe_dump({
        "TEST-IND": dict(IND, status="Individual"),
        "TEST-HUF": dict(HUF, status="HUF"),
    }), encoding="utf-8")
    out = tmp_path / "out"
    out.mkdir()
    monkeypatch.setattr(_generic._config, "data_root_dir", lambda: data)
    monkeypatch.setattr(_generic._config, "output_dir", lambda: out)
    return out


def _picker():
    sk = next(s for s in discover(refresh=True)
              if Path(s.manifest_path).parent.name == "skill_partner_comp_recon")
    return next(i for i in sk.inputs if i.name == "xlsx_26as")


def _vals(choices):
    return [v for _, v in choices]


def test_huf_26as_never_appears_on_the_individuals_form(world):          # NEGATIVE
    ind = _make_26as(world / "1-ind-26AS.xlsx", IND["name"], IND["pan"], mtime=1_000_000)
    huf = _make_26as(world / "2-huf-26AS.xlsx", HUF["name"], HUF["pan"], mtime=2_000_000)
    choices, _value, _msg = _generic._scoped_picker_state(_picker(), "TEST-IND")
    assert str(huf) not in _vals(choices)
    assert str(ind) in _vals(choices)


def test_newer_other_entity_file_never_becomes_the_default(world):       # NEGATIVE
    ind = _make_26as(world / "1-ind-26AS.xlsx", IND["name"], IND["pan"], mtime=1_000_000)
    _make_26as(world / "2-huf-26AS.xlsx", HUF["name"], HUF["pan"], mtime=9_000_000)
    _c, value, _m = _generic._scoped_picker_state(_picker(), "TEST-IND")
    assert value == str(ind)          # the entity's own newest, not the newest overall


def test_default_is_the_entitys_own_newest(world):
    _make_26as(world / "old-26AS.xlsx", IND["name"], IND["pan"], mtime=1_000_000)
    new = _make_26as(world / "new-26AS.xlsx", IND["name"], IND["pan"], mtime=5_000_000)
    _c, value, _m = _generic._scoped_picker_state(_picker(), "TEST-IND")
    assert value == str(new)


def test_no_file_for_the_entity_offers_none_and_says_why(world):         # NEGATIVE
    huf = _make_26as(world / "huf-26AS.xlsx", HUF["name"], HUF["pan"])
    choices, value, msg = _generic._scoped_picker_state(_picker(), "TEST-IND")
    assert _vals(choices) == [""]     # only "(none)"; the HUF file is not offered
    assert str(huf) not in _vals(choices)
    assert value == ""
    assert "No 26AS" in msg and "26AS Convert" in msg


def test_none_is_a_real_deliberate_choice(world):
    _make_26as(world / "a-26AS.xlsx", IND["name"], IND["pan"])
    choices, _v, _m = _generic._scoped_picker_state(_picker(), "TEST-IND")
    assert ("(none)", "") in choices


def test_no_entity_selected_offers_nothing_but_none(world):              # NEGATIVE
    _make_26as(world / "a-26AS.xlsx", IND["name"], IND["pan"])
    choices, value, msg = _generic._scoped_picker_state(_picker(), None)
    assert _vals(choices) == [""] and value == "" and "entity" in msg.lower()


def test_match_by_name_when_the_workbook_has_no_pan(world):
    f = _make_26as(world / "n-26AS.xlsx", "TEST INDIVIDUAL", IND["pan"], with_pan=False)
    choices, _v, _m = _generic._scoped_picker_state(_picker(), "TEST-IND")
    assert str(f) in _vals(choices)


def test_same_name_different_pan_is_not_a_match(world):                  # NEGATIVE
    f = _make_26as(world / "n-26AS.xlsx", IND["name"], "ZZZZZ9999Z")
    choices, _v, _m = _generic._scoped_picker_state(_picker(), "TEST-IND")
    assert str(f) not in _vals(choices)


def test_unreadable_workbook_is_nobodys(world):                          # NEGATIVE
    bad = world / "junk-26AS.xlsx"
    bad.write_bytes(b"not a workbook")
    choices, _v, _m = _generic._scoped_picker_state(_picker(), "TEST-IND")
    assert str(bad) not in _vals(choices)


def test_fy_filter_when_a_fy_is_given(world):
    a = _make_26as(world / "a-26AS.xlsx", IND["name"], IND["pan"], fy="2024-25")
    b = _make_26as(world / "b-26AS.xlsx", IND["name"], IND["pan"], fy="2025-26")
    ch = entity_scope.filter_choices(
        [(p.name, str(p)) for p in (a, b)], "TEST-IND",
        _generic._entities_yaml_path(), fy="2025-26")
    assert _vals(ch) == [str(b)]


# ---- run-time refusal ------------------------------------------------------

def test_another_entitys_26as_is_refused_with_a_clear_message(world):    # NEGATIVE
    huf = _make_26as(world / "huf-26AS.xlsx", HUF["name"], HUF["pan"])
    verdict, why = entity_scope.check_26as_owner(huf, "TEST-IND", IND["name"], IND["pan"])
    assert verdict == "mismatch"
    assert "Test Family HUF" in why and "TEST-IND" in why


def test_owner_check_accepts_the_entitys_own_26as(world):
    own = _make_26as(world / "own-26AS.xlsx", IND["name"], IND["pan"])
    assert entity_scope.check_26as_owner(own, "TEST-IND", IND["name"], IND["pan"]) == ("ok", "")


def test_owner_check_unknown_when_workbook_states_no_identity(world):
    f = world / "blank-26AS.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Part I"
    wb.save(f)
    verdict, why = entity_scope.check_26as_owner(f, "TEST-IND", IND["name"], IND["pan"])
    assert verdict == "unknown" and "TEST-IND" in why


def test_partner_run_is_wired_to_the_owner_check():
    src = (ROOT / "src/agents/skill_partner_comp_recon/agent.py").read_text(encoding="utf-8")
    assert "check_owner_26as(" in src and '"mismatch"' in src


def test_26as_journal_is_wired_to_the_owner_check():
    src = (ROOT / "src/agents/skill_26as_journal/agent.py").read_text(encoding="utf-8")
    assert "check_26as_owner(" in src and '"mismatch"' in src


# ---- manifest sweep --------------------------------------------------------

# Every `match:` glob in a skill.yaml is either entity-scoped or listed here.
UNSCOPED_MATCHES = {
    # KRC workbooks are named after the client report and carry no assessee
    # identity that maps to a registry entity: deliberately unscoped.
    ("skill_krc_gnucash", "*-KRC-Bills-Recon.xlsx"),
    ("skill_krc_recon", "*-KRC-Ledger.xlsx"),
}


def test_every_26as_picker_is_entity_scoped_and_every_other_match_is_listed():
    seen_scoped = set()
    for s in discover(refresh=True):
        folder = Path(s.manifest_path).parent.name
        for i in s.inputs:
            if i.type != "output_file" or not i.match:
                continue
            if i.match in entity_scope.SCOPABLE_MATCHES:
                assert i.entity_from, f"{folder}.{i.name} lists 26AS workbooks without entity_from"
                assert any(j.name == i.entity_from and j.type == "select" for j in s.inputs)
                seen_scoped.add(folder)
            else:
                assert (folder, i.match) in UNSCOPED_MATCHES, (
                    f"{folder}.{i.name} match {i.match!r} is neither scoped nor listed")
    assert seen_scoped == {"skill_partner_comp_recon", "skill_26as_journal", "skill_ais_reconcile"}


def test_entity_from_requires_a_scopable_match():
    for s in discover(refresh=True):
        for i in s.inputs:
            if i.entity_from:
                assert i.type == "output_file" and i.match in entity_scope.SCOPABLE_MATCHES
