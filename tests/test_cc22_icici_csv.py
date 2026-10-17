"""tests/test_cc22_icici_csv.py -- CC-22: ICICI statement CSV exports.

Everything is synthetic: made-up card numbers, merchants and amounts.
"""
from __future__ import annotations

import importlib.util
import sys
from datetime import date, datetime
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
SCRIPT = SRC / "agents" / "skill_cc_transactions" / "scripts" / "create_cc_transaction_list.py"
SORT_SCRIPT = SRC / "agents" / "skill_cc_sort" / "scripts" / "extract_sort_cc_pdfs.py"
for _p in (str(SRC), str(ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from agents.skill_cc_sort import completeness as comp  # noqa: E402
from agents.skill_cc_sort import icici_csv as ic  # noqa: E402


def _load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


m = _load("create_cc_transaction_list_cc22", SCRIPT)

HEADER = "Date,Sr.No.,Transaction Details,Reward Point Header,Intl.Amount,Amount(in Rs),BillingAmountSign"
MC = "5100 XXXX XXXX 1111"
AMEX = "3700 XXXX XXXX 222"
SOLO = "4000 XXXX XXXX 3333"


def _csv(blocks, header=HEADER, footer=True, sign_override=None):
    """blocks: [(mask, [(date, desc, amount_text)])]"""
    lines = ["Statement,Export", "Account,Synthetic", "Name,Test", "", "", "Preamble note", header]
    n = 0
    for mask, rows in blocks:
        lines.append(f'"{mask}"')
        for d, desc, amt in rows:
            n += 1
            sign = amt if sign_override is None else sign_override
            lines.append(f'{d},{n},{desc},5,0.00,"{amt}","{sign}"')
    if footer:
        lines.append('"Important notes"')
        lines.append('"Please read the following"')
        for i in range(1, 10):
            lines.append(f'{i},"Synthetic footer note {i} 01-Jan-25 100.00",x')
    return "\n".join(lines) + "\n"


def _write(tmp, folder, name, text):
    d = Path(tmp) / folder
    d.mkdir(parents=True, exist_ok=True)
    p = d / name
    p.write_text(text, encoding="utf-8")
    return p


SAPPH = [
    (MC, [("05-Apr-25", "SHOP ONE MUMBAI", "1,200.50"),
          ("07-Apr-25", "PAYMENT RECEIVED THANK YOU", "-5,000.00")]),
    (AMEX, [("06-Apr-25", "CAFE TWO PUNE", "300.00")]),
]


# ---- parser ----

def test_parses_rows_with_direction_from_amount(tmp_path):
    p = _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH))
    st = ic.parse_icici_csv(p)
    assert len(st.rows) == 3
    assert [r.direction for r in st.rows] == ["Dr", "Cr", "Dr"]
    assert all(r.amount > 0 for r in st.rows)
    assert st.rows[0].amount == 1200.50
    assert st.first == date(2025, 4, 5) and st.last == date(2025, 4, 7)
    assert (st.debits, st.credits) == (2, 1)
    assert {r.network for r in st.rows} == {"Mastercard", "Amex"}


def test_footer_notes_never_become_rows(tmp_path):
    with_footer = ic.parse_icici_csv(_write(tmp_path, "a", "a.csv", _csv(SAPPH)))
    without = ic.parse_icici_csv(_write(tmp_path, "b", "b.csv", _csv(SAPPH, footer=False)))
    assert len(with_footer.rows) == len(without.rows) == 3


def test_dated_row_after_footer_refuses_file(tmp_path):
    text = _csv(SAPPH) + '01-Apr-25,99,LATE ROW,5,0.00,"10.00","10.00"\n'
    with pytest.raises(ic.CsvRefused):
        ic.parse_icici_csv(_write(tmp_path, "a", "a.csv", text))


def test_header_mismatch_is_refused_not_half_read(tmp_path):
    bad = HEADER.replace("BillingAmountSign", "Sign")
    p = _write(tmp_path, "a", "a.csv", _csv(SAPPH, header=bad))
    with pytest.raises(ic.CsvRefused):
        ic.parse_icici_csv(p)
    six = ",".join(HEADER.split(",")[:6])
    with pytest.raises(ic.CsvRefused):
        ic.parse_icici_csv(_write(tmp_path, "b", "b.csv", _csv(SAPPH, header=six)))


def test_billing_sign_differing_from_amount_refuses_file(tmp_path):
    p = _write(tmp_path, "a", "a.csv", _csv(SAPPH, sign_override="9.99"))
    with pytest.raises(ic.CsvRefused) as e:
        ic.parse_icici_csv(p)
    assert "BillingAmountSign" in str(e.value)


def test_card_identity_from_masks_not_filename_or_rows(tmp_path):
    assert ic.identify_card([MC, AMEX])[0] == "Sapphiro"
    assert ic.identify_card([SOLO])[0] == "Amazon"
    assert ic.identify_card([MC, MC])[0] == "Amazon"      # one distinct card
    assert ic.identify_card([])[0] is None
    assert ic.identify_card([MC, AMEX, SOLO])[0] is None
    one_row_two_cards = _csv([(MC, [("05-Apr-25", "X", "1.00")]), (AMEX, [])])
    p = _write(tmp_path, "amazon-named", "amazon_icici.csv", one_row_two_cards)
    assert ic.card_folder(p) == ("ICICI", "Sapphiro")
    many_rows_one_card = _csv([(SOLO, [("05-Apr-25", f"M{i}", "1.00") for i in range(9)])])
    p2 = _write(tmp_path, "x", "sapphiro.csv", many_rows_one_card)
    assert ic.card_folder(p2) == ("ICICI", "Amazon")


def test_unreadable_csv_routes_unknown(tmp_path):
    p = _write(tmp_path, "a", "a.csv", _csv(SAPPH, header="a,b,c"))
    assert ic.card_folder(p) == ("Unknown", "Unknown")


# ---- Credit Card Transactions ----

def _run(tmp, **kw):
    return m.run_extraction(tmp, Path(tmp) / "out.xlsx", kw.get("start"), kw.get("end"))


def test_both_networks_land_on_one_account_with_network_column(tmp_path):
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH))
    res = _run(tmp_path)
    assert len(res.rows) == 3
    assert {r.card for r in res.rows} == {"ICICI-Sapphiro"}        # never a second account
    assert {r.network for r in res.rows} == {"Mastercard", "Amex"}
    from openpyxl import load_workbook
    ws = load_workbook(res.output)["Transactions"]
    head = [c.value for c in ws[1]]
    assert "Network" in head
    nets = {row[head.index("Network")].value for row in ws.iter_rows(min_row=2)}
    assert nets == {"Mastercard", "Amex"}
    cards = {(row[0].value, row[1].value) for row in ws.iter_rows(min_row=2)}
    assert cards == {("ICICI", "Sapphiro")}


def test_payment_row_is_never_a_spend(tmp_path):
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH))
    res = _run(tmp_path)
    pay = [r for r in res.rows if "PAYMENT" in r.description][0]
    assert pay.direction == "Cr" and pay.kind != "spend"
    assert sum(1 for r in res.rows if r.kind == "spend") == 2
    assert all(r.direction == "Dr" for r in res.rows if r.kind == "spend")


def test_output_says_no_totals_and_reports_coverage_and_counts(tmp_path):
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH))
    res = _run(tmp_path)
    assert "no statement totals" in res.message
    assert "2 debit(s), 1 credit(s)" in res.message
    assert "05 Apr 2025 to 07 Apr 2025" in res.message
    assert res.csv_ties and res.csv_ties[0].result == m.CSV_NO_TOTALS
    assert res.ties == []                       # never passed through a balance tie-out
    assert not any("tie-out" in i for i in res.issues)


def test_footer_notes_are_not_in_the_output(tmp_path):
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH))
    res = _run(tmp_path)
    assert not any("footer" in r.description.lower() for r in res.rows)
    assert len(res.rows) == 3


def test_refused_header_gives_issue_and_no_rows(tmp_path):
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH, header=HEADER.replace("Sr.No.", "No")))
    res = _run(tmp_path)
    assert res.rows == []
    assert any("s.csv" in i and "refused" in i for i in res.issues)
    assert "completed successfully" not in res.message


def test_sign_mismatch_gives_issue_and_no_rows(tmp_path):
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH, sign_override="1.00"))
    res = _run(tmp_path)
    assert res.rows == []
    assert any("BillingAmountSign" in i for i in res.issues)


def test_csv_in_wrong_card_folder_is_refused(tmp_path):
    _write(tmp_path, "ICICI-Amazon", "s.csv", _csv(SAPPH))          # two networks = Sapphiro
    res = _run(tmp_path)
    assert res.rows == []
    assert any("say ICICI-Sapphiro" in i for i in res.issues)


def test_csv_with_no_rows_is_not_clean(tmp_path):
    _write(tmp_path, "ICICI-Amazon", "s.csv", _csv([(SOLO, [])]))
    res = _run(tmp_path)
    assert any("holds no transactions" in i for i in res.issues)
    assert "completed successfully" not in res.message


def test_range_filters_csv_rows(tmp_path):
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH))
    res = _run(tmp_path, start=date(2025, 4, 6), end=date(2025, 4, 30))
    assert len(res.rows) == 2
    assert all(r.date.date() >= date(2025, 4, 6) for r in res.rows)


def test_same_csv_twice_is_counted_once(tmp_path):
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH))
    _write(tmp_path, "ICICI-Sapphiro", "s copy.csv", _csv(SAPPH))
    res = _run(tmp_path)
    assert len(res.rows) == 3
    assert res.duplicates


def test_overlapping_csvs_book_each_row_once(tmp_path):
    a = _csv([(MC, [("05-Apr-25", "SHOP ONE", "100.00"), ("06-Apr-25", "SHOP TWO", "200.00")]), (AMEX, [])])
    b = _csv([(MC, [("06-Apr-25", "SHOP TWO", "200.00"), ("07-Apr-25", "SHOP THREE", "300.00")]), (AMEX, [])])
    _write(tmp_path, "ICICI-Sapphiro", "a.csv", a)
    _write(tmp_path, "ICICI-Sapphiro", "b.csv", b)
    res = _run(tmp_path)
    assert sorted(r.amount for r in res.rows) == [100.0, 200.0, 300.0]


def test_two_genuine_identical_rows_are_both_kept(tmp_path):
    rows = [(MC, [("05-Apr-25", "VENDING A", "50.00"), ("05-Apr-25", "VENDING B", "50.00")]), (AMEX, [])]
    _write(tmp_path, "ICICI-Sapphiro", "a.csv", _csv(rows))
    res = _run(tmp_path)
    assert len(res.rows) == 2


# ---- CSV + PDF overlap, 0-row PDF ----

def _fake_pdf_statement(rows_spec, period, summary=None):
    def fake_parse(lines, bank, card, source):
        rows = [m._mk_row(bank, card, datetime(*d), desc, amt, direction) for d, desc, amt, direction in rows_spec]
        p = m.Parsed(rows=rows, summary=dict(summary or {}))
        st = m.Statement(bank, card, source, period, "", p, list(lines))
        for r in rows:
            r.period, r.source = st.label, source
            r.statement = f"{st.card} | {st.label}"
        return st
    return fake_parse


def _patch_pdf(monkeypatch, parse):
    monkeypatch.setattr(m, "read_lines", lambda pdf: ["synthetic statement text"])
    monkeypatch.setattr(m, "non_statement_reason", lambda lines, bank: None)
    monkeypatch.setattr(m, "parse_statement", parse)


def test_row_in_both_pdf_and_csv_is_booked_once(tmp_path, monkeypatch):
    d = _write(tmp_path, "ICICI-Sapphiro", "apr.pdf", "x").parent
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH))
    spec = [((2025, 4, 5), "SHOP ONE MUMBAI", 1200.50, "Dr"),
            ((2025, 4, 6), "CAFE TWO PUNE", 300.0, "Dr"),
            ((2025, 4, 7), "PAYMENT RECEIVED", 5000.0, "Cr")]
    _patch_pdf(monkeypatch, _fake_pdf_statement(spec, (date(2025, 4, 1), date(2025, 4, 30))))
    res = _run(d.parent)
    assert len(res.rows) == 3                                   # not 6
    assert sum(1 for r in res.rows if r.direction == "Cr") == 1
    assert all(r.source == "apr.pdf" for r in res.rows)         # the tied-out PDF wins
    assert not any("not in it" in i for i in res.issues)


def test_csv_only_row_inside_pdf_period_is_kept_and_flagged(tmp_path, monkeypatch):
    d = _write(tmp_path, "ICICI-Sapphiro", "apr.pdf", "x").parent
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH))
    spec = [((2025, 4, 5), "SHOP ONE MUMBAI", 1200.50, "Dr")]
    _patch_pdf(monkeypatch, _fake_pdf_statement(spec, (date(2025, 4, 1), date(2025, 4, 30)),
                                                summary={"total": 1.0}))
    res = _run(d.parent)
    assert len(res.rows) == 3                                   # 1 PDF + 2 CSV-only, no repeats
    assert any("not in it" in i for i in res.issues)


def test_csv_dates_outside_pdf_period_are_added(tmp_path, monkeypatch):
    d = _write(tmp_path, "ICICI-Sapphiro", "mar.pdf", "x").parent
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH))
    spec = [((2025, 3, 10), "OLD SHOP", 10.0, "Dr")]
    _patch_pdf(monkeypatch, _fake_pdf_statement(spec, (date(2025, 3, 1), date(2025, 3, 31))))
    res = _run(d.parent)
    assert len(res.rows) == 4
    assert not any("not in it" in i for i in res.issues)


def test_zero_row_pdf_without_summary_is_never_clean(tmp_path, monkeypatch):
    d = _write(tmp_path, "ICICI-Sapphiro", "apr.pdf", "x").parent
    _patch_pdf(monkeypatch, _fake_pdf_statement([], (date(2025, 4, 1), date(2025, 4, 30))))
    res = _run(d.parent)
    assert any("apr.pdf" in i and "NOT clean" in i for i in res.issues)
    assert "completed successfully" not in res.message


def test_zero_row_pdf_with_csv_still_flags_the_pdf(tmp_path, monkeypatch):
    d = _write(tmp_path, "ICICI-Sapphiro", "apr.pdf", "x").parent
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH))
    _patch_pdf(monkeypatch, _fake_pdf_statement([], (date(2025, 4, 1), date(2025, 4, 30))))
    res = _run(d.parent)
    assert len(res.rows) == 3
    assert any("apr.pdf" in i and "NOT clean" in i for i in res.issues)
    assert "completed successfully" not in res.message


def test_zero_row_pdf_with_zero_summary_is_still_no_activity(tmp_path, monkeypatch):
    d = _write(tmp_path, "ICICI-Sapphiro", "apr.pdf", "x").parent
    _patch_pdf(monkeypatch, _fake_pdf_statement([], (date(2025, 4, 1), date(2025, 4, 30)),
                                                summary={"prev": 0.0, "total": 0.0}))
    res = _run(d.parent)
    assert any("no activity" in i for i in res.issues)
    assert not any("NOT clean" in i for i in res.issues)


# ---- CC Sort / completeness ----

def test_completeness_counts_csv_as_a_statement(tmp_path):
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH))
    rep = comp.check_folder(tmp_path, (date(2025, 4, 1), date(2025, 4, 30), "FY"),
                            reader=lambda p: None)
    card = [c for c in rep.cards if c.card == "ICICI-Sapphiro"][0]
    assert card.statements == 1
    assert not card.not_statements


def test_completeness_unreadable_csv_is_not_coverage(tmp_path):
    _write(tmp_path, "ICICI-Sapphiro", "s.csv", _csv(SAPPH, header="a,b"))
    rep = comp.check_folder(tmp_path, (date(2025, 4, 1), date(2025, 4, 30), "FY"),
                            reader=lambda p: None)
    card = [c for c in rep.cards if c.card == "ICICI-Sapphiro"][0]
    assert card.statements == 0 and card.not_statements == ["s.csv"]


def test_sort_routes_csv_by_card_numbers(tmp_path):
    sort = _load("extract_sort_cc_pdfs_cc22", SORT_SCRIPT)
    inp = tmp_path / "in"
    inp.mkdir()
    (inp / "random_name.csv").write_text(_csv(SAPPH), encoding="utf-8")
    (inp / "other.csv").write_text("a,b\n1,2\n", encoding="utf-8")        # not a statement: ignored
    assert sort.detect_input_type(inp) == "csv"
    out = tmp_path / "out"
    res = sort.organize_pdfs(str(inp), str(out), password=None)
    assert (out / "Decrypted_PDFs_Correct" / "ICICI-Sapphiro" / "random_name.csv").exists()
    assert res["organization"] == {"ICICI-Sapphiro": 1}
    assert not (out / "Decrypted_PDFs_Correct" / "Unknown-Unknown").exists()


# ---- real-export preamble (single-cell "Transaction Details:" line) ----

def _real_shape(header=HEADER, body_blocks=SAPPH, eol="\n", rows_before_header=False):
    base = _csv(body_blocks, header=header)
    lines = base.split("\n")
    i = lines.index(header)
    pre = lines[:i] + ['"Transaction Details:"']
    if rows_before_header:
        pre.append('01-Apr-25,1,EARLY ROW,5,0.00,"10.00","10.00"')
    return eol.join(pre + lines[i:])


def test_single_cell_transaction_details_preamble_parses(tmp_path):
    p = _write(tmp_path, "a", "a.csv", _real_shape())
    st = ic.parse_icici_csv(p)
    assert len(st.rows) == 3 and len(st.masks) == 2


def test_crlf_no_bom_with_preamble_parses(tmp_path):
    d = Path(tmp_path) / "a"
    d.mkdir()
    p = d / "a.csv"
    p.write_bytes(_real_shape(eol="\r\n").encode("utf-8"))
    assert not p.read_bytes().startswith(b"\xef\xbb\xbf")
    st = ic.parse_icici_csv(p)
    assert len(st.rows) == 3 and len(st.masks) == 2


def test_near_miss_only_is_refused_naming_the_line(tmp_path):
    six = ",".join(HEADER.split(",")[:6])
    text = _real_shape(header=six)
    with pytest.raises(ic.CsvRefused) as e:
        ic.parse_icici_csv(_write(tmp_path, "a", "a.csv", text))
    assert "expected 7 columns" in str(e.value)
    assert f"line {text.split(chr(10)).index('\"Transaction Details:\"') + 1}:" in str(e.value)


def test_preamble_then_renamed_header_is_refused_with_no_rows(tmp_path):
    bad = HEADER.replace("BillingAmountSign", "Sign")
    p = _write(tmp_path, "a", "a.csv", _real_shape(header=bad))
    with pytest.raises(ic.CsvRefused):
        ic.parse_icici_csv(p)
    r = _run(tmp_path)
    assert not any(getattr(x, "desc", "") == "SHOP ONE MUMBAI" for x in r.rows)


def test_exact_header_only_after_dated_rows_is_refused(tmp_path):
    p = _write(tmp_path, "a", "a.csv", _real_shape(rows_before_header=True))
    with pytest.raises(ic.CsvRefused) as e:
        ic.parse_icici_csv(p)
    assert "before any statement header" in str(e.value)


# ---- one-card identity: Amazon Visa 4315 is never Sapphiro ----

def test_one_card_visa_4315_is_amazon_never_sapphiro(tmp_path):
    visa = "4315 XXXX XXXX 0000"
    assert ic.identify_card([visa]) == ("Amazon", "")
    assert ic.identify_card([visa])[0] != "Sapphiro"                  # NEGATIVE
    p = _write(tmp_path, "x", "v.csv", _csv([(visa, [("05-Apr-25", "SHOP", "10.00")])]))
    assert ic.card_folder(p) == ("ICICI", "Amazon")


def test_one_card_amex_3747_is_sapphiro(tmp_path):
    amex = "3747 XXXX XXXX 000"
    assert ic.identify_card([amex])[0] == "Sapphiro"
    p = _write(tmp_path, "x", "a.csv", _csv([(amex, [("05-Apr-25", "SHOP", "10.00")])]))
    assert ic.card_folder(p) == ("ICICI", "Sapphiro")


def test_amex_plus_mastercard_is_still_sapphiro():
    assert ic.identify_card([MC, AMEX])[0] == "Sapphiro"


def test_no_visa_first_group_is_in_the_sapphiro_digit_set():          # NEGATIVE guard
    assert ic.SAPPHIRO_DIGITS == ("3747",)
    assert "4315" not in ic.SAPPHIRO_DIGITS
    assert all(not d.startswith("4") for d in ic.SAPPHIRO_DIGITS)
    # a last group equal to a Sapphiro digit group does not make a card Sapphiro
    assert ic.identify_card(["4000 XXXX XXXX 3747"])[0] == "Amazon"
