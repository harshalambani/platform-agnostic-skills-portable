"""CC-03: Credit Card - Book Spends (cash basis). Synthetic data only; the
statement extraction, the book and the account mapper are replaced by fakes."""
from __future__ import annotations

import csv
import sys
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
import yaml

ROOT = Path(__file__).resolve().parent.parent
for _p in (str(ROOT / "src"), str(ROOT), str(ROOT / "src" / "agents" / "skill_itr_workbook" / "scripts")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from agents.skill_cc_spend_booking import agent as ccsb          # noqa: E402
from agents.skill_cc_spend_booking import journal as J           # noqa: E402
from agents.skill_cc_spend_booking import matcher as M           # noqa: E402
from agents.skill_cc_spend_booking import settings_io            # noqa: E402
from agents.skill_cc_spend_booking.bookio import BookTxn, BookView, CCP_PATH   # noqa: E402
from agents.skill_cc_transactions.agent import _load_script      # noqa: E402

MOD = _load_script()
BANK = "Assets:Bank:Synthetic"
FOOD, SHOP, BSC, DRAW = "Expenses:Food", "Expenses:Shopping", "Expenses:Bank Service Charge", "Equity:Drawings"
EMI_INT = "Expenses:EMI Interest"
PATHS = {CCP_PATH, BANK, FOOD, SHOP, BSC, DRAW, EMI_INT}


# ---------------------------------------------------------------------------
# builders
# ---------------------------------------------------------------------------

def row(desc, amt, d, kind="spend", direction="Dr", card="Regalia", bank="HDFC", when=None):
    return MOD.Row(bank, card, when or datetime(2025, *d), desc, amt, direction, kind)


def stmt(period, rows, card="Regalia", bank="HDFC"):
    return MOD.Statement(bank=bank, card_type=card, source=f"{bank}-{card}/{period[0]:%b%y}.pdf", period=period,
                         period_note="", parsed=MOD.Parsed(rows=rows), lines=[])


def tie(st, result="PASS", printed=0.0, gap=""):
    return MOD.TieOut(card=st.card, period=st.label, source=st.source, previous=0.0, payments=0.0, credits=0.0,
                      spends=0.0, fees=0.0, other=0.0, computed=printed, printed=printed, difference=0.0,
                      result=result, note="", gap=gap)


APR = (date(2025, 4, 1), date(2025, 4, 30))
MAY = (date(2025, 5, 1), date(2025, 5, 31))


def april_rows():
    return [row("SWIGGY ORDER", 1500, (4, 5, 0, 0)[:2] if False else (4, 5)),
            row("AMAZON PURCHASE", 500, (4, 9)),
            row("ANNUAL FEE", 100, (4, 10), kind="fee"),
            row("AMAZON REFUND", 200, (4, 20), kind="refund", direction="Cr")]


def may_rows(pay=1900.0, extra=None):
    return [row("PAYMENT RECEIVED", pay, (5, 12), kind="payment", direction="Cr"),
            row("SWIGGY ORDER", 300, (5, 3))] + (extra or [])


def book_pay(guid, amount, d, desc="HDFC CC PAYMENT"):
    return BookTxn(guid, d, desc, "", ((CCP_PATH, round(amount, 2)), (BANK, -round(amount, 2))))


class Guard:
    blocked: dict = {}

    def blocked_target_reason(self, path):
        return self.blocked.get(path)


def fake_mapper_factory(calls):
    def fake(gnucash, cin, cout, config, drawings, default):
        rows = list(csv.DictReader(open(cin, encoding="utf-8")))
        calls.append([r["Description"] for r in rows])
        for r in rows:
            d = r["Description"].upper()
            r["Account"] = SHOP if any(k in d for k in ("AMAZON", "PHONE", "RING", "TV")) else FOOD
            r["Confidence"], r["MatchReason"] = "high", "fake"
        with open(cout, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
            w.writeheader()
            w.writerows(rows)
    return fake


def go(tmp_path, monkeypatch, stmts, ties, txns, emi=None, entity_extra=None, threshold=None,
       period="Custom date range", start="2025-04-01", end="2025-05-31"):
    book = tmp_path / "book.gnucash"
    book.write_text("x", encoding="utf-8")
    pdfs = tmp_path / "pdfs"
    pdfs.mkdir(exist_ok=True)
    ent = {"name": "Synthetic", "pan": "AAAAA0000A", "status": "Individual", "residency": "Resident",
           "default_regime": "new", "bank_service_charge_account": BSC, "drawings_accounts": [DRAW],
           "card_emi_interest_account": EMI_INT}
    ent.update(entity_extra or {})
    ent = {k: v for k, v in ent.items() if v is not None}
    ents = tmp_path / "entities.yaml"
    ents.write_text(yaml.safe_dump({"SYN": ent}), encoding="utf-8")
    settings = tmp_path / "config.yaml"
    if threshold is not None:
        settings.write_text(yaml.safe_dump({"cc_large_spend_threshold": threshold}), encoding="utf-8")
    result = MOD.Result(message="", issues=[], rows=[], ties=ties, fees=[], duplicates=[], statements=stmts,
                        emi_rows=emi or [], pool=list(stmts))
    monkeypatch.setattr(ccsb, "_extract", lambda *a, **k: (MOD, result))
    monkeypatch.setattr(ccsb, "read_book", lambda p: BookView(txns=list(txns), paths=set(PATHS)))
    from agents.gnucash_accounts import TargetGuard
    monkeypatch.setattr(TargetGuard, "from_book", classmethod(lambda cls, p: Guard()))
    calls = []
    monkeypatch.setattr(ccsb, "_run_mapper", fake_mapper_factory(calls))
    out = tmp_path / "out" / "x_CC-Spend_GnuCash_import_ready.csv"
    reply = ccsb.run(str(pdfs), "SYN", str(book), str(out), period, start, end, str(ents), str(settings))
    rows = list(csv.DictReader(open(out, encoding="utf-8"))) if out.is_file() else []
    return reply, rows, calls, out


def basic(tmp_path, monkeypatch, **kw):
    s1, s2 = stmt(APR, april_rows()), stmt(MAY, may_rows())
    txns = kw.pop("txns", [book_pay("p1", 1900, date(2025, 5, 12))])
    return go(tmp_path, monkeypatch, [s1, s2], [tie(s1, printed=1900.0), tie(s2, printed=300.0)], txns, **kw)


def amount_col(r):
    return r["Amount (Withdrawal)"] or r["Amount Negated (Deposit)"]


# ---------------------------------------------------------------------------
# matcher (pure)
# ---------------------------------------------------------------------------

def L(amount, d, desc="PAYMENT RECEIVED", card="C1"):
    st = SimpleNamespace(card=card, label=f"stmt {card}")
    return M.PayLine(st, SimpleNamespace(description=desc), d, amount)


def P(guid, amount, d, desc="HDFC CC PAYMENT"):
    return M.BookPayment(guid, d, amount, desc, BANK, True)


def test_exact_payment_pairs_one_to_one():
    out = M.match_payments([L(1000, date(2025, 5, 10))], [P("a", 1000, date(2025, 5, 12))])
    assert len(out.pairings) == 1 and out.pairings[0].kind == "exact" and not out.payments_unmatched


@pytest.mark.parametrize("over,ok", [(0.5, True), (15.0, True), (15.5, False)])
def test_cred_overage_is_zero_to_fifteen_rupees(over, ok):
    out = M.match_payments([L(1000 + over, date(2025, 5, 10), desc="CRED PAYMENT")],
                           [P("a", 1000, date(2025, 5, 10), desc="CRED CLUB")])
    assert bool(out.pairings) is ok
    if not ok:                                                                  # NEGATIVE
        assert out.payments_unmatched and not out.pairings


def test_an_overage_without_a_cred_marker_is_never_accepted():                 # NEGATIVE
    out = M.match_payments([L(1010, date(2025, 5, 10))], [P("a", 1000, date(2025, 5, 10))])
    assert not out.pairings and out.payments_unmatched


def test_a_cred_payment_below_the_statement_line_is_not_an_overage():          # NEGATIVE
    out = M.match_payments([L(990, date(2025, 5, 10), desc="CRED PAYMENT")],
                           [P("a", 1000, date(2025, 5, 10), desc="CRED")])
    assert not out.pairings


@pytest.mark.parametrize("days,ok", [(5, True), (6, False)])
def test_five_day_window(days, ok):
    out = M.match_payments([L(1000, date(2025, 5, 10))], [P("a", 1000, date(2025, 5, 10 + days))])
    assert bool(out.pairings) is ok


def test_part_payments_pair_when_the_solution_is_unique():
    out = M.match_payments([L(1000, date(2025, 5, 10))],
                           [P("a", 600, date(2025, 5, 8)), P("b", 400, date(2025, 5, 20))])
    assert len(out.pairings) == 1 and out.pairings[0].kind == "part" and len(out.pairings[0].payments) == 2


def test_part_payments_with_two_solutions_are_ambiguous_never_guessed():       # NEGATIVE
    out = M.match_payments([L(1000, date(2025, 5, 10))],
                           [P("a", 600, date(2025, 5, 8)), P("b", 400, date(2025, 5, 9)),
                            P("c", 600, date(2025, 5, 11)), P("d", 400, date(2025, 5, 12))])
    assert not out.pairings and out.ambiguous
    assert len(out.payments_unmatched) == 4


def test_part_payment_outside_fifteen_days_is_not_used():                      # NEGATIVE
    out = M.match_payments([L(1000, date(2025, 5, 10))],
                           [P("a", 600, date(2025, 5, 8)), P("b", 400, date(2025, 6, 10))])
    assert not out.pairings


def test_two_lines_one_payment_is_ambiguous_and_nothing_pairs():               # NEGATIVE (one-to-one)
    out = M.match_payments([L(1000, date(2025, 5, 10), card="C1"), L(1000, date(2025, 5, 11), card="C2")],
                           [P("a", 1000, date(2025, 5, 10))])
    assert not out.pairings and out.ambiguous and out.payments_unmatched


def test_one_cred_payment_covering_two_cards_stays_unmatched():                # decision 6.6
    out = M.match_payments([L(600, date(2025, 5, 10), desc="CRED PAYMENT", card="C1"),
                            L(400, date(2025, 5, 10), desc="CRED PAYMENT", card="C2")],
                           [P("a", 1000, date(2025, 5, 10), desc="CRED CLUB")])
    assert not out.pairings                                                      # never split between the cards
    assert [p.guid for p in out.payments_unmatched] == ["a"]
    assert len(out.lines_unmatched) == 2


def test_reason_for_a_payment_near_a_line_with_the_wrong_amount():
    out = M.match_payments([L(1000, date(2025, 5, 10))], [P("a", 1500, date(2025, 5, 10))])
    assert M.R_AMOUNT_OFF in out.unmatched_reason["a"]
    out = M.match_payments([L(1000, date(2025, 5, 10))], [P("a", 1000, date(2025, 8, 10))])
    assert out.unmatched_reason["a"] == M.R_NO_STATEMENT


# ---------------------------------------------------------------------------
# journals (pure)
# ---------------------------------------------------------------------------

def test_prior_fy_row_is_booked_on_1_april_with_the_real_date_in_the_description():   # 6.2
    s = stmt((date(2025, 3, 1), date(2025, 3, 31)), [row("SWIGGY ORDER", 100, (3, 28))])
    js, _e, _i = J.journals_for_settlement(s, date(2025, 4, 10), MOD.FEE_REVERSAL_RX, 20000)
    assert js[0].date == date(2025, 4, 1) and js[0].real_date == date(2025, 3, 28)
    assert "28-Mar-2025" in js[0].description


def test_row_inside_the_fy_keeps_its_own_date():                                # NEGATIVE
    s = stmt(APR, [row("SWIGGY ORDER", 100, (4, 5))])
    js, _e, _i = J.journals_for_settlement(s, date(2025, 5, 10), MOD.FEE_REVERSAL_RX, 20000)
    assert js[0].date == date(2025, 4, 5) and "[" not in js[0].description


def test_num_is_deterministic_prefixed_and_distinct_for_identical_rows():
    s = stmt(APR, [row("SWIGGY ORDER", 100, (4, 5)), row("SWIGGY ORDER", 100, (4, 5))])
    a, _e, _i = J.journals_for_settlement(s, date(2025, 5, 10), MOD.FEE_REVERSAL_RX, 20000)
    b, _e, _i = J.journals_for_settlement(s, date(2025, 5, 10), MOD.FEE_REVERSAL_RX, 20000)
    assert [j.num for j in a] == [j.num for j in b]
    assert a[0].num != a[1].num and all(j.num.startswith("CCSB-") for j in a)


def _j(account=FOOD, d=date(2025, 4, 5), amount=100.0, side="W", num="CCSB-x"):
    return J.Journal(num=num, kind=J.K_SPEND, date=d, real_date=d, description="d", amount=amount, side=side,
                     card="C", statement="s", account=account)


def test_rf1_a_bill_payment_never_proves_a_journal_is_booked():                 # NEGATIVE
    j = _j()
    payment = BookTxn("t1", date(2025, 4, 5), "card bill", "", ((CCP_PATH, 100.0), (BANK, -100.0)))
    J.mark_booked([j], [payment], CCP_PATH)
    assert j.status == J.ST_READY


def test_r1_hand_entry_dr_expense_cr_ccp_is_withheld():
    j = _j()
    J.mark_booked([j], [BookTxn("t1", date(2025, 4, 5), "hand", "", ((FOOD, 100.0), (CCP_PATH, -100.0)))], CCP_PATH)
    assert j.status == J.ST_POSSIBLY and "05 Apr 2025" in j.why and FOOD in j.why


def test_r1_hand_entry_two_days_off_is_withheld():
    j = _j()
    J.mark_booked([j], [BookTxn("t1", date(2025, 4, 7), "hand", "", ((FOOD, 100.0), (BANK, -100.0)))], CCP_PATH)
    assert j.status == J.ST_POSSIBLY


def test_r1_hand_entry_on_a_different_expense_account_is_withheld():
    j = _j(account=FOOD)
    J.mark_booked([j], [BookTxn("t1", date(2025, 4, 5), "hand", "", ((SHOP, 100.0), (BANK, -100.0)))], CCP_PATH)
    assert j.status == J.ST_POSSIBLY and SHOP in j.why


def test_r1_six_days_off_is_not_matched():                                     # NEGATIVE
    j = _j()
    J.mark_booked([j], [BookTxn("t1", date(2025, 4, 11), "hand", "", ((FOOD, 100.0), (BANK, -100.0)))], CCP_PATH)
    assert j.status == J.ST_READY
    assert J.ALREADY_BOOKED_WINDOW_DAYS == 5


def test_r1_two_identical_spends_one_hand_entry_withholds_exactly_one():
    a, b = _j(num="CCSB-a"), _j(num="CCSB-b")
    J.mark_booked([a, b], [BookTxn("t1", date(2025, 4, 6), "hand", "", ((FOOD, 100.0), (BANK, -100.0)))], CCP_PATH)
    assert sorted([a.status, b.status]) == sorted([J.ST_READY, J.ST_POSSIBLY])


def test_r1_nearest_date_is_paired_first():
    a = _j(num="CCSB-a", d=date(2025, 4, 5))
    b = _j(num="CCSB-b", d=date(2025, 4, 10))
    J.mark_booked([a, b], [BookTxn("t1", date(2025, 4, 10), "hand", "", ((FOOD, 100.0), (BANK, -100.0)))], CCP_PATH)
    assert b.status == J.ST_POSSIBLY and a.status == J.ST_READY


def test_r1_an_unmapped_spend_still_gets_the_check():
    j = _j(account="")
    J.mark_booked([j], [BookTxn("t1", date(2025, 4, 5), "hand", "", ((FOOD, 100.0), (BANK, -100.0)))], CCP_PATH)
    assert j.status == J.ST_POSSIBLY


def test_r1_credit_side_kinds_match_the_credit_leg():
    j = _j(side="D")
    J.mark_booked([j], [BookTxn("t1", date(2025, 4, 5), "refund", "", ((BANK, 100.0), (FOOD, -100.0)))], CCP_PATH)
    assert j.status == J.ST_POSSIBLY
    k = _j(side="D")
    J.mark_booked([k], [BookTxn("t2", date(2025, 4, 5), "wrong way", "", ((FOOD, 100.0), (BANK, -100.0)))], CCP_PATH)
    assert k.status == J.ST_READY                                              # NEGATIVE


def test_rf1_a_transaction_without_ccp_withholds_one_to_one():
    a, b = _j(num="CCSB-a"), _j(num="CCSB-b")
    t = BookTxn("t1", date(2025, 4, 5), "manual", "", ((FOOD, 100.0), (BANK, -100.0)))
    J.mark_booked([a, b], [t], CCP_PATH)
    assert [a.status, b.status].count(J.ST_POSSIBLY) == 1 and [a.status, b.status].count(J.ST_READY) == 1


def test_num_match_marks_already_booked():
    j = _j(num="CCSB-abc")
    J.mark_booked([j], [BookTxn("t", date(2025, 4, 5), "x", "CCSB-abc", ((FOOD, 100.0), (CCP_PATH, -100.0)))], CCP_PATH)
    assert j.status == J.ST_BOOKED


# ---------------------------------------------------------------------------
# the whole run
# ---------------------------------------------------------------------------

def test_happy_path_books_the_settled_statement_against_ccp(tmp_path, monkeypatch):
    reply, rows, calls, _ = basic(tmp_path, monkeypatch)
    assert "completed successfully" in reply and "NOT complete" not in reply
    assert len(rows) == 4 and all(r["Transfer Account"] == CCP_PATH for r in rows)
    by = {r["Description"]: r for r in rows}
    assert by["SWIGGY ORDER"]["Account"] == FOOD and by["SWIGGY ORDER"]["Amount (Withdrawal)"] == "1500.00"
    assert by["AMAZON REFUND"]["Account"] == SHOP and by["AMAZON REFUND"]["Amount Negated (Deposit)"] == "200.00"
    assert by["ANNUAL FEE"]["Account"] == BSC and by["ANNUAL FEE"]["Amount (Withdrawal)"] == "100.00"
    assert all(r["Transaction ID"].startswith("CCSB-") for r in rows)
    assert "after: 0.00" in reply


def test_the_latest_statement_is_awaiting_payment_and_books_nothing(tmp_path, monkeypatch):    # NEGATIVE
    reply, rows, _c, _ = basic(tmp_path, monkeypatch)
    assert "AWAITING PAYMENT" in reply
    assert not any(r["Amount (Withdrawal)"] == "300.00" for r in rows)


def test_rerun_gives_the_same_ids_and_a_booked_run_is_not_booked_twice(tmp_path, monkeypatch):
    _r, rows1, _c, _ = basic(tmp_path / "a" if (tmp_path / "a").mkdir() is None else tmp_path, monkeypatch)
    ids = [r["Transaction ID"] for r in rows1]
    booked = [BookTxn(f"b{i}", date(2025, 5, 12), "x", n, ((FOOD, 1.0), (CCP_PATH, -1.0))) for i, n in enumerate(ids)]
    d2 = tmp_path / "b"
    d2.mkdir()
    reply, rows2, _c, _ = basic(d2, monkeypatch, txns=[book_pay("p1", 1900, date(2025, 5, 12))] + booked)
    assert rows2 == []                                                           # NEGATIVE: nothing re-booked
    assert "already booked" in reply


def test_journals_are_deterministic_across_runs(tmp_path, monkeypatch):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    _r, r1, _c, _ = basic(tmp_path / "a", monkeypatch)
    _r, r2, _c, _ = basic(tmp_path / "b", monkeypatch)
    assert [x["Transaction ID"] for x in r1] == [x["Transaction ID"] for x in r2]


def test_refusal_of_a_non_pass_statement_shows_its_gap_report(tmp_path, monkeypatch):       # 6.4
    s1, s2 = stmt(APR, april_rows()), stmt(MAY, may_rows())
    reply, rows, _c, _ = go(tmp_path, monkeypatch, [s1, s2],
                            [tie(s1, result="FAIL", printed=1900.0, gap="printed 1,900.00 vs rows 1,700.00"),
                             tie(s2, printed=300.0)], [book_pay("p1", 1900, date(2025, 5, 12))])
    assert rows == [] and "printed 1,900.00 vs rows 1,700.00" in reply
    assert "REFUSED" in reply and M.R_REFUSED in reply
    assert "completed successfully" not in reply                                  # NEGATIVE


def test_pass_chained_counts_as_pass(tmp_path, monkeypatch):
    s1, s2 = stmt(APR, april_rows()), stmt(MAY, may_rows())
    _r, rows, _c, _ = go(tmp_path, monkeypatch, [s1, s2],
                         [tie(s1, result="PASS (chained)", printed=1900.0), tie(s2, printed=300.0)],
                         [book_pay("p1", 1900, date(2025, 5, 12))])
    assert len(rows) == 4


def test_a_revolving_partial_payment_books_nothing_and_is_partly_settled(tmp_path, monkeypatch):   # 6.5
    s1, s2 = stmt(APR, april_rows()), stmt(MAY, may_rows(pay=1000.0))
    reply, rows, _c, _ = go(tmp_path, monkeypatch, [s1, s2], [tie(s1, printed=1900.0), tie(s2, printed=1200.0)],
                            [book_pay("p1", 1000, date(2025, 5, 12))])
    assert rows == [] and "PARTLY SETTLED" in reply and "paid 1,000.00 of 1,900.00" in reply
    assert "completed successfully" not in reply


def test_fee_reversal_is_paired_with_its_fee_and_both_go_to_bank_service_charge(tmp_path, monkeypatch):   # 6.7
    extra = [row("ANNUAL FEE REVERSAL", 100, (4, 25), kind="refund", direction="Cr")]
    s1, s2 = stmt(APR, april_rows() + extra), stmt(MAY, may_rows(pay=1800.0))
    reply, rows, _c, _ = go(tmp_path, monkeypatch, [s1, s2], [tie(s1, printed=1800.0), tie(s2, printed=300.0)],
                            [book_pay("p1", 1800, date(2025, 5, 12))])
    rev = next(r for r in rows if "REVERSAL" in r["Description"])
    assert rev["Account"] == BSC and rev["Amount Negated (Deposit)"] == "100.00"
    assert "reverses" in reply and "net 0.00" in reply
    assert not any(r["Account"] == SHOP and "REVERSAL" in r["Description"] for r in rows)   # NEGATIVE


def test_cashback_goes_to_drawings(tmp_path, monkeypatch):
    extra = [row("CASHBACK CREDIT", 50, (4, 28), kind="cashback", direction="Cr")]
    s1, s2 = stmt(APR, april_rows() + extra), stmt(MAY, may_rows(pay=1850.0))
    _r, rows, _c, _ = go(tmp_path, monkeypatch, [s1, s2], [tie(s1, printed=1850.0), tie(s2, printed=300.0)],
                         [book_pay("p1", 1850, date(2025, 5, 12))])
    cb = next(r for r in rows if "CASHBACK" in r["Description"])
    assert cb["Account"] == DRAW and cb["Amount Negated (Deposit)"] == "50.00"


def test_cred_overage_is_dr_ccp_cr_drawings(tmp_path, monkeypatch):
    r1 = april_rows()
    r1[0] = row("SWIGGY ORDER", 1510, (4, 5))
    s1 = stmt(APR, r1)
    s2 = stmt(MAY, [row("CRED PAYMENT RECEIVED", 1910, (5, 12), kind="payment", direction="Cr")])
    reply, rows, _c, _ = go(tmp_path, monkeypatch, [s1, s2], [tie(s1, printed=1910.0), tie(s2, printed=0.0)],
                            [book_pay("p1", 1900, date(2025, 5, 12), desc="CRED CLUB")])
    ov = next(r for r in rows if "CRED payment" in r["Description"])
    assert ov["Account"] == DRAW and ov["Amount Negated (Deposit)"] == "10.00"
    assert "after: 0.00" in reply or "after: 0.0" in reply


def test_missing_bank_service_charge_key_is_a_named_error_not_a_guess(tmp_path, monkeypatch):
    reply, rows, _c, out = basic(tmp_path, monkeypatch, entity_extra={"bank_service_charge_account": None})
    assert "bank_service_charge_account" in reply and reply.startswith("ERROR")
    assert rows == [] and not out.exists()                                       # NEGATIVE
    assert reply.withhold_primary


def test_blocked_bank_service_charge_account_stops_the_run(tmp_path, monkeypatch):
    Guard.blocked = {BSC: "placeholder"}
    try:
        reply, rows, _c, _ = basic(tmp_path, monkeypatch)
    finally:
        Guard.blocked = {}
    assert reply.startswith("ERROR") and "placeholder" in reply and rows == []


def test_a_blocked_mapped_account_is_withheld_never_written(tmp_path, monkeypatch):
    Guard.blocked = {SHOP: "hidden"}
    try:
        reply, rows, _c, _ = basic(tmp_path, monkeypatch)
    finally:
        Guard.blocked = {}
    assert all(r["Account"] != SHOP for r in rows) and "WITHHELD" in reply
    assert "completed successfully" not in reply


# ---- large-spend flag -------------------------------------------------------

def _big(amount, threshold=None, desc="PHONE STORE"):
    return amount, threshold, desc


def run_big(tmp_path, monkeypatch, amount, threshold=None):
    rows_ = [row("PHONE STORE", amount, (4, 5))]
    s1, s2 = stmt(APR, rows_), stmt(MAY, [row("PAYMENT RECEIVED", amount, (5, 12), kind="payment", direction="Cr")])
    return go(tmp_path, monkeypatch, [s1, s2], [tie(s1, printed=amount), tie(s2, printed=0.0)],
              [book_pay("p1", amount, date(2025, 5, 12))], threshold=threshold)


def test_threshold_is_flagged_and_one_below_is_not(tmp_path, monkeypatch):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    at, _r, _c, _ = run_big(tmp_path / "a", monkeypatch, 20000.0)[:1] + (None, None, None)
    below = run_big(tmp_path / "b", monkeypatch, 19999.0)[0]
    assert "LARGE SPENDS (1)" in at and J.LARGE_SPEND_FLAG in at
    assert "LARGE SPENDS" not in below                                           # NEGATIVE


def test_a_changed_setting_is_honoured(tmp_path, monkeypatch):
    reply = run_big(tmp_path, monkeypatch, 6000.0, threshold=5000)[0]
    assert "LARGE SPENDS (1)" in reply and "5,000.00" in reply
    (tmp_path / "n").mkdir()
    again = run_big(tmp_path / "n", monkeypatch, 6000.0, threshold=7000)[0]
    assert "LARGE SPENDS" not in again                                            # NEGATIVE


def test_a_flag_never_changes_the_account(tmp_path, monkeypatch):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    _r1, flagged, _c, _ = run_big(tmp_path / "a", monkeypatch, 25000.0)
    _r2, plain, _c, _ = run_big(tmp_path / "b", monkeypatch, 25000.0, threshold=50000)
    assert flagged[0]["Account"] == plain[0]["Account"] == SHOP
    assert not flagged[0]["Account"].startswith("Assets")                         # NEGATIVE: never an asset account


def test_large_spend_flag_is_in_the_workbook(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    reply = run_big(tmp_path, monkeypatch, 25000.0)[0]
    wb = load_workbook(reply.extra_outputs[0]["path"])
    assert "Large spends" in wb.sheetnames
    flags = [c.value for row_ in wb["Journals"].iter_rows(min_row=2) for c in row_ if c.value == J.LARGE_SPEND_FLAG]
    assert flags


# ---- EMI --------------------------------------------------------------------

def emi_stmts(extra=None):
    rows_ = [row("PHONE STORE", 5000, (4, 3)),
             row("SWIGGY ORDER", 400, (4, 4)),
             row("EMI CONVERSION PHONE STORE", 5000, (4, 6), kind="emi_conversion", direction="Cr"),
             row("EMI PRINCIPAL INSTALMENT 1", 800, (4, 7), kind="emi_principal"),
             row("EMI SOMETHING ELSE", 55, (4, 8), kind="emi_unclassified"),
             row("EMI INTEREST", 120, (4, 30), kind="emi_interest"),
             row("EMI PROCESSING FEE", 99, (4, 30), kind="emi_processing_fee"),
             row("GST ON EMI PROCESSING FEE", 18, (4, 30), kind="gst_on_emi"),
             row("ANNUAL FEE", 100, (4, 10), kind="fee")] + (extra or [])
    s1 = stmt(APR, rows_)
    total = 5000 + 400 - 5000 + 800 + 55 + 120 + 99 + 18 + 100
    s2 = stmt(MAY, [row("PAYMENT RECEIVED", total, (5, 12), kind="payment", direction="Cr")])
    emi = [r for r in rows_ if r.kind.startswith("emi") or r.kind == "gst_on_emi"]
    return s1, s2, emi, float(total)


def run_emi(tmp_path, monkeypatch, entity_extra=None, txns_extra=None):
    s1, s2, emi, total = emi_stmts()
    return go(tmp_path, monkeypatch, [s1, s2], [tie(s1, printed=total), tie(s2, printed=0.0)],
              [book_pay("p1", total, date(2025, 5, 12))] + (txns_extra or []), emi=emi, entity_extra=entity_extra)


def test_conversion_principal_and_unclassified_are_never_booked(tmp_path, monkeypatch):      # NEGATIVE
    reply, rows, calls, _ = run_emi(tmp_path, monkeypatch)
    descs = [r["Description"] for r in rows]
    for bad in ("CONVERSION", "PRINCIPAL", "SOMETHING ELSE"):
        assert not any(bad in d for d in descs)
        assert not any(bad in d for batch in calls for d in batch)               # never even sent to the mapper
    assert "PARTLY BOOKED" in reply and J.EMI_UNBOOKED_TITLE.upper() in reply
    assert "EMI CONVERSION PHONE STORE" in reply and "EMI PRINCIPAL INSTALMENT 1" in reply
    assert "completed successfully" not in reply


def test_unbooked_emi_rows_stay_unbooked_on_a_rerun(tmp_path, monkeypatch):                  # NEGATIVE
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    _r, rows1, _c, _ = run_emi(tmp_path / "a", monkeypatch)
    ids = [r["Transaction ID"] for r in rows1]
    booked = [BookTxn(f"b{i}", date(2025, 5, 12), "x", n, ((FOOD, 1.0), (CCP_PATH, -1.0))) for i, n in enumerate(ids)]
    reply, rows2, _c, _ = run_emi(tmp_path / "b", monkeypatch, txns_extra=booked)
    assert rows2 == []
    assert "EMI PRINCIPAL INSTALMENT 1" in reply and "PARTLY BOOKED" in reply


def test_emi_interest_is_booked_to_the_entity_interest_account_once_and_marked_unvalidated(tmp_path, monkeypatch):
    reply, rows, calls, _ = run_emi(tmp_path, monkeypatch)
    ints = [r for r in rows if r["Description"] == "EMI INTEREST"]
    assert len(ints) == 1 and ints[0]["Account"] == EMI_INT and ints[0]["Amount (Withdrawal)"] == "120.00"
    assert ints[0]["Transfer Account"] == CCP_PATH
    assert "UNVALIDATED" in reply
    assert not any("EMI INTEREST" in d for batch in calls for d in batch)       # NEGATIVE: never via the mapper
    assert ints[0]["Account"] != BSC                                              # NEGATIVE: never Bank Service Charge
    assert sum(1 for r in rows if "INTEREST" in r["Description"]) == 1           # NEGATIVE: not twice


def test_emi_interest_already_in_the_book_is_not_booked_twice(tmp_path, monkeypatch):        # NEGATIVE
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    _r, rows1, _c, _ = run_emi(tmp_path / "a", monkeypatch)
    n = next(r["Transaction ID"] for r in rows1 if r["Description"] == "EMI INTEREST")
    t = BookTxn("bx", date(2025, 4, 30), "x", n, ((EMI_INT, 120.0), (CCP_PATH, -120.0)))
    _r, rows2, _c, _ = run_emi(tmp_path / "b", monkeypatch, txns_extra=[t])
    assert not any(r["Description"] == "EMI INTEREST" for r in rows2)


def test_processing_fee_and_gst_go_to_bsc_and_always_in_the_fee_block(tmp_path, monkeypatch):
    reply, rows, _c, _ = run_emi(tmp_path, monkeypatch)
    by = {r["Description"]: r for r in rows}
    for d, amt in (("EMI PROCESSING FEE", "99.00"), ("GST ON EMI PROCESSING FEE", "18.00"), ("ANNUAL FEE", "100.00")):
        assert by[d]["Account"] == BSC and by[d]["Amount (Withdrawal)"] == amt
    fee_block = reply[reply.index("FEES BOOKED"):].split("\n\n")[0]
    assert "FEES BOOKED (3): charged 217.00" in fee_block
    assert "EMI PROCESSING FEE" in fee_block and "GST ON EMI PROCESSING FEE" in fee_block


def test_emi_interest_never_enters_the_fee_total_or_the_bsc_total(tmp_path, monkeypatch):     # NEGATIVE
    reply, rows, _c, _ = run_emi(tmp_path, monkeypatch)
    assert "EMI interest (its own subtotal" in reply and "FY 2025-26 | 120.00" in reply
    assert "charged 217.00" in reply and "charged 337.00" not in reply
    bsc = sum(float(r["Amount (Withdrawal)"] or 0) for r in rows if r["Account"] == BSC)
    assert bsc == 217.0


def test_missing_emi_interest_account_books_no_interest_and_reports_it(tmp_path, monkeypatch):   # NEGATIVE
    reply, rows, calls, _ = run_emi(tmp_path, monkeypatch, entity_extra={"card_emi_interest_account": None})
    assert not any(r["Description"] == "EMI INTEREST" for r in rows)
    assert not any("EMI INTEREST" in d for batch in calls for d in batch)
    assert all(r["Account"] != EMI_INT for r in rows)
    assert "EMI INTEREST NOT BOOKED" in reply and "card_emi_interest_account" in reply
    assert "EMI INTEREST" in reply.split(J.EMI_UNBOOKED_TITLE.upper())[1]        # sits in the awaiting block
    assert "completed successfully" not in reply


def test_hidden_or_placeholder_emi_interest_account_books_no_interest(tmp_path, monkeypatch):    # NEGATIVE
    Guard.blocked = {EMI_INT: "placeholder"}
    try:
        reply, rows, calls, _ = run_emi(tmp_path, monkeypatch)
    finally:
        Guard.blocked = {}
    assert not any(r["Description"] == "EMI INTEREST" for r in rows)
    assert all(r["Account"] != EMI_INT for r in rows)
    assert "EMI INTEREST NOT BOOKED" in reply and "placeholder" in reply


def test_original_purchase_and_its_conversion_reversal_are_never_both_booked(tmp_path, monkeypatch):
    reply, rows, _c, _ = run_emi(tmp_path, monkeypatch)
    assert not any(r["Description"] == "PHONE STORE" for r in rows)               # NEGATIVE
    assert any(r["Description"] == "SWIGGY ORDER" for r in rows)                  # other rows still book
    assert "converted to EMI" in reply


def test_emi_summary_keeps_interest_out_of_fee_totals():
    s1, _s2, emi, _t = emi_stmts()
    js, _e, _i = J.journals_for_settlement(s1, date(2025, 5, 12), MOD.FEE_REVERSAL_RX, 20000)
    assert J.fee_totals(js)["fees"] == 217.0
    assert J.emi_summary(emi)["interest"] == {("HDFC-Regalia", "2025-26"): 120.0}


def test_journals_for_settlement_never_journals_principal_conversion_or_unclassified():       # NEGATIVE
    s1, _s2, _emi, _t = emi_stmts()
    for book_interest in (True, False):
        js, left, _i = J.journals_for_settlement(s1, date(2025, 5, 12), MOD.FEE_REVERSAL_RX, 20000, book_interest)
        assert {r.kind for r in left} >= {"emi_conversion", "emi_principal", "emi_unclassified"}
        assert not any(j.kind == J.K_EMI_INT for j in js) == (not book_interest)


# ---- payments without a statement (RED FLAG) --------------------------------

def test_payment_with_no_statement_is_a_red_flag_at_the_top_and_never_dropped(tmp_path, monkeypatch):
    txns = [book_pay("p1", 1900, date(2025, 5, 12)), book_pay("p2", 777, date(2025, 5, 20), desc="CRED CLUB")]
    reply, rows, _c, _ = basic(tmp_path, monkeypatch, txns=txns)
    assert reply.startswith("RED FLAG")
    assert "777.00" in reply and "CRED" in reply and M.R_NO_STATEMENT in reply
    assert BANK in reply and "20 May 2025" in reply
    assert reply.index("RED FLAG") < reply.index("FEES BOOKED")
    assert "completed successfully" not in reply                                   # NEGATIVE: never success
    assert "of which payments with no statement: 777.00" in reply
    assert "after: 777.00" in reply


def test_a_payment_off_tolerance_is_not_matched_to_the_nearest_statement(tmp_path, monkeypatch):   # NEGATIVE
    reply, rows, _c, _ = basic(tmp_path, monkeypatch, txns=[book_pay("p1", 1950, date(2025, 5, 12))])
    assert rows == [] and reply.startswith("RED FLAG") and M.R_AMOUNT_OFF in reply
    assert "1,900.00" in reply                                                      # names the nearest line, still flags it


def test_a_payment_settling_a_statement_from_before_the_range_is_listed(tmp_path, monkeypatch):
    s1, s2 = stmt(APR, april_rows()), stmt(MAY, may_rows())
    reply, rows, _c, _ = go(tmp_path, monkeypatch, [s2], [tie(s2, printed=300.0)],
                            [book_pay("p1", 1900, date(2025, 5, 12))], start="2025-05-01", end="2025-05-31")
    # S1 is not among the statements read in this run, so find_prior finds nothing in the pool
    assert reply.startswith("RED FLAG") and rows == []


def test_payment_on_a_refused_statement_is_flagged_with_that_reason(tmp_path, monkeypatch):
    s1, s2 = stmt(APR, april_rows()), stmt(MAY, may_rows())
    reply, _r, _c, _ = go(tmp_path, monkeypatch, [s1, s2], [tie(s1, result="FAIL", printed=1900.0), tie(s2, printed=300.0)],
                          [book_pay("p1", 1900, date(2025, 5, 12))])
    assert reply.startswith("RED FLAG") and M.R_REFUSED in reply.split("\n\n")[0]


def test_payments_outside_the_range_are_not_flagged(tmp_path, monkeypatch):    # NEGATIVE
    txns = [book_pay("p1", 1900, date(2025, 5, 12)), book_pay("p9", 555, date(2025, 8, 1))]
    reply, _r, _c, _ = basic(tmp_path, monkeypatch, txns=txns)
    assert "555.00" not in reply and "completed successfully" in reply


def test_ccsb_journals_already_in_the_book_are_not_payments(tmp_path, monkeypatch):    # NEGATIVE
    spend = BookTxn("z", date(2025, 5, 12), "x", "CCSB-zzz", ((FOOD, 50.0), (CCP_PATH, -50.0)))
    refund = BookTxn("y", date(2025, 5, 12), "x", "CCSB-yyy", ((CCP_PATH, 50.0), (FOOD, -50.0)))
    reply, _r, _c, _ = basic(tmp_path, monkeypatch,
                             txns=[book_pay("p1", 1900, date(2025, 5, 12)), spend, refund])
    assert not reply.startswith("RED FLAG")


# ---- book / inputs ----------------------------------------------------------

def test_unreadable_book_is_a_named_error(tmp_path, monkeypatch):
    s1 = stmt(APR, april_rows())
    book = tmp_path / "b.gnucash"
    book.write_text("not a book", encoding="utf-8")
    ents = tmp_path / "e.yaml"
    ents.write_text(yaml.safe_dump({"SYN": {"name": "S", "pan": "AAAAA0000A", "status": "Individual"}}), encoding="utf-8")
    (tmp_path / "pdfs").mkdir()
    reply = ccsb.run(str(tmp_path / "pdfs"), "SYN", str(book), str(tmp_path / "o.csv"), "Custom date range",
                     "2025-04-01", "2025-05-31", str(ents), "")
    assert reply.startswith("ERROR") and "could not read" in reply and reply.withhold_primary


def test_unknown_entity_is_a_named_error(tmp_path, monkeypatch):
    reply = ccsb.run(str(tmp_path), "NOPE", str(tmp_path / "x"), str(tmp_path / "o.csv"), "Custom date range",
                     "2025-04-01", "2025-05-31", "", "")
    assert reply.startswith("ERROR")


def test_nothing_is_written_when_nothing_is_ready(tmp_path, monkeypatch):
    reply, rows, _c, out = basic(tmp_path, monkeypatch, txns=[book_pay("p1", 1950, date(2025, 5, 12))])
    assert not out.exists() and reply.withhold_primary
    assert reply.extra_outputs[0]["path"]                                           # the workbook still ships


# ---------------------------------------------------------------------------
# settings, entity keys, manifest
# ---------------------------------------------------------------------------

def test_threshold_default_and_reads(tmp_path):
    assert settings_io.read_large_spend_threshold(None) == 20000.0
    assert settings_io.read_large_spend_threshold(tmp_path / "missing.yaml") == 20000.0
    p = tmp_path / "c.yaml"
    p.write_text("cc_large_spend_threshold: 7500\n", encoding="utf-8")
    assert settings_io.read_large_spend_threshold(p) == 7500.0


@pytest.mark.parametrize("bad", ["abc", "-5", "0", "nan", None, "inf"])
def test_invalid_threshold_values_fall_back_to_the_default(tmp_path, bad):      # NEGATIVE
    p = tmp_path / "c.yaml"
    p.write_text(yaml.safe_dump({"cc_large_spend_threshold": bad}), encoding="utf-8")
    assert settings_io.read_large_spend_threshold(p) == 20000.0


def test_settings_save_refuses_an_invalid_value_and_writes_nothing(tmp_path, monkeypatch):   # NEGATIVE
    from ui import _config
    from ui.tabs import settings as st
    cfg = tmp_path / "config.yaml"
    cfg.write_text("active_endpoint: x\n", encoding="utf-8")
    monkeypatch.setattr(_config, "PORTABLE_CONFIG_PATH", cfg)
    for bad in ("abc", -1, 0, None):
        assert "Not saved" in st._save_large_spend_threshold(bad)
    assert cfg.read_text(encoding="utf-8") == "active_endpoint: x\n"
    assert "Saved" in st._save_large_spend_threshold("25,000")
    got = yaml.safe_load(cfg.read_text(encoding="utf-8"))
    assert got["cc_large_spend_threshold"] == 25000.0 and got["active_endpoint"] == "x"
    assert st._load_large_spend_threshold() == 25000.0


def test_entity_keys_round_trip_and_stay_empty_when_absent(tmp_path):
    import configs
    p = tmp_path / "e.yaml"
    base = {"name": "S", "pan": "AAAAA0000A", "status": "Individual"}
    p.write_text(yaml.safe_dump({"A": dict(base, bank_service_charge_account="Expenses:BSC",
                                           card_emi_interest_account="Expenses:EMI Interest"),
                                 "B": dict(base)}), encoding="utf-8")
    ents = configs.load_entities(p)
    assert ents["A"].bank_service_charge_account == "Expenses:BSC"
    assert ents["A"].card_emi_interest_account == "Expenses:EMI Interest"
    assert ents["B"].bank_service_charge_account == "" == ents["B"].card_emi_interest_account
    out = yaml.safe_load(configs.dump_entities(ents))
    assert out["A"]["bank_service_charge_account"] == "Expenses:BSC"
    assert "bank_service_charge_account" not in out["B"] and "card_emi_interest_account" not in out["B"]


def test_entities_form_does_not_drop_the_new_keys():
    from ui.tabs import itr_entities
    assert {"bank_service_charge_account", "card_emi_interest_account"} <= set(itr_entities._FORM_HIDDEN_FIELDS)


def test_manifest_shape():
    sk = yaml.safe_load((ROOT / "src/agents/skill_cc_spend_booking/skill.yaml").read_text(encoding="utf-8"))
    assert sk["category"] == "gnucash" and sk["entry_point"] == "agent:run"
    names = {i["name"]: i for i in sk["inputs"]}
    assert names["gnucash_path"]["book_from"] == "entity" and names["gnucash_path"]["fy_from"] == "fy"
    assert names["fy"]["options_from"] == "itr_ay_years" and names["period"]["options_from"] == "report_periods"
    assert sk["run_args"]["settings_path"] == "{data_root}/settings/config.yaml"
    assert sk["output"]["extension"] == ".csv" and "GnuCash_import_ready" in sk["output"]["suffix"]
    assert [e["key"] for e in sk["output"]["extra_outputs"]] == ["workbook"]
    import inspect
    params = set(inspect.signature(ccsb.run).parameters)
    for k in sk["run_args"]:
        assert k in params, k
