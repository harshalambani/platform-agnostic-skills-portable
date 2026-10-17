"""
tests/test_cc20_fee_block.py -- CC-20: reversed fees stay visible in the fee block
but are not totalled, and fee lines print once (the issues list carries a pointer
for unreversed fees only).

Everything is synthetic.
"""
from __future__ import annotations

import importlib.util
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
SCRIPT = SRC / "agents" / "skill_cc_transactions" / "scripts" / "create_cc_transaction_list.py"
for _p in (str(SRC), str(ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)


def _load_module():
    spec = importlib.util.spec_from_file_location("create_cc_transaction_list", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


m = _load_module()


def _row(desc, amt, direction, kind, day, card="Sample"):
    return m.Row("Axis", card, datetime(2025, 4, day), desc, amt, direction, kind)


def _fees(rows):
    st = m.Statement("Axis", "Sample", "a.pdf", None, "", m.Parsed(rows=rows))
    return m.collect_fees([st])


def _total_line(block, card="Axis-Sample"):
    return next(ln for ln in block.splitlines() if ln.strip().startswith(f"Total fees {card}"))


def test_an_unreversed_fee_is_counted_in_full():
    block = m.fee_block(_fees([_row("ANNUAL FEE", 500.0, "Dr", "fee", 25),
                               _row("IGST @18%", 90.0, "Dr", "fee", 25)]))
    assert "Total fees Axis-Sample: 590.00" in block
    assert "reversed" not in block


def test_a_reversed_fee_stays_visible_but_is_not_totalled():
    block = m.fee_block(_fees([_row("ANNUAL FEE", 500.0, "Dr", "fee", 25),
                               _row("ANNUAL FEE REVERSAL", 500.0, "Cr", "refund", 30),
                               _row("LATE PAYMENT FEE", 100.0, "Dr", "fee", 26)]))
    assert "ANNUAL FEE" in block and "[reversed:" in block          # the original line is still there
    assert "Total fees Axis-Sample: 100.00 (excludes 500.00 reversed)" in block


def test_a_fully_reversed_card_totals_zero_never_negative():
    block = m.fee_block(_fees([_row("ANNUAL FEE", 500.0, "Dr", "fee", 25),
                               _row("ANNUAL FEE WAIVER", 500.0, "Cr", "refund", 30)]))
    line = _total_line(block)
    assert "Total fees Axis-Sample: 0.00" in line
    assert "-" not in line.split(":", 1)[1].split("(")[0]
    assert "ANNUAL FEE" in block                                     # never hidden


def test_a_credit_that_is_not_a_fee_reversal_does_not_cancel_a_fee():
    fees = _fees([_row("ANNUAL FEE", 500.0, "Dr", "fee", 25),
                  _row("EXAMPLE MART REFUND", 500.0, "Cr", "refund", 30)])
    block = m.fee_block(fees)
    assert fees[0].reversed_by is None
    assert "Total fees Axis-Sample: 500.00" in block
    assert "[reversed" not in block


def test_a_reversal_of_a_different_amount_does_not_cancel_the_fee():
    fees = _fees([_row("ANNUAL FEE", 500.0, "Dr", "fee", 25),
                  _row("ANNUAL FEE REVERSAL", 400.0, "Cr", "refund", 30)])
    assert fees[0].reversed_by is None
    assert "Total fees Axis-Sample: 500.00" in m.fee_block(fees)


def test_a_reversal_on_another_card_does_not_cancel_the_fee():
    st1 = m.Statement("Axis", "Sample", "a.pdf", None, "", m.Parsed(rows=[_row("ANNUAL FEE", 500.0, "Dr", "fee", 25)]))
    st2 = m.Statement("Axis", "Other", "b.pdf", None, "",
                      m.Parsed(rows=[_row("ANNUAL FEE REVERSAL", 500.0, "Cr", "refund", 30, card="Other")]))
    fees = m.collect_fees([st1, st2])
    assert fees[0].reversed_by is None


# --- once only ---------------------------------------------------------------

def _run(tmp_path, monkeypatch, lines):
    d = tmp_path / "in" / "Axis-Sample"
    d.mkdir(parents=True, exist_ok=True)
    (d / "a.pdf").write_bytes(b"%PDF-1.4 synthetic")
    monkeypatch.setattr(m, "read_lines", lambda p: lines)
    return m.run_extraction(tmp_path / "in", tmp_path / "out.xlsx")


STMT = [
    "01/04/2025 - 30/04/2025 15/05/2025 25/05/2025",
    "0.00 Dr 0.00 0.00 1,000.00 0.00 590.00 1,590.00 Dr",
    "05/04/2025 EXAMPLE MART MUMBAI DEPT STORES 1,000.00 Dr",
    "25/04/2025 ANNUAL FEE FEES 500.00 Dr",
    "25/04/2025 IGST @18% FEES 90.00 Dr",
]


def test_fee_lines_print_once_and_the_issue_is_a_pointer(tmp_path, monkeypatch):
    res = _run(tmp_path, monkeypatch, STMT)
    assert res.message.count("ANNUAL FEE") == 1
    assert res.message.count("IGST @18%") == 1
    fee_issues = [i for i in res.issues if i.startswith("fee charged")]
    assert len(fee_issues) == 1
    assert "ANNUAL FEE" not in fee_issues[0]
    assert "2 card fee(s) not reversed, total 590.00" in fee_issues[0]


def test_a_reversed_fee_is_not_listed_as_an_issue(tmp_path, monkeypatch):
    lines = STMT[:1] + ["0.00 Dr 0.00 0.00 1,000.00 0.00 90.00 1,090.00 Dr",
                        "05/04/2025 EXAMPLE MART MUMBAI DEPT STORES 1,000.00 Dr",
                        "25/04/2025 ANNUAL FEE FEES 500.00 Dr",
                        "26/04/2025 ANNUAL FEE REVERSAL FEES 500.00 Cr",
                        "25/04/2025 IGST @18% FEES 90.00 Dr"]
    res = _run(tmp_path, monkeypatch, lines)
    fee_issues = [i for i in res.issues if i.startswith("fee charged")]
    assert len(fee_issues) == 1 and "1 card fee(s) not reversed, total 90.00" in fee_issues[0]
    assert "[reversed:" in res.message                                  # still visible in the block
    assert "Total fees Axis-Sample: 90.00 (excludes 500.00 reversed)" in res.message


def test_the_fees_sheet_still_lists_every_fee_reversed_or_not(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    lines = STMT[:1] + ["0.00 Dr 0.00 0.00 1,000.00 0.00 90.00 1,090.00 Dr",
                        "05/04/2025 EXAMPLE MART MUMBAI DEPT STORES 1,000.00 Dr",
                        "25/04/2025 ANNUAL FEE FEES 500.00 Dr",
                        "26/04/2025 ANNUAL FEE REVERSAL FEES 500.00 Cr",
                        "25/04/2025 IGST @18% FEES 90.00 Dr"]
    res = _run(tmp_path, monkeypatch, lines)
    wb = load_workbook(res.output)
    rows = list(wb["Fees"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    assert {r[5] for r in rows} == {"Yes", "No"}
