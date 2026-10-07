"""
tests/test_ui18_entity_fy_matching.py -- UI-18: entity AND financial-year
matching on the partner and AIS forms. The pickers and the run path agree:
what is not offered is also refused when passed directly. Everything here is
synthetic (made-up names and PANs).
"""
import sys
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
import yaml

ROOT = Path(__file__).resolve().parent.parent
for _p in (ROOT, ROOT / "src"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

from agents import entity_scope  # noqa: E402
from agents.registry import discover  # noqa: E402
from agents.skill_ais_reconcile import agent as ais_agent  # noqa: E402
from agents.skill_partner_comp_recon import agent as partner_agent  # noqa: E402
from ui.tabs import _generic  # noqa: E402

IND = {"name": "Test Individual", "pan": "AAAAA0000A"}
HUF = {"name": "Test Family HUF", "pan": "BBBBB1111B"}
FY = "2025-26"


def make_26as(path, name, pan, fy=FY, with_identity=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Part I"
    ws.cell(row=1, column=1, value="PART I")
    if with_identity:
        ws.cell(row=2, column=1,
                value=f"Assessee Name: {name}  |  PAN: {pan}  |  Financial Year: {fy}")
    wb.save(path)
    return path


def _skill(folder):
    return next(s for s in discover(refresh=True) if Path(s.manifest_path).parent.name == folder)


@pytest.fixture
def world(tmp_path, monkeypatch):
    data = tmp_path / "Data"
    (data / "itr").mkdir(parents=True)
    (data / "itr" / "entities.yaml").write_text(yaml.safe_dump({
        "TEST-IND": dict(IND, status="Individual"),
        "TEST-HUF": dict(HUF, status="HUF"),
    }), encoding="utf-8")
    out = tmp_path / "out"
    out.mkdir()
    monkeypatch.setattr(_generic._config, "data_root_dir", lambda: data)
    monkeypatch.setattr(_generic._config, "output_dir", lambda: out)
    return out


# ---- (a) the partner form has an FY, and the picker follows it -------------

def test_partner_form_has_a_required_fy_right_after_the_entity():
    sk = _skill("skill_partner_comp_recon")
    names = [i.name for i in sk.inputs]
    assert names[:2] == ["entity", "fy"]
    fy = sk.inputs[1]
    assert fy.type == "select" and fy.options_from == "itr_ay_years" and fy.required
    assert sk.run_args["financial_year"] == "{inputs.fy}"
    assert next(i for i in sk.inputs if i.name == "xlsx_26as").fy_from == "fy"


def test_picker_offers_only_this_entitys_file_for_this_fy(world):
    mine = make_26as(world / "a-26AS.xlsx", IND["name"], IND["pan"], fy=FY)
    other_fy = make_26as(world / "b-26AS.xlsx", IND["name"], IND["pan"], fy="2024-25")
    other_ent = make_26as(world / "c-26AS.xlsx", HUF["name"], HUF["pan"], fy=FY)
    inp = next(i for i in _skill("skill_partner_comp_recon").inputs if i.name == "xlsx_26as")
    choices, value, _m = _generic._scoped_picker_state(inp, "TEST-IND", FY)
    offered = [v for _, v in choices]
    assert str(mine) in offered and value == str(mine)
    assert str(other_fy) not in offered and str(other_ent) not in offered     # NEGATIVE


def test_picker_with_no_file_for_the_fy_says_which_year(world):
    make_26as(world / "b-26AS.xlsx", IND["name"], IND["pan"], fy="2024-25")
    inp = next(i for i in _skill("skill_partner_comp_recon").inputs if i.name == "xlsx_26as")
    choices, value, msg = _generic._scoped_picker_state(inp, "TEST-IND", FY)
    assert [v for _, v in choices] == [""] and value == "" and FY in msg


def test_picker_without_a_fy_is_not_narrowed_by_year(world):
    a = make_26as(world / "a-26AS.xlsx", IND["name"], IND["pan"], fy="2024-25")
    inp = next(i for i in _skill("skill_partner_comp_recon").inputs if i.name == "xlsx_26as")
    choices, _v, _m = _generic._scoped_picker_state(inp, "TEST-IND")
    assert str(a) in [v for _, v in choices]


def test_check_26as_fy(world):
    f = make_26as(world / "a-26AS.xlsx", IND["name"], IND["pan"], fy="2024-25")
    assert entity_scope.check_26as_fy(f, FY)[0] == "mismatch"
    assert entity_scope.check_26as_fy(f, "2024-25") == ("ok", "")
    assert entity_scope.check_26as_fy(f, "") == ("ok", "")


# ---- partner run path ------------------------------------------------------

def _partner_run(tmp_path, monkeypatch, xlsx, *, financial_year=FY, advisory_fy=FY):
    d = tmp_path / "advices"
    d.mkdir()
    (d / "a.pdf").write_bytes(b"%PDF-1.4 fake")
    adv = tmp_path / "adv.pdf"
    adv.write_bytes(b"%PDF-1.4 fake")
    profile = SimpleNamespace(name=IND["name"], pan=IND["pan"], extra_items={},
                              partner_comp_accounts={})
    monkeypatch.setattr(partner_agent._advisory_parser, "parse",
                        lambda p, pw: {"financial_year": advisory_fy})
    monkeypatch.setattr(partner_agent._payout_advice_parser, "parse",
                        lambda p, pw: {"month": "2025-04", "source_name": "a.pdf",
                                       "total_paid": 480000.0, "remuneration": 200000.0,
                                       "share_of_profit_gross": 300000.0,
                                       "additional_share_of_profit": 0.0, "tds": -20000.0})
    monkeypatch.setattr(partner_agent, "_resolve_entity_config", lambda e, c: (profile, None))
    out = tmp_path / "o.xlsx"
    res = partner_agent.run(
        entity="TEST-IND", advices_dir=str(d), advisory_path=str(adv),
        xlsx_26as=str(xlsx) if xlsx else "", output_path=str(out), journal_path="",
        financial_year=financial_year)
    return res, out


def test_partner_refuses_another_entitys_26as(tmp_path, monkeypatch):     # NEGATIVE
    f = make_26as(tmp_path / "h-26AS.xlsx", HUF["name"], HUF["pan"])
    res, out = _partner_run(tmp_path, monkeypatch, f)
    assert res.startswith("ERROR") and "Test Family HUF" in res and not out.exists()


def test_partner_refuses_another_fys_26as(tmp_path, monkeypatch):         # NEGATIVE
    f = make_26as(tmp_path / "o-26AS.xlsx", IND["name"], IND["pan"], fy="2024-25")
    res, out = _partner_run(tmp_path, monkeypatch, f)
    assert res.startswith("ERROR") and "2024-25" in res and FY in res and not out.exists()


def test_partner_refuses_a_26as_that_states_no_identity(tmp_path, monkeypatch):   # NEGATIVE
    f = make_26as(tmp_path / "x-26AS.xlsx", "", "", with_identity=False)
    res, out = _partner_run(tmp_path, monkeypatch, f)
    assert res.startswith("ERROR") and "does not state whose it is" in res and not out.exists()


def test_partner_refuses_documents_of_another_year_than_selected(tmp_path, monkeypatch):  # NEGATIVE
    res, out = _partner_run(tmp_path, monkeypatch, None, financial_year="2024-25", advisory_fy=FY)
    assert res.startswith("ERROR") and not out.exists()


def test_partner_correct_entity_and_fy_still_passes(tmp_path, monkeypatch):
    f = make_26as(tmp_path / "ok-26AS.xlsx", IND["name"], IND["pan"], fy=FY)
    res, out = _partner_run(tmp_path, monkeypatch, f)
    assert "ownership" not in res and "belongs to" not in res
    assert not res.startswith("ERROR"), res
    assert out.exists()


def test_partner_without_a_selected_fy_still_checks_against_the_documents(tmp_path, monkeypatch):
    f = make_26as(tmp_path / "o-26AS.xlsx", IND["name"], IND["pan"], fy="2024-25")
    res, out = _partner_run(tmp_path, monkeypatch, f, financial_year="")
    assert res.startswith("ERROR") and "2024-25" in res and not out.exists()


# ---- (c) AIS: Entity required, and the export must be that entity's ---------

def test_ais_entity_is_required_and_the_picker_stays_entity_scoped():
    sk = _skill("skill_ais_reconcile")
    ent = next(i for i in sk.inputs if i.name == "entity")
    assert ent.required and sk.inputs[0].name == "entity"
    assert sk.run_args["entity_key"] == "{inputs.entity}"
    assert next(i for i in sk.inputs if i.name == "xlsx_path").entity_from == "entity"


PAN = "ABCDE1234F"
DOB = "1985-06-15"


def _ais_world(tmp_path, key="syn-one"):
    ent = tmp_path / "entities.yaml"
    ent.write_text(yaml.safe_dump({
        key: {"name": "Synthetic Taxpayer", "pan": PAN, "status": "Individual", "dob": DOB},
        "syn-two": {"name": "Other Taxpayer", "pan": "ZZZZZ9999Z", "status": "Individual",
                    "dob": DOB},
    }), encoding="utf-8")
    ais = tmp_path / "XXXDE1234X_2025-26_AIS.json"
    from agents.skill_ais_reconcile import decrypt as D
    from tests.test_ais_reconcile import synthetic_ais
    ais.write_text(D._encrypt_for_test(synthetic_ais(), D.derive_password_from_iso_date(PAN, DOB)),
                   encoding="utf-8")
    return ent, ais


def test_ais_cannot_run_without_an_entity(tmp_path):                      # NEGATIVE
    ent, ais = _ais_world(tmp_path)
    out = tmp_path / "o.xlsx"
    res = ais_agent.run(str(ais), str(out), entity_key="", entities_path=str(ent))
    assert res.startswith("ERROR") and "Entity" in res and not out.exists()


def test_ais_export_of_a_different_entity_is_refused(tmp_path):           # NEGATIVE
    ent, ais = _ais_world(tmp_path)
    out = tmp_path / "o.xlsx"
    res = ais_agent.run(str(ais), str(out), entity_key="syn-two", entities_path=str(ent))
    assert res.startswith("ERROR") and "syn-one" in res and "syn-two" in res
    assert not out.exists()


def test_ais_matching_entity_runs(tmp_path):
    ent, ais = _ais_world(tmp_path)
    out = tmp_path / "o.xlsx"
    res = ais_agent.run(str(ais), str(out), entity_key="syn-one", entities_path=str(ent))
    assert not res.startswith("ERROR"), res
    assert out.exists()


def test_ais_refuses_another_entitys_or_another_years_26as(tmp_path):     # NEGATIVE
    ent, ais = _ais_world(tmp_path)
    for name, pan, fy in (("Other Taxpayer", "ZZZZZ9999Z", FY),
                          ("Synthetic Taxpayer", PAN, "2024-25")):
        x = make_26as(tmp_path / f"{fy}-{pan}-26AS.xlsx", name, pan, fy=fy)
        out = tmp_path / f"o-{fy}-{pan}.xlsx"
        res = ais_agent.run(str(ais), str(out), entity_key="syn-one",
                            xlsx_path=str(x), entities_path=str(ent))
        assert res.startswith("ERROR") and not out.exists()


def test_ais_refuses_a_26as_without_identity(tmp_path):                   # NEGATIVE
    ent, ais = _ais_world(tmp_path)
    x = make_26as(tmp_path / "blank-26AS.xlsx", "", "", with_identity=False)
    out = tmp_path / "o.xlsx"
    res = ais_agent.run(str(ais), str(out), entity_key="syn-one",
                        xlsx_path=str(x), entities_path=str(ent))
    assert res.startswith("ERROR") and not out.exists()
