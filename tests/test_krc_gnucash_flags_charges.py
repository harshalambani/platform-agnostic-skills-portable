"""
KRC-01 (demat charges), KRC-03 (new-security / unbookable-sale flags),
UI-27 (Review tab lists every run) and UI-28 (one button per file the run
wrote) -- all synthetic data, nothing from Data/.
"""
from __future__ import annotations

import csv
import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "ui"))
sys.path.insert(0, str(ROOT / "src"))
SCRIPT = ROOT / "src" / "agents" / "skill_krc_gnucash" / "scripts" / "build_krc_gnucash.py"

_spec = importlib.util.spec_from_file_location("build_krc_gnucash_t", SCRIPT)
bk = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(bk)

_NS = ('xmlns:gnc="http://www.gnucash.org/XML/gnc" xmlns:act="http://www.gnucash.org/XML/act" '
       'xmlns:slot="http://www.gnucash.org/XML/slot"')


def _acct(name, aid, typ, parent, slots=()):
    s = [f'<gnc:account version="2.0.0"><act:name>{name}</act:name>',
         f'<act:id type="guid">{aid}</act:id><act:type>{typ}</act:type>']
    if parent:
        s.append(f'<act:parent type="guid">{parent}</act:parent>')
    if slots:
        s.append("<act:slots>" + "".join(
            f'<slot><slot:key>{k}</slot:key><slot:value type="string">{v}</slot:value></slot>'
            for k, v in slots) + "</act:slots>")
    s.append("</gnc:account>")
    return "".join(s)


def _book(path: Path):
    accts = [
        _acct("Root Account", "root", "ROOT", None),
        _acct("Assets", "as", "ASSET", "root", [("placeholder", "true")]),
        _acct("Investments", "inv", "ASSET", "as", [("placeholder", "true")]),
        _acct("Stocks", "stk", "ASSET", "inv", [("placeholder", "true")]),
        _acct("ALPHA STEEL", "s1", "STOCK", "stk"),
        _acct("BETA FOODS", "s2", "STOCK", "stk"),
        _acct("Expenses", "ex", "EXPENSE", "root", [("placeholder", "true")]),
        _acct("Demat Charges", "ex1", "EXPENSE", "ex"),
        _acct("Old Demat", "ex2", "EXPENSE", "ex", [("hidden", "true")]),
        _acct("Header Exp", "ex3", "EXPENSE", "ex", [("placeholder", "true")]),
    ]
    path.write_text(f'<?xml version="1.0"?><gnc-v2 {_NS}><gnc:book version="2.0.0">'
                    + "".join(accts) + "</gnc:book></gnc-v2>", encoding="utf-8")
    return path


BROKER = "Assets:Broker"
BILL_HDR = ["CN No", "Type", "Date", "Settlement No", "Security", "Quantity",
            "Net/Bill Amount", "Direction"]
REC_HDR = ["Date", "V.No", "Particulars", "ChqNo", "Debit", "Credit", "Balance", "Dr/Cr",
           "Tag (Part I)", "Category", "Bill CN", "Bill Settlement", "Match Status", "Notes"]


def _bills_xlsx(path: Path, bills, ledger, code=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bills"
    ws.append(BILL_HDR)
    for b in bills:
        ws.append(b)
    ws3 = wb.create_sheet("Reconciliation")
    ws3.append(["Synthetic reconciliation"])
    ws3.append([])
    ws3.append(REC_HDR)
    for date, part, deb, cred, bal, drcr, cat in ledger:
        ws3.append([date, "", part, "", deb, cred, bal, drcr, "", cat, "", "", "", ""])
    if code:
        w = wb.create_sheet("Run Info")
        w.append(["Field", "Value"])
        w.append(["Client Code", code])
    wb.save(path)
    return path


def _cfg(tmp_path):
    p = tmp_path / "cfg.yaml"
    if p.exists():
        return p
    p.write_text(f"accounts:\n  broker: {BROKER}\n  slbs_income: Income:SLBS\n"
                 "  ltcg: Income:LTCG\n  stcg: Income:STCG\n", encoding="utf-8")
    return p


def _run(tmp_path, bills, ledger, extra=(), code=None, outname="run-KRC-GnuCash"):
    book = _book(tmp_path / "b.gnucash")
    x = _bills_xlsx(tmp_path / "bills.xlsx", bills, ledger, code)
    out = tmp_path / outname
    r = subprocess.run([sys.executable, str(SCRIPT), str(x), str(book), str(out),
                        str(_cfg(tmp_path)), *extra], capture_output=True, text=True)
    return r, out


def _buy(cn, sec, qty, net):
    return [cn, "TRADE", "2025-06-01", "S1", sec, qty, net, "Payable"]


def _sell(cn, sec, qty, net):
    return [cn, "TRADE", "2025-07-01", "S2", sec, qty, net, "Receivable"]


# ledger: opening 0, pay-in 1000 (credit), buy ALPHA STEEL 600 (debit), demat 50 (debit)
GOOD_LEDGER = [
    ("2025-05-31", "Opening balance", None, None, 0.0, "Cr", "Opening Balance"),
    ("2025-06-01", "Bank receipt", None, 1000.0, 1000.0, "Cr", "Bank Pay-In"),
    ("2025-06-02", "Bill purchase", 600.0, None, 400.0, "Cr", "Trade Bill"),
    ("2025-06-03", "DP charges synthetic", 50.0, None, 350.0, "Cr", "Demat Charge"),
]
BUY = [_buy("C1", "ALPHA STEEL", 10, 600.0)]


# ---------------- KRC-01 ----------------

def test_charges_written_with_account_and_balanced(tmp_path):
    r, out = _run(tmp_path, BUY, GOOD_LEDGER,
                  ["--entity", "ent1", "--demat-account", "Expenses:Demat Charges"])
    assert r.returncode == 0, r.stdout + r.stderr
    rows = list(csv.reader(open(out / "Charges.csv", encoding="utf-8")))
    flat = " ".join(" ".join(x) for x in rows)
    assert "Expenses:Demat Charges" in flat and BROKER in flat
    assert "RED FLAG" not in r.stdout
    assert "closes to the ledger" in r.stdout


def test_no_charges_csv_without_account_and_red_flag(tmp_path):
    r, out = _run(tmp_path, BUY, GOOD_LEDGER)
    assert not (out / "Charges.csv").exists()
    assert "RED FLAG - Demat charges NOT booked" in r.stdout
    assert "will NOT close" in r.stdout or "not close to zero" in r.stdout
    assert r.returncode == 2
    assert r.stdout.lstrip().startswith("RED FLAG")  # flags come first


@pytest.mark.parametrize("acct", ["Expenses:Old Demat", "Expenses:Header Exp",
                                  "Expenses:Nope"])
def test_hidden_placeholder_or_unknown_account_never_used(tmp_path, acct):
    r, out = _run(tmp_path, BUY, GOOD_LEDGER, ["--demat-account", acct])
    assert not (out / "Charges.csv").exists()
    assert "not a usable" in r.stdout


def test_no_demat_rows_gives_no_charges_csv(tmp_path):
    led = [x for x in GOOD_LEDGER if x[6] != "Demat Charge"]
    r, out = _run(tmp_path, BUY, led, ["--demat-account", "Expenses:Demat Charges"])
    assert not (out / "Charges.csv").exists()
    assert r.returncode == 0, r.stdout


def test_charge_rows_never_in_other_files_and_bank_rows_never_charges(tmp_path):
    r, out = _run(tmp_path, BUY, GOOD_LEDGER, ["--demat-account", "Expenses:Demat Charges"])
    for n in ("Purchase.csv", "SLBM.csv", "Sale.csv"):
        f = out / n
        if f.exists():
            assert "Demat" not in f.read_text(encoding="utf-8")
    ch = (out / "Charges.csv").read_text(encoding="utf-8")
    assert "Bank receipt" not in ch and len({l.split(",")[1] for l in ch.splitlines()[1:]}) == 1


def test_remembered_per_entity_then_used_without_form_value(tmp_path):
    r1, _ = _run(tmp_path, BUY, GOOD_LEDGER,
                 ["--entity", "ent1", "--demat-account", "Expenses:Demat Charges"],
                 outname="r1-KRC-GnuCash")
    assert r1.returncode == 0
    r2, out2 = _run(tmp_path, BUY, GOOD_LEDGER, ["--entity", "ent1"], outname="r2-KRC-GnuCash")
    assert (out2 / "Charges.csv").exists()
    r3, out3 = _run(tmp_path, BUY, GOOD_LEDGER, ["--entity", "other"], outname="r3-KRC-GnuCash")
    assert not (out3 / "Charges.csv").exists()


def test_unbalanced_run_raises_red_flag_balanced_does_not(tmp_path):
    bad = list(GOOD_LEDGER) + [("2025-06-04", "Bill purchase 2", 100.0, None, 250.0, "Cr",
                                "Trade Bill")]
    r, _ = _run(tmp_path, BUY, bad, ["--demat-account", "Expenses:Demat Charges"])
    assert "RED FLAG - Broker account will NOT close" in r.stdout
    assert "-100.00" in r.stdout or "+100.00" in r.stdout
    ok, _ = _run(tmp_path, BUY, GOOD_LEDGER, ["--demat-account", "Expenses:Demat Charges"],
                 outname="ok-KRC-GnuCash")
    assert "Broker account will NOT close" not in ok.stdout


# ---------------- KRC-03 ----------------

def test_new_security_flagged_not_booked_and_listed(tmp_path):
    bills = [_buy("N1", "GAMMA PHARMA", 5, 300.0)]
    led = [("2025-06-01", "Bank receipt", None, 300.0, 300.0, "Cr", "Bank Pay-In"),
           ("2025-06-02", "Bill purchase", 300.0, None, 0.0, "Cr", "Trade Bill")]
    r, out = _run(tmp_path, bills, led)
    assert "FLAG - New security - create it in GnuCash first: GAMMA PHARMA" in r.stdout
    assert not (out / "Purchase.csv").exists()
    ns = list(csv.reader(open(out / "NewSecurities.csv", encoding="utf-8")))
    assert ns[1][0] == "GAMMA PHARMA" and ns[1][1] == "N1"
    assert ns[1][3] == "Assets:Investments:Stocks:GAMMA PHARMA"  # derived parent
    assert ns[1][4] == "Stock"
    info = json.loads((out / "run_info.json").read_text(encoding="utf-8"))
    assert info["new_securities"][0]["security"] == "GAMMA PHARMA"


def test_one_shared_token_is_not_booked_to_existing_stock(tmp_path):
    # shares only "ALPHA" with the existing ALPHA STEEL account
    bills = [_buy("N2", "ALPHA TEXTILES", 5, 300.0)]
    led = [("2025-06-02", "Bill purchase", 300.0, None, 0.0, "Cr", "Trade Bill")]
    r, out = _run(tmp_path, bills, led)
    assert not (out / "Purchase.csv").exists()
    assert "New security" in r.stdout and "ALPHA TEXTILES" in r.stdout
    pth, _ = bk.match_security("ALPHA TEXTILES", ["A:ALPHA STEEL"], {})
    assert pth is None


def test_fuzzy_match_is_written_but_listed_with_score(tmp_path):
    bills = [_buy("F1", "BETA FOODS PRODUCTS", 5, 300.0)]  # 2 of 3 tokens
    led = [("2025-06-02", "Bill purchase", 300.0, None, 0.0, "Cr", "Trade Bill")]
    r, out = _run(tmp_path, bills, led)
    assert (out / "Purchase.csv").exists()
    assert "matched by name similarity - check" in r.stdout
    assert "score 0.67" in r.stdout


def test_oversold_sale_not_written_and_red_flag(tmp_path):
    bills = [_buy("B1", "ALPHA STEEL", 10, 600.0), _sell("S9", "ALPHA STEEL", 50, 900.0)]
    led = [("2025-06-02", "Bill", 600.0, None, 0.0, "Cr", "Trade Bill")]
    r, out = _run(tmp_path, bills, led)
    assert "RED FLAG - Sale NOT booked: ALPHA STEEL, CN S9" in r.stdout
    assert "insufficient FIFO lots" in r.stdout
    assert not (out / "Sale.csv").exists()


def test_clean_sale_has_no_sale_flag(tmp_path):
    bills = [_buy("B1", "ALPHA STEEL", 10, 600.0), _sell("S1", "ALPHA STEEL", 10, 900.0)]
    led = [("2025-06-02", "Bill", 600.0, None, 0.0, "Cr", "Trade Bill")]
    r, out = _run(tmp_path, bills, led)
    assert "Sale NOT booked" not in r.stdout
    assert (out / "Sale.csv").exists()


def test_unreadable_quantity_sale_never_written_as_zero(tmp_path):
    bills = [_buy("B1", "ALPHA STEEL", 10, 600.0), _sell("S3", "ALPHA STEEL", None, 900.0)]
    led = [("2025-06-02", "Bill", 600.0, None, 0.0, "Cr", "Trade Bill")]
    r, out = _run(tmp_path, bills, led)
    assert "RED FLAG - Sale NOT booked: ALPHA STEEL, CN S3" in r.stdout
    assert not (out / "Sale.csv").exists()


def test_inconsistent_proceeds_sale_not_written(tmp_path):
    bills = [_buy("B1", "ALPHA STEEL", 10, 600.0), _sell("S4", "ALPHA STEEL", 5, -10.0)]
    led = [("2025-06-02", "Bill", 600.0, None, 0.0, "Cr", "Trade Bill")]
    r, out = _run(tmp_path, bills, led)
    assert "Sale NOT booked" in r.stdout and "inconsistent" in r.stdout
    assert not (out / "Sale.csv").exists()


def test_unmatched_sale_security_is_red_flag_not_new_security(tmp_path):
    bills = [_sell("S5", "OMEGA CORP", 5, 100.0)]
    led = [("2025-06-02", "Bill", None, 100.0, 100.0, "Cr", "Trade Bill")]
    r, out = _run(tmp_path, bills, led)
    assert "RED FLAG - Sale NOT booked: OMEGA CORP, CN S5" in r.stdout
    assert not (out / "NewSecurities.csv").exists()


def test_suggest_parent_derived_from_book_never_hardcoded():
    assert bk.suggest_parent(["X:Y:A", "X:Y:B"]) == "X:Y"
    assert bk.suggest_parent(["X:Y:A", "X:Z:B"]) == "X"
    assert bk.suggest_parent([]) is None


def test_client_code_recorded_in_run_info(tmp_path):
    r, out = _run(tmp_path, BUY, GOOD_LEDGER, ["--demat-account", "Expenses:Demat Charges"],
                  code="ZZ99")
    assert json.loads((out / "run_info.json").read_text(encoding="utf-8"))["client_code"] == "ZZ99"
    r2, out2 = _run(tmp_path, BUY, GOOD_LEDGER, outname="n-KRC-GnuCash")
    assert json.loads((out2 / "run_info.json").read_text(encoding="utf-8"))["client_code"] is None


# ---------------- UI-27 ----------------

def _mk_run(base: Path, name: str, mtime: int, review=False, info=None):
    d = base / name
    d.mkdir()
    (d / "Purchase.csv").write_text("x", encoding="utf-8")
    if review:
        (d / "Review.csv").write_text("CN No,Type,Security,Net,Reason\n", encoding="utf-8")
    if info is not None:
        (d / "run_info.json").write_text(json.dumps(info), encoding="utf-8")
    import os
    os.utime(d, (mtime, mtime))
    return d


def test_review_tab_lists_every_run_newest_first_and_labels_code(tmp_path, monkeypatch):
    from ui.tabs import krc_gnucash_review as m
    monkeypatch.setattr(m._config_mod, "output_dir", lambda: tmp_path)
    _mk_run(tmp_path, "old-KRC-GnuCash", 1_000_000, review=True, info={"client_code": "AA11"})
    _mk_run(tmp_path, "new-KRC-GnuCash", 2_000_000, review=False, info={"client_code": "BB22"})
    found = m._scan_review_csvs()
    assert [Path(v).name for _l, v in found] == ["new-KRC-GnuCash", "Review.csv"]
    assert "BB22" in found[0][0] and "nothing to review" in found[0][0]  # newer run not hidden
    assert "AA11" in found[1][0]
    assert found[0][0].startswith("new-")  # default (first) is the newest run


def test_review_tab_code_not_recorded_for_old_run(tmp_path, monkeypatch):
    from ui.tabs import krc_gnucash_review as m
    monkeypatch.setattr(m._config_mod, "output_dir", lambda: tmp_path)
    _mk_run(tmp_path, "r-KRC-GnuCash", 1_000_000, review=True)
    assert "code not recorded" in m._scan_review_csvs()[0][0]


def test_review_tab_run_without_review_says_nothing_to_review_and_lists_files(tmp_path, monkeypatch):
    from ui.tabs import krc_gnucash_review as m
    monkeypatch.setattr(m._config_mod, "output_dir", lambda: tmp_path)
    d = _mk_run(tmp_path, "n-KRC-GnuCash", 2_000_000,
                info={"client_code": None, "red_flags": ["RED FLAG - synthetic"], "flags": []})
    html = m._load_review_data(str(d), "")
    assert "Nothing to review in this run" in html
    assert "Purchase.csv" in html and "RED FLAG - synthetic" in html
    assert "code not recorded" in html


def test_older_review_not_default_when_newer_run_exists(tmp_path, monkeypatch):
    from ui.tabs import krc_gnucash_review as m
    monkeypatch.setattr(m._config_mod, "output_dir", lambda: tmp_path)
    _mk_run(tmp_path, "a-KRC-GnuCash", 1_000_000, review=True)
    _mk_run(tmp_path, "b-KRC-GnuCash", 3_000_000, review=False)
    default = m._scan_review_csvs()[0][1]
    assert not default.endswith("Review.csv")
    assert Path(default).name == "b-KRC-GnuCash"


# ---------------- UI-28 ----------------

def _skill():
    import agents.registry as reg
    sk = reg.get("KRC GnuCash")
    assert sk is not None
    return sk


def test_extra_outputs_declared_for_every_file():
    sk = _skill()
    keys = [x.key for x in sk.output.extra_outputs]
    assert keys == ["purchase_csv", "slbm_csv", "sale_csv", "charges_csv",
                    "new_securities_csv", "review_csv"]


def test_only_written_files_get_a_button_and_none_from_other_runs(tmp_path, monkeypatch):
    from ui.tabs import _generic as g
    from agents.outputs import ReplyWithOutputs, extra_output
    sk = _skill()
    root = tmp_path / "outputs"
    run = root / "r1-KRC-GnuCash"
    other = root / "r0-KRC-GnuCash"
    run.mkdir(parents=True)
    other.mkdir()
    (run / "Purchase.csv").write_text("p", encoding="utf-8")
    (other / "Sale.csv").write_text("OTHER RUN", encoding="utf-8")
    staging = tmp_path / "stage"
    staging.mkdir()
    monkeypatch.setattr(g._config, "output_dir", lambda: root)
    monkeypatch.setattr(g._config, "download_staging_dir", lambda: staging)
    reply = ReplyWithOutputs("x", (
        extra_output("purchase_csv", run / "Purchase.csv"),
        extra_output("sale_csv", other / "Sale.csv"),        # a different run's file
        extra_output("charges_csv", None, "Not written by this run."),
    ))
    res = g._stage_extra_outputs(sk, reply, run_dir=run, heading="Files written by this run")
    ups = res["updates"]
    keys = [x.key for x in sk.output.extra_outputs]
    on = {k for k, u in zip(keys, ups) if u.get("interactive")}
    assert on == {"purchase_csv"}                      # Sale from another run: no button
    assert "OTHER RUN" not in "".join(p.read_text() for p in staging.iterdir())
    assert "Files written by this run" in res["block"]
    assert not any((staging / n).exists() for n in ("Sale.csv", "Charges.csv"))
