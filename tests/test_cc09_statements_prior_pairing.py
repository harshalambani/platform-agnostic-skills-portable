"""CC-09: one Statements row per statement with its real tie-out; find_prior accepts a gap of 0;
same-amount payment lines pair one-to-one by date when the counts are equal. Synthetic only."""
from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from tests.test_cc_spend_booking import (APR, MAY, L, MOD, M, P, april_rows, book_pay, go, may_rows, stmt, tie)


# ---- (1) Statements sheet ---------------------------------------------------------------

def _statement_rows(out):
    wb = load_workbook(out.with_name(out.stem + "_report.xlsx"))
    ws = wb["Statements"]
    return [tuple(c.value for c in r) for r in ws.iter_rows(min_row=2)]


def test_an_untied_statement_has_one_row_and_never_shows_pass(tmp_path, monkeypatch):      # NEGATIVE
    s1, s2 = stmt(APR, april_rows()), stmt(MAY, may_rows())
    _r, _rows, _c, out = go(tmp_path, monkeypatch, [s1, s2],
                            [tie(s1, result="FAIL", printed=1900.0, gap="printed 1,900.00 vs rows 1,700.00"),
                             tie(s2, printed=300.0)], [book_pay("p1", 1900, date(2025, 5, 12))])
    rows = _statement_rows(out)
    apr = [r for r in rows if r[1] == s1.label]
    assert len(apr) == 1 and apr[0][2] == "FAIL"
    assert apr[0][3] == "SETTLED"                               # CC-26: an untied statement is imported, its payment settles it
    assert len({(r[0], r[1]) for r in rows}) == len(rows)       # one row per statement


def test_a_passing_chained_statement_shows_its_real_result(tmp_path, monkeypatch):
    s1, s2 = stmt(APR, april_rows()), stmt(MAY, may_rows())
    _r, _rows, _c, out = go(tmp_path, monkeypatch, [s1, s2],
                            [tie(s1, result="PASS (chained)", printed=1900.0), tie(s2, printed=300.0)],
                            [book_pay("p1", 1900, date(2025, 5, 12))])
    apr = [r for r in _statement_rows(out) if r[1] == s1.label]
    assert len(apr) == 1 and apr[0][2] == "PASS (chained)"


# ---- (2) find_prior ---------------------------------------------------------------------

def _prior(a, b):
    s_a, s_b = stmt(a, [], card="C"), stmt(b, [], card="C")
    return MOD.find_prior(s_b, [s_a, s_b]) is s_a


def test_a_gap_of_zero_cycle_change_is_a_predecessor():
    assert _prior((date(2025, 4, 1), date(2025, 4, 30)), (date(2025, 4, 30), date(2025, 5, 31)))
    assert _prior((date(2025, 4, 1), date(2025, 4, 30)), (date(2025, 5, 1), date(2025, 5, 31)))     # gap 1 still works


def test_a_two_day_overlap_is_not_a_predecessor():                                       # NEGATIVE
    assert not _prior((date(2025, 4, 1), date(2025, 4, 30)), (date(2025, 4, 29), date(2025, 5, 31)))
    assert not _prior((date(2025, 4, 1), date(2025, 4, 30)), (date(2025, 4, 10), date(2025, 5, 31)))


def test_a_long_gap_is_still_not_a_predecessor():                                        # NEGATIVE
    assert not _prior((date(2025, 1, 1), date(2025, 1, 31)), (date(2025, 6, 1), date(2025, 6, 30)))


# ---- (3) same-amount pairing ------------------------------------------------------------

def test_equal_counts_of_the_same_amount_pair_one_to_one_by_date():
    out = M.match_payments([L(1000, date(2025, 5, 10), card="C1"), L(1000, date(2025, 5, 20), card="C2")],
                           [P("b", 1000, date(2025, 5, 21)), P("a", 1000, date(2025, 5, 11))])
    pairs = {pr.line.stmt.card: pr.payments[0].guid for pr in out.pairings}
    assert pairs == {"C1": "a", "C2": "b"} and not out.ambiguous and not out.payments_unmatched


def test_equal_counts_never_use_one_bank_payment_twice():                                # NEGATIVE
    out = M.match_payments([L(1000, date(2025, 5, 10)), L(1000, date(2025, 5, 12)), L(1000, date(2025, 5, 14))],
                           [P("a", 1000, date(2025, 5, 10)), P("b", 1000, date(2025, 5, 12)),
                            P("c", 1000, date(2025, 5, 14))])
    guids = [pr.payments[0].guid for pr in out.pairings]
    assert sorted(guids) == ["a", "b", "c"] and len(set(guids)) == 3


def test_unequal_counts_stay_ambiguous_and_say_why():                                    # NEGATIVE
    out = M.match_payments([L(1000, date(2025, 5, 10), card="C1"), L(1000, date(2025, 5, 11), card="C2")],
                           [P("a", 1000, date(2025, 5, 10)), P("b", 1000, date(2025, 5, 11)),
                            P("c", 1000, date(2025, 5, 12))])
    assert not out.pairings and out.ambiguous
    assert any("counts differ" in a["reason"] for a in out.ambiguous)


def test_a_pair_outside_the_window_is_not_force_paired():                                # NEGATIVE
    out = M.match_payments([L(1000, date(2025, 5, 10)), L(1000, date(2025, 5, 11))],
                           [P("a", 1000, date(2025, 5, 10)), P("b", 1000, date(2025, 6, 30))])
    assert len(out.pairings) <= 1 and not any(pr.payments[0].guid == "b" for pr in out.pairings)
