"""
tests/test_imp08_bank_account_pick.py -- IMP-08: several of the owner's own
accounts match the bank name, and the pipeline must NEVER guess one.

Synthetic book, real naming shape (see gnc_book_fixture): four HDFC accounts
(three postable, one hidden), two HSBC accounts, "Root Account:" root.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
for _p in (ROOT.parent / "src", ROOT.parent / "src" / "agents"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import gnc_book_fixture as fx  # noqa: E402
from agents.skill_gnucash_pipeline import agent as pipe  # noqa: E402
from agents.skill_hsbc.acct_number import find_account_number  # noqa: E402

Y1234, Y5678 = fx.P_HDFC1, fx.P_HDFC2


def _bal(acc, paise, date="2025-03-01"):
    return fx.txn_xml("OPENING", date, [(acc, paise), ("obe", -paise)])


def _book(tmp_path, txns, accounts=None, name="b.gnucash"):
    return fx.write_book(tmp_path / name, accounts or fx.standard_accounts(), txns)


def _rows(opening=100000.00, narr="NEFT SALARY CREDIT"):
    return [{"Date": "2025-04-01", "Transaction ID": "T1", "Description": narr,
             "Account": "", "Deposit": "500.00", "Withdrawal": "",
             "Balance": f"{opening + 500.00:.2f}", "Currency": "INR"}]


def _hdfc_csv(tmp_path):
    sp = str(ROOT / "skill_hdfc")
    if sp not in sys.path:
        sys.path.insert(0, sp)
    import hdfc_fixture_gen
    stmt = tmp_path / "stmt.csv"
    stmt.write_text(hdfc_fixture_gen.build_csv_text(), encoding="utf-8")
    return stmt


def test_two_postable_matches_no_number_no_evidence_stops(tmp_path):
    book = _book(tmp_path, [_bal(fx.HDFC1, 25000000), _bal(fx.HDFC2, 7000000)])
    res = pipe._get_gnucash_account_balance(book, "HDFC")
    assert res["found"] is False and res["ambiguous"] is True
    assert res["account_name"] == ""
    assert sorted(res["candidates"]) == sorted([Y1234, Y5678, fx.P_HDFC4])
    assert not any("(old)" in c for c in res["candidates"])


def test_reconcile_stop_returns_unfiltered_rows_and_candidates(tmp_path):
    book = _book(tmp_path, [_bal(fx.HDFC1, 25000000)])
    rec = pipe._reconcile_opening_balance(_rows(), book, "HDFC", None)
    assert rec["ok"] is False and rec["stop"] is True
    assert rec["filtered_rows"] == _rows()
    assert len(rec["candidates"]) == 3


def test_opening_balance_picks_the_one_matching_candidate(tmp_path):
    book = _book(tmp_path, [_bal(fx.HDFC1, 25000000), _bal(fx.HDFC2, 10000000)])
    res = pipe._get_gnucash_account_balance(
        book, "HDFC", opening_balance=100000.00, start_date="2025-04-01")
    assert res["found"] is True and res["ambiguous"] is False
    assert res["account_name"].endswith("094XXXX5678")
    assert "opening balance" in res["match_note"]


def test_evidence_fitting_two_candidates_still_stops(tmp_path):
    book = _book(tmp_path, [_bal(fx.HDFC1, 10000000), _bal(fx.HDFC2, 10000000)])
    res = pipe._get_gnucash_account_balance(
        book, "HDFC", opening_balance=100000.00, start_date="2025-04-01")
    assert res["found"] is False and res["ambiguous"] is True


def test_hidden_candidate_is_never_picked_even_first_in_book_order(tmp_path):
    accs = fx.standard_accounts()
    hidden = next(a for a in accs if "094XXXX9012" in a)
    accs.remove(hidden)
    accs.insert(4, hidden)
    # only the hidden account carries the opening balance the statement shows
    book = _book(tmp_path, [_bal(fx.HDFC_OLD, 10000000), _bal(fx.HDFC1, 25000000)],
                 accounts=accs)
    res = pipe._get_gnucash_account_balance(
        book, "HDFC", opening_balance=100000.00, start_date="2025-04-01")
    assert "(old)" not in res["account_name"]
    assert res["found"] is False and res["ambiguous"] is True
    assert all("(old)" not in c for c in res["candidates"])


def test_opening_balance_disagreeing_with_chosen_account_is_not_accepted(tmp_path):
    book = _book(tmp_path, [_bal(fx.HDFC1, 25000000), _bal(fx.HDFC2, 10000000)])
    rec = pipe._reconcile_opening_balance(
        _rows(100000.00), book, "HDFC", None, chosen_account=Y1234)
    assert rec["ok"] is False
    assert "OPENING BALANCE MISMATCH" in rec["message"]


def test_narration_carrying_one_accounts_number_picks_it(tmp_path):
    book = _book(tmp_path, [_bal(fx.HDFC1, 1), _bal(fx.HDFC2, 2)])
    res = pipe._get_gnucash_account_balance(
        book, "HDFC", narrations=["NEFT FROM SELF A/C 3456 REF"])
    assert res["found"] is True and res["account_name"].endswith("094XXXX3456")


def test_narrations_naming_two_accounts_still_stop(tmp_path):
    book = _book(tmp_path, [_bal(fx.HDFC1, 1)])
    res = pipe._get_gnucash_account_balance(
        book, "HDFC", narrations=["XFER 1234 TO 5678"])
    assert res["found"] is False and res["ambiguous"] is True


def test_year_like_digits_in_narration_are_not_evidence(tmp_path):
    accs = [a for a in fx.standard_accounts() if "094XXXX3456" not in a] + [
        fx.account_xml("y2025", "HDFC Bank - 094XXXX2025", "BANK", "cab")]
    book = _book(tmp_path, [], accounts=accs)
    res = pipe._get_gnucash_account_balance(
        book, "HDFC", narrations=["INTEREST FOR FY 2025"])
    assert res["found"] is False and res["ambiguous"] is True


def test_single_account_resolves_without_prompt(tmp_path):
    accs = [a for a in fx.standard_accounts() if "HDFC" not in a] + [
        fx.account_xml("only", "HDFC Bank - 094XXXX7777", "BANK", "cab")]
    book = _book(tmp_path, [_bal("only", 500000)], accounts=accs)
    res = pipe._get_gnucash_account_balance(book, "HDFC")
    assert res["found"] is True and res["ambiguous"] is False
    assert res["match_warning"] is None and res["balance"] == 5000.00


def test_account_number_still_resolves_directly(tmp_path):
    book = _book(tmp_path, [_bal(fx.HDFC1, 1), _bal(fx.HDFC2, 2)])
    res = pipe._get_gnucash_account_balance(book, "HDFC", "094XXXX5678")
    assert res["found"] is True and res["account_name"].endswith("5678")


def test_account_number_matching_nothing_does_not_silently_fall_back(tmp_path):
    book = _book(tmp_path, [_bal(fx.HDFC1, 1)])
    res = pipe._get_gnucash_account_balance(book, "HDFC", "99999999")
    assert res["found"] is False and res["account_name"] == ""
    assert "not resolving" in res["match_warning"]


@pytest.mark.parametrize("choice,expect", [
    (fx.P_HDFC_OLD, "hidden"),
    ("Assets:Current Assets:Cash and Bank", "not an account at"),
    ("Expenses:Food:Dining", "not an account at"),
    ("Assets:Nowhere:HDFC Bank - 1", "not an account in the GnuCash book"),
    (fx.P_HSBC1, "not an account at"),
])
def test_supplied_bank_account_that_is_not_acceptable_is_refused(tmp_path, choice, expect):
    book = _book(tmp_path, [_bal(fx.HDFC1, 1)])
    res = pipe._get_gnucash_account_balance(book, "HDFC", chosen_account=choice)
    assert res["found"] is False and res["refused"] is True, res
    assert res["account_name"] == ""
    assert expect in res["match_warning"].lower() or expect in res["match_warning"]


def test_supplied_valid_bank_account_wins_over_contrary_evidence(tmp_path):
    book = _book(tmp_path, [_bal(fx.HDFC1, 25000000), _bal(fx.HDFC2, 10000000)])
    res = pipe._get_gnucash_account_balance(
        book, "HDFC", chosen_account="Root Account:" + Y1234,
        opening_balance=100000.00, start_date="2025-04-01")
    assert res["found"] is True and res["account_name"].endswith("1234")


def test_postable_bank_accounts_lists_only_postable_at_that_bank(tmp_path):
    book = _book(tmp_path, [])
    got = pipe.postable_bank_accounts(book, "HDFC")
    assert sorted(got) == sorted([Y1234, Y5678, fx.P_HDFC4])
    assert all("(old)" not in g for g in got)
    assert not any("HSBC" in g for g in got)
    assert pipe.postable_bank_accounts(str(tmp_path / "missing.gnucash"), "HDFC") == []


def test_run_stops_and_asks_writes_no_output(tmp_path):
    stmt = _hdfc_csv(tmp_path)
    book = _book(tmp_path, [_bal(fx.HDFC1, 25000000), _bal(fx.HDFC2, 7000000)])
    out = tmp_path / "out.csv"
    msg = pipe.run(bank="HDFC", statement_files=str(stmt), gnucash_file=book,
                   output_path=str(out), config_path=None)
    assert not out.exists()
    assert "Which HDFC account" in msg and "Nothing was written" in msg
    assert "094XXXX5678" in msg and "(old)" not in msg


def test_run_refuses_hidden_bank_account_and_writes_no_output(tmp_path):
    stmt = _hdfc_csv(tmp_path)
    book = _book(tmp_path, [_bal(fx.HDFC1, 25000000)])
    out = tmp_path / "out.csv"
    msg = pipe.run(bank="HDFC", statement_files=str(stmt), gnucash_file=book,
                   output_path=str(out), config_path=None,
                   bank_account=fx.P_HDFC_OLD)
    assert not out.exists()
    assert "hidden" in msg.lower()


# -- HSBC statement account-number reading ---------------------------------

def test_hsbc_labelled_account_number_is_read():
    assert find_account_number(["Account Number 013065123456"]) == "013065123456"
    assert find_account_number(["A/C No. 013 065 123 456"]) == "013065123456"


def test_hsbc_bare_digit_run_is_never_guessed():
    assert find_account_number(["REF 013065123456 something"]) is None


def test_hsbc_two_different_labelled_numbers_is_ambiguous_none():
    assert find_account_number(
        ["Account Number 013065123456", "Account No 013065999999"]) is None


def test_hsbc_short_number_is_ignored():
    assert find_account_number(["Account Number 1234"]) is None


def test_hsbc_enriched_workbook_summary_number_is_read(tmp_path):
    import openpyxl
    from agents.skill_hsbc.agent import _read_account_number
    wb = openpyxl.Workbook()
    wb.active.title = "Transactions"
    s = wb.create_sheet("Summary")
    s.append(["Account Number", "013065 123456"])
    p = tmp_path / "e.xlsx"
    wb.save(p)
    assert _read_account_number(str(p)) == "013065123456"
    wb2 = openpyxl.Workbook()
    wb2.active.append(["nothing here"])
    q = tmp_path / "n.xlsx"
    wb2.save(q)
    assert _read_account_number(str(q)) is None


# -- wiring -------------------------------------------------------------

def test_skill_yaml_declares_optional_bank_account_input_and_run_arg():
    from agents.registry import get
    sk = get("gnucash_pipeline")
    inp = next(i for i in sk.inputs if i.name == "bank_account")
    assert inp.required is False and inp.options_from == "bank_accounts"
    assert inp.depends_on == ("gnucash_file", "bank")
    assert sk.run_args["bank_account"] == "{inputs.bank_account}"
