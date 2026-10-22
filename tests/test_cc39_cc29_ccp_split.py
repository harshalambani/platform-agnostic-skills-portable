"""CC-39 (per-card Credit Card Payment sub-accounts) and CC-29 (corporate card prepayment to Drawings).
Synthetic data only; neutral last4 values. The statements, the book and the mapper are the fakes of
test_cc_spend_booking."""
from __future__ import annotations

import datetime
from datetime import date

import pytest
from openpyxl import load_workbook

from agents import cc_cards as cc
from agents.skill_cc_spend_booking import agent as ccsb
from agents.skill_cc_spend_booking import bookio
from agents.skill_cc_spend_booking import journal as J
from agents.skill_cc_spend_booking import split as SP
from agents.skill_cc_spend_booking.bookio import BookTxn, CCP_PATH
from tests.test_cc_spend_booking import (APR, BANK, DRAW, FOOD, MAY, PATHS, Guard, april_rows, book_pay, go,
                                         may_rows, stmt, tie)
from ui.tabs import _entity_cards as ec

SUB = CCP_PATH + ":Regalia"
SUB2 = CCP_PATH + ":Other"
ENT = {"SYN"}
_ORIG_RUN = ccsb.run


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _cards_text(account=SUB, split="2025-05-01", prepay=DRAW, last4="1111", extra_card=""):
    top = ""
    if split:
        top += f'ccp_split_from: "{split}"\n'
    if prepay:
        top += f'prepayment_account: "{prepay}"\n'
    card = f'  - bank: "HDFC"\n    card: "Regalia"\n    last4: ["{last4}"]\n    holder: ""\n    active: true\n'
    if account:
        card += f'    account: "{account}"\n'
    return top + "cards:\n" + card + extra_card


@pytest.fixture
def paths(monkeypatch):
    """The fake book's account paths and the fake guard's blocks, restored after each test."""
    added = {SUB, SUB2}
    before = set(PATHS)
    PATHS.update(added)
    monkeypatch.setattr(Guard, "blocked", {})
    yield PATHS
    PATHS.clear()
    PATHS.update(before)


def _run(tmp_path, monkeypatch, cards_text=None, txns=None, marks="", s1_rows=None):
    if cards_text is not None:
        (tmp_path / "cc_cards.yaml").write_text(cards_text, encoding="utf-8")
    monkeypatch.setattr(ccsb, "run", lambda *a, **k: _ORIG_RUN(*a, prepayment_marks=marks, **k))
    s1, s2 = stmt(APR, s1_rows or april_rows()), stmt(MAY, may_rows())
    txns = [book_pay("p1", 1900, date(2025, 5, 12))] if txns is None else txns
    return go(tmp_path, monkeypatch, [s1, s2], [tie(s1, printed=1900.0), tie(s2, printed=300.0)], txns)


def _amt(r):
    return float(r["Amount (Withdrawal)"] or 0) - float(r["Amount Negated (Deposit)"] or 0)


def _moves(rows):
    return [r for r in rows if r["Description"].startswith("Moves card payment")]


def _report(out, sheet):
    ws = load_workbook(out.with_name(out.stem + "_report.xlsx"))[sheet]
    return [[c.value for c in r] for r in ws.iter_rows(min_row=1)]


def _core(rows):
    return [(r["Date"], r["Transaction ID"], r["Description"], r["Account"], r["Transfer Account"],
             r["Amount (Withdrawal)"], r["Amount Negated (Deposit)"]) for r in rows]


# ---------------------------------------------------------------------------
# constants and the loader
# ---------------------------------------------------------------------------

def test_the_pool_path_is_the_same_in_the_loader_and_the_book_reader():
    assert cc.CCP_PATH == bookio.CCP_PATH


def test_new_keys_load():
    cards = cc.parse_text(_cards_text(), ENT)
    assert cards[0].account == SUB
    assert cards.split_from == datetime.date(2025, 5, 1)
    assert cards.prepayment_account == DRAW


def test_an_old_file_loads_exactly_as_before():
    cards = cc.parse_text('cards:\n  - bank: "HDFC"\n    card: "Regalia"\n    last4: ["1111"]\n'
                          '    holder: ""\n    active: true\n', ENT)
    assert cards[0].account == "" and cards.split_from is None and cards.prepayment_account == ""


def _refused(text):
    with pytest.raises(cc.CardListError) as e:
        cc.parse_text(text, ENT)
    assert "1111" not in str(e.value)
    return str(e.value)


def test_an_account_outside_the_pool_is_refused():                                      # NEGATIVE
    _refused(_cards_text(account="Expense:Food"))
    _refused(_cards_text(account=CCP_PATH))                  # the parent itself is not a sub-account
    _refused(_cards_text(account="Expense:Withdrawals:Credit Card Payments:X"))   # look-alike prefix
    _refused(_cards_text(account=CCP_PATH + ":"))


def test_a_non_text_account_is_refused():                                               # NEGATIVE
    _refused(_cards_text().replace(f'account: "{SUB}"', "account: 12"))
    _refused(_cards_text().replace(f'account: "{SUB}"', "account: true"))


def test_two_cards_sharing_one_account_are_refused():                                   # NEGATIVE
    second = ('  - bank: "ICICI"\n    card: "Amazon"\n    last4: ["2222"]\n    holder: ""\n    active: true\n'
              f'    account: "{SUB}"\n')
    msg = _refused(_cards_text(extra_card=second))
    assert "share an account" in msg


def test_a_bad_split_date_is_refused():                                                 # NEGATIVE
    _refused(_cards_text().replace('ccp_split_from: "2025-05-01"', "ccp_split_from: 2025-05-01"))   # unquoted date
    _refused(_cards_text(split="2025-13-01"))
    _refused(_cards_text(split="01/05/2025"))
    _refused(_cards_text().replace('ccp_split_from: "2025-05-01"', "ccp_split_from: 20250501"))


def test_a_prepayment_account_inside_the_pool_is_refused():                             # NEGATIVE
    _refused(_cards_text(prepay=SUB))
    _refused(_cards_text(prepay=CCP_PATH))


def test_an_unknown_top_level_key_is_still_refused():                                   # NEGATIVE
    _refused("surprise: 1\n" + _cards_text())


def test_render_round_trips_the_new_keys():
    rows = [{"bank": "HDFC", "card": "Regalia", "last4": ["1111"], "holder": "", "active": True, "account": SUB},
            {"bank": "ICICI", "card": "Amazon", "last4": ["2222"], "holder": "", "active": True}]
    text = cc.render_yaml(rows, cc.TOP_COMMENT, "2025-05-01", DRAW)
    cards = cc.parse_text(text, ENT)
    assert [c.account for c in cards] == [SUB, ""] and cards.split_from == datetime.date(2025, 5, 1)
    assert "account:" not in text.split('card: "Amazon"')[1]                      # a card without one gets no key


# ---------------------------------------------------------------------------
# the Cards screen
# ---------------------------------------------------------------------------

def test_the_screen_keeps_every_key_on_save(tmp_path):
    p = tmp_path / "cc_cards.yaml"
    p.write_text(_cards_text(), encoding="utf-8")
    assert ec.save_card(p, ENT, 0, "HDFC-Regalia", "1111", "", False).startswith("**Saved")     # account untouched
    cards = cc.parse_text(p.read_text(encoding="utf-8"), ENT)
    assert cards[0].account == SUB and cards[0].active is False
    assert cards.split_from == datetime.date(2025, 5, 1) and cards.prepayment_account == DRAW
    assert list(tmp_path.glob("cc_cards.yaml.bak-*"))                                           # timestamped backup


def test_the_screen_edits_the_account_and_the_top_settings(tmp_path):
    p = tmp_path / "cc_cards.yaml"
    p.write_text(_cards_text(account="", split="", prepay=""), encoding="utf-8")
    assert ec.save_card(p, ENT, 0, "HDFC-Regalia", "1111", "", True, SUB).startswith("**Saved")
    assert ec.save_top(p, ENT, "2025-05-01", DRAW).startswith("**Saved")
    cards = cc.parse_text(p.read_text(encoding="utf-8"), ENT)
    assert cards[0].account == SUB and cards.split_from == datetime.date(2025, 5, 1)
    assert ec.read_top(p, ENT) == ("2025-05-01", DRAW)
    assert ec.save_top(p, ENT, "", "").startswith("**Saved")                                    # blank clears
    assert ec.read_top(p, ENT) == ("", "")
    assert cc.parse_text(p.read_text(encoding="utf-8"), ENT)[0].account == SUB


@pytest.mark.parametrize("acct", ["Expense:Food", CCP_PATH, 'X"\nbad: 1'])
def test_the_screen_refuses_an_account_outside_the_pool_and_leaves_the_file(tmp_path, acct):  # NEGATIVE
    p = tmp_path / "cc_cards.yaml"
    p.write_text(_cards_text(), encoding="utf-8")
    before = p.read_bytes()
    assert "Not saved" in ec.save_card(p, ENT, 0, "HDFC-Regalia", "1111", "", True, acct)
    assert p.read_bytes() == before


@pytest.mark.parametrize("split,prepay", [("2025-5-1", DRAW), ("2025-05-01", SUB), ("not a date", "")])
def test_the_screen_refuses_bad_top_settings_and_leaves_the_file(tmp_path, split, prepay):   # NEGATIVE
    p = tmp_path / "cc_cards.yaml"
    p.write_text(_cards_text(), encoding="utf-8")
    before = p.read_bytes()
    assert "Not saved" in ec.save_top(p, ENT, split, prepay)
    assert p.read_bytes() == before


def test_the_screen_refuses_a_shared_account(tmp_path):                                      # NEGATIVE
    p = tmp_path / "cc_cards.yaml"
    p.write_text(_cards_text(), encoding="utf-8")
    assert "Not saved" in ec.save_card(p, ENT, None, "ICICI-Amazon", "2222", "", True, SUB)
    assert cc.parse_text(p.read_text(encoding="utf-8"), ENT)[0].account == SUB and len(
        cc.parse_text(p.read_text(encoding="utf-8"), ENT)) == 1


# ---------------------------------------------------------------------------
# CC-39 in the run
# ---------------------------------------------------------------------------

def test_a_paired_payment_after_the_cutoff_is_moved_and_the_statement_books_to_the_sub_account(
        tmp_path, monkeypatch, paths):
    _r, rows, _c, out = _run(tmp_path, monkeypatch, _cards_text())
    mv = _moves(rows)
    assert len(mv) == 1
    m = mv[0]
    assert m["Account"] == SUB and m["Transfer Account"] == CCP_PATH
    assert _amt(m) == 1900.0 and m["Date"] == "12/05/2025"
    spends = [r for r in rows if r not in mv]
    assert spends and all(r["Transfer Account"] == SUB for r in spends)
    assert not any(r["Transfer Account"] == CCP_PATH for r in spends)
    # the pool ties: parent + sub-account = what the parent alone showed
    t = {r[0]: r[1] for r in _report(out, "CCP tie-out") if r[0]}
    assert t["Difference (must be 0)"] == 0


def test_the_pool_tie_out_reads_the_same_as_the_unsplit_run(tmp_path, monkeypatch, paths):
    _r, _rows, _c, out = _run(tmp_path, monkeypatch, _cards_text())
    split_tie = {r[0]: r[1] for r in _report(out, "CCP tie-out") if r[0]}
    d2 = tmp_path / "plain"
    d2.mkdir()
    _r2, _rows2, _c2, out2 = _run(d2, monkeypatch, None)
    plain_tie = {r[0]: r[1] for r in _report(out2, "CCP tie-out") if r[0]}
    assert split_tie["These journals"] == plain_tie["These journals"]


def test_before_the_cutoff_nothing_changes(tmp_path, monkeypatch, paths):                    # NEGATIVE + golden
    _r, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(split="2025-06-01"))
    d2 = tmp_path / "plain"
    d2.mkdir()
    _r2, rows2, _c2, _o2 = _run(d2, monkeypatch, None)
    assert not _moves(rows) and not any(SUB in r["Transfer Account"] + r["Account"] for r in rows)
    assert _core(rows) == _core(rows2)


def test_no_cutoff_set_means_no_split(tmp_path, monkeypatch, paths):                         # NEGATIVE
    _r, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(split=""))
    assert not _moves(rows) and not any(SUB in r["Transfer Account"] for r in rows)


def test_a_card_without_an_account_books_exactly_as_today(tmp_path, monkeypatch, paths):     # NEGATIVE + golden
    _r, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(account=""))
    d2 = tmp_path / "plain"
    d2.mkdir()
    _r2, rows2, _c2, _o2 = _run(d2, monkeypatch, None)
    assert _core(rows) == _core(rows2)
    assert all(r["Transfer Account"] == CCP_PATH for r in rows)


def test_no_pairing_no_moving_row(tmp_path, monkeypatch, paths):                              # NEGATIVE
    # the only bank payment matches no statement line (wrong amount), so nothing is settled or moved
    _r, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(), txns=[book_pay("p1", 1234, date(2025, 5, 12))])
    assert not _moves(rows) and not any(SUB in r["Transfer Account"] for r in rows)


def test_a_payment_that_is_in_the_pool_but_unpaired_stays_in_the_pool_and_is_reported(
        tmp_path, monkeypatch, paths):                                                        # NEGATIVE
    reply, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(),
                               txns=[book_pay("p1", 1900, date(2025, 5, 12)), book_pay("stray", 777, date(2025, 5, 20))])
    assert len(_moves(rows)) == 1 and _amt(_moves(rows)[0]) == 1900.0
    assert "777" in reply and reply.startswith("RED FLAG")           # reported, not guessed onto the card


@pytest.mark.parametrize("why", ["hidden", "placeholder", "dormant"])
def test_an_unusable_sub_account_is_refused_and_books_nothing_for_the_card(
        tmp_path, monkeypatch, paths, why):                                                   # NEGATIVE
    Guard.blocked = {SUB: f"it is {why}"}
    reply, rows, _c, out = _run(tmp_path, monkeypatch, _cards_text())
    assert "RED FLAG" in reply and SUB in reply
    assert rows == []                                              # no fallback to the parent, no moving row
    assert not any(r["Transfer Account"] == CCP_PATH for r in rows)
    refused = _report(out, "Card sub-accounts") if out.with_name(out.stem + "_report.xlsx").is_file() else []
    assert not refused or any("REFUSED" in str(c) for r in refused for c in r)


def test_a_sub_account_that_is_not_in_the_book_is_refused(tmp_path, monkeypatch, paths):      # NEGATIVE
    PATHS.discard(SUB)
    reply, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text())
    assert "RED FLAG" in reply and rows == []


def test_mixed_dates_are_not_split_and_are_reported(tmp_path, monkeypatch, paths):            # NEGATIVE
    # cut-off falls after the payment: the (single) payment is before it, so the statement stays on the parent
    _r, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(split="2025-05-13"))
    assert not _moves(rows) and all(r["Transfer Account"] == CCP_PATH for r in rows)


def test_a_statement_already_booked_against_the_parent_is_not_split_later(
        tmp_path, monkeypatch, paths):                                                        # NEGATIVE
    _r, rows, _c, _o = _run(tmp_path, monkeypatch, None)         # first run, unsplit
    booked = [BookTxn(f"b{i}", date(2025, 5, 12), r["Description"], r["Transaction ID"],
                      ((r["Account"], 1.0), (CCP_PATH, -1.0)))
              for i, r in enumerate(rows)]
    d2 = tmp_path / "second"
    d2.mkdir()
    _r2, rows2, _c2, _o2 = _run(d2, monkeypatch, _cards_text(),
                                txns=[book_pay("p1", 1900, date(2025, 5, 12))] + booked)
    assert not _moves(rows2) and not any(SUB in r["Transfer Account"] for r in rows2)


def test_a_rerun_does_not_book_the_split_twice(tmp_path, monkeypatch, paths):                 # NEGATIVE
    _r, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text())
    booked = [BookTxn(f"b{i}", date(2025, 5, 12), r["Description"], r["Transaction ID"],
                      ((r["Account"], _amt(r)), (r["Transfer Account"], -_amt(r))))
              for i, r in enumerate(rows)]
    d2 = tmp_path / "again"
    d2.mkdir()
    _r2, rows2, _c2, _o2 = _run(d2, monkeypatch, _cards_text(), txns=[book_pay("p1", 1900, date(2025, 5, 12))] + booked)
    assert rows2 == []


def test_the_pool_tie_out_fails_loud_when_parent_plus_subs_do_not_equal_the_old_figure(
        tmp_path, monkeypatch, paths):                                                        # NEGATIVE
    real = J.move_journal

    def off_by_one(*a, **k):
        j = real(*a, **k)
        j.amount = round(j.amount + 1.0, 2)
        return j

    monkeypatch.setattr(J, "move_journal", off_by_one)
    reply, rows, _c, out = _run(tmp_path, monkeypatch, _cards_text())
    assert reply.startswith("RED FLAG") and "does not tie" in reply
    assert rows == [] and not out.is_file()                         # no import file is written


def test_pool_tie_unit_catches_a_leg_that_leaves_the_pool():                                 # NEGATIVE
    j = J.Journal(num="CCSB-x", kind=J.K_SPEND, date=date(2025, 5, 1), real_date=date(2025, 5, 1), description="d",
                  amount=100.0, side="W", card="C", statement="s", account=FOOD, transfer="Assets:Elsewhere")
    t = SP.pool_tie([j], CCP_PATH)
    assert t["ok"] is False and t["gap"] == 100.0
    j.transfer = SUB
    assert SP.pool_tie([j], CCP_PATH)["ok"] is True


def test_a_mapper_choice_inside_the_pool_is_never_booked(tmp_path, monkeypatch, paths):      # NEGATIVE
    from tests import test_cc_spend_booking as T
    monkeypatch.setitem(T.MAP_OVERRIDE, "SWIGGY", (SUB, "high", "fake"))
    _r, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text())
    assert not any(r["Account"] == SUB and r["Description"].upper().startswith("SWIGGY") for r in rows)


# ---------------------------------------------------------------------------
# CC-29 in the run
# ---------------------------------------------------------------------------

STRAY = book_pay("stray", 777, date(2025, 5, 20))
CODE = J.prepay_code("stray")


def _txns(*extra):
    return [book_pay("p1", 1900, date(2025, 5, 12)), *extra]


def test_an_unmarked_unmatched_payment_produces_no_row_and_shows_its_code(tmp_path, monkeypatch, paths):  # NEGATIVE
    reply, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(), txns=_txns(STRAY))
    assert CODE in reply
    assert not any(_amt(r) == 777.0 for r in rows) and not any(r["Account"] == DRAW for r in rows)


def test_a_marked_payment_gets_one_drawings_row(tmp_path, monkeypatch, paths):
    reply, rows, _c, out = _run(tmp_path, monkeypatch, _cards_text(), txns=_txns(STRAY), marks=CODE)
    pre = [r for r in rows if r["Account"] == DRAW]
    assert len(pre) == 1 and pre[0]["Transfer Account"] == CCP_PATH and _amt(pre[0]) == 777.0
    assert pre[0]["Date"] == "20/05/2025"
    assert "Corporate card prepayments marked (1)" in reply and not reply.startswith("RED FLAG")
    assert [r for r in _report(out, "Prepayments")][1:]
    # not also moved to a card
    assert not any(_amt(r) == 777.0 and r["Account"] == SUB for r in rows)
    t = {r[0]: r[1] for r in _report(out, "CCP tie-out") if r[0]}
    assert t["Difference (must be 0)"] == 0


def test_marking_works_without_any_card_having_an_account(tmp_path, monkeypatch, paths):
    _r, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(account=""), txns=_txns(STRAY), marks=CODE)
    assert len([r for r in rows if r["Account"] == DRAW]) == 1


def test_a_paired_payment_cannot_be_marked(tmp_path, monkeypatch, paths):                    # NEGATIVE
    reply, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(), txns=_txns(STRAY),
                               marks=J.prepay_code("p1"))
    assert not any(r["Account"] == DRAW for r in rows)
    assert "paired to a statement" in reply
    assert len(_moves(rows)) == 1                                   # still paired and moved, once


def test_a_marked_payment_is_never_also_paired_or_moved(tmp_path, monkeypatch, paths):       # NEGATIVE
    # the stray payment has the same amount as the statement line, but the statement already took p1
    twin = book_pay("twin", 1900, date(2025, 5, 13))
    _r, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(), txns=_txns(twin), marks=J.prepay_code("twin"))
    uses = [r for r in rows if _amt(r) == 1900.0 and (r["Account"] == DRAW or r["Account"] == SUB)]
    assert len(uses) == 2                                           # one move (p1) and one prepayment (twin)
    assert sorted(r["Account"] for r in uses) == sorted([DRAW, SUB])
    nums = [r["Transaction ID"] for r in rows]
    assert len(nums) == len(set(nums))


def test_a_payment_before_the_cutoff_cannot_be_marked(tmp_path, monkeypatch, paths):         # NEGATIVE
    early = book_pay("early", 555, date(2025, 4, 20))
    reply, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(), txns=_txns(early), marks=J.prepay_code("early"))
    assert not any(r["Account"] == DRAW for r in rows) and "before ccp_split_from" in reply


def test_an_unknown_or_malformed_code_books_nothing(tmp_path, monkeypatch, paths):           # NEGATIVE
    reply, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(), txns=_txns(STRAY), marks="deadbeef")
    assert not any(r["Account"] == DRAW for r in rows) and "deadbeef" in reply
    d2 = tmp_path / "bad"
    d2.mkdir()
    reply2, rows2, _c2, _o2 = _run(d2, monkeypatch, _cards_text(), txns=_txns(STRAY), marks="not-a-code")
    assert reply2.startswith("ERROR") and rows2 == []


@pytest.mark.parametrize("why", ["hidden", "placeholder", "dormant"])
def test_an_unusable_prepayment_account_is_refused_and_books_nothing(tmp_path, monkeypatch, paths, why):  # NEGATIVE
    Guard.blocked = {DRAW: f"it is {why}"}
    reply, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(), txns=_txns(STRAY), marks=CODE)
    assert reply.startswith("ERROR") and rows == []


def test_a_prepayment_account_missing_from_the_book_is_refused(tmp_path, monkeypatch, paths):  # NEGATIVE
    reply, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(prepay="Equity:Nowhere"), txns=_txns(STRAY),
                               marks=CODE)
    assert reply.startswith("ERROR") and rows == []


def test_marks_need_the_cutoff_and_the_account_to_be_set(tmp_path, monkeypatch, paths):      # NEGATIVE
    reply, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(prepay=""), txns=_txns(STRAY), marks=CODE)
    assert reply.startswith("ERROR") and "prepayment_account" in reply and rows == []
    d2 = tmp_path / "nocut"
    d2.mkdir()
    reply2, rows2, _c2, _o2 = _run(d2, monkeypatch, _cards_text(split=""), txns=_txns(STRAY), marks=CODE)
    assert reply2.startswith("ERROR") and "ccp_split_from" in reply2 and rows2 == []


def test_a_prepayment_already_in_the_book_is_read_back_not_rebooked(tmp_path, monkeypatch, paths):  # NEGATIVE
    done = BookTxn("pre1", date(2025, 5, 20), "Corporate card prepayment", J.prepay_num("stray"),
                   ((DRAW, 777.0), (CCP_PATH, -777.0)))
    prior = BookTxn("stray", date(2025, 5, 20), "HDFC CC PAYMENT", "", ((CCP_PATH, 777.0), (BANK, -777.0)))
    reply, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(), txns=_txns(prior, done), marks=CODE)
    assert not any(r["Account"] == DRAW for r in rows)             # not booked a second time
    assert not reply.startswith("RED FLAG") and "Already booked as prepayments earlier (1)" in reply


def test_changing_the_prepayment_account_cannot_double_book(tmp_path, monkeypatch, paths):   # NEGATIVE
    assert J.prepay_num("stray") == J.prepay_num("stray")
    assert DRAW not in J.prepay_num("stray") and "Other" not in J.prepay_num("stray")


def test_no_marks_means_the_run_is_unchanged(tmp_path, monkeypatch, paths):                  # NEGATIVE + golden
    _r, rows, _c, _o = _run(tmp_path, monkeypatch, _cards_text(account="", split="", prepay=""), txns=_txns(STRAY))
    d2 = tmp_path / "plain"
    d2.mkdir()
    _r2, rows2, _c2, _o2 = _run(d2, monkeypatch, None, txns=_txns(STRAY))
    assert _core(rows) == _core(rows2)


def test_marks_parser():
    assert SP.parse_marks(" AB12CD34, ab12cd34;ef567890 ") == ["ab12cd34", "ef567890"]
    assert SP.parse_marks("") == [] and SP.valid_code("ab12cd34") and not SP.valid_code("ab12cd3")
