"""tests/test_cc27_csv_payment_match.py -- CC-27: the payment rows of an ICICI CSV export are matched to the
bank payments (same amount tolerance, same date window, one-to-one) so a bank payment in a CSV-only period
is no longer reported as "payment with no statement". Payment rows are still never booked.

Synthetic only: made-up card groups and amounts.
"""
from __future__ import annotations

import re
from datetime import date

from tests import test_cc22_icici_csv as t22
from tests.test_cc24_csv_booking import BLOCKS, _book, _res, _src
from tests.test_cc_spend_booking import MAY, book_pay, go, may_rows, stmt, tie

MC = t22.MC
PAY = "PAYMENT RECEIVED THANK YOU"


def _blocks(*pays):
    return [(MC, [("05-Apr-25", "SHOP ONE MUMBAI", "1,200.50")] + [(d, PAY, f"-{a}") for d, a in pays]),
            (t22.AMEX, [("06-Apr-25", "CAFE TWO PUNE", "300.00")])]


def _run(tmp_path, monkeypatch, blocks, txns):
    d = _src(tmp_path, blocks=blocks)
    return _book(tmp_path / "w", monkeypatch, _res(d), txns)


def _unmatched_total(reply):
    m = re.search(r"of which payments with no statement: ([\d,\.\-]+)", reply)
    return float(m.group(1).replace(",", "")) if m else None


def test_a_bank_payment_matching_a_csv_payment_row_is_matched(tmp_path, monkeypatch):
    txns = [book_pay("p1", 5000, date(2025, 4, 8))]
    reply, rows, _c, _o = _run(tmp_path / "a", monkeypatch, _blocks(("07-Apr-25", "5,000.00")), txns)
    assert "RED FLAG" not in reply
    assert "CSV PAYMENT ROWS MATCHED TO BANK PAYMENTS" in reply
    assert (_unmatched_total(reply) or 0.0) == 0.0
    assert "[matched to a CSV payment row]" in reply
    assert len(rows) == 2                                          # only the spend; the payment is not booked


def test_without_the_csv_row_the_same_payment_is_unmatched(tmp_path, monkeypatch):      # NEGATIVE
    txns = [book_pay("p1", 5000, date(2025, 4, 8))]
    reply, _rows, _c, _o = _run(tmp_path, monkeypatch, _blocks(), txns)
    assert "RED FLAG" in reply and "MATCHED TO BANK PAYMENTS" not in reply
    assert _unmatched_total(reply) == 5000.0


def test_a_csv_payment_row_with_no_bank_payment_is_reported_and_never_booked(tmp_path, monkeypatch):      # NEGATIVE
    reply, rows, _c, _o = _run(tmp_path, monkeypatch, _blocks(("07-Apr-25", "5,000.00")), [])
    assert "CSV PAYMENT ROWS WITH NO BANK PAYMENT IN THE BOOK" in reply
    assert "NOT complete" in reply
    assert len(rows) == 2
    assert not any("PAYMENT" in r["Description"].upper() for r in rows)


def test_one_bank_payment_never_matches_two_csv_rows(tmp_path, monkeypatch):            # NEGATIVE
    txns = [book_pay("p1", 5000, date(2025, 4, 8))]
    blocks = _blocks(("07-Apr-25", "5,000.00"), ("08-Apr-25", "5,000.00"))
    reply, rows, _c, _o = _run(tmp_path, monkeypatch, blocks, txns)
    matched = reply.split("CSV PAYMENT ROWS MATCHED TO BANK PAYMENTS")[-1].split("\n\n")[0]
    assert matched.count("<-") <= 1
    assert "AMBIGUOUS" in reply or "NO BANK PAYMENT" in reply     # the other row is not silently matched
    assert len(rows) == 2


def test_one_bank_payment_never_matches_a_pdf_line_and_a_csv_row(tmp_path, monkeypatch):       # NEGATIVE
    s2 = stmt(MAY, may_rows())
    pdf_pay = next(r for r in s2.parsed.rows if r.kind == "payment")
    amt = pdf_pay.amount
    res = _res(_src(tmp_path, blocks=[(MC, [("05-Apr-25", "SHOP ONE MUMBAI", "10.00")]),
                                      (t22.AMEX, [("12-May-25", PAY, f"-{amt:.2f}")])]))
    res.statements, res.ties, res.pool = [s2], [tie(s2, printed=0.0)], [s2]
    txns = [book_pay("p1", amt, date(2025, 5, 12))]
    (tmp_path / "w").mkdir(exist_ok=True)
    reply, rows, _c, _o = go(tmp_path / "w", monkeypatch, [], [], txns, result_override=res)
    # the one bank payment fits two payment lines of different sources: it is never paired with both
    assert "AMBIGUOUS payments" in reply and "the bank payment could settle 2 statement lines" in reply
    assert "CSV PAYMENT ROWS MATCHED TO BANK PAYMENTS" not in reply
    assert not any("PAYMENT" in r["Description"].upper() for r in rows)


def test_an_amount_just_outside_the_tolerance_is_not_matched(tmp_path, monkeypatch):    # NEGATIVE
    txns = [book_pay("p1", 5000.02, date(2025, 4, 8))]
    reply, _rows, _c, _o = _run(tmp_path, monkeypatch, _blocks(("07-Apr-25", "5,000.00")), txns)
    assert "MATCHED TO BANK PAYMENTS" not in reply
    assert "CSV PAYMENT ROWS WITH NO BANK PAYMENT" in reply and "RED FLAG" in reply


def test_a_date_outside_the_window_is_not_matched(tmp_path, monkeypatch):               # NEGATIVE
    txns = [book_pay("p1", 5000, date(2025, 4, 20))]
    reply, _rows, _c, _o = _run(tmp_path, monkeypatch, _blocks(("07-Apr-25", "5,000.00")), txns)
    assert "MATCHED TO BANK PAYMENTS" not in reply
    assert "CSV PAYMENT ROWS WITH NO BANK PAYMENT" in reply


def test_no_payment_row_is_ever_booked(tmp_path, monkeypatch):                          # NEGATIVE
    txns = [book_pay("p1", 5000, date(2025, 4, 8))]
    _r, rows, _c, out = _run(tmp_path, monkeypatch, _blocks(("07-Apr-25", "5,000.00")), txns)
    assert all("PAYMENT" not in r["Description"].upper() for r in rows)
    assert all(r["Amount Negated (Deposit)"] in ("", None) or float(r["Amount Negated (Deposit)"]) != 5000.0
               for r in rows)


def test_the_payments_sheet_marks_the_csv_source(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    txns = [book_pay("p1", 5000, date(2025, 4, 8))]
    _r, _rows, _c, out = _run(tmp_path, monkeypatch, _blocks(("07-Apr-25", "5,000.00")), txns)
    ws = load_workbook(out.with_name(out.stem + "_report.xlsx"))["Payments matched"]
    vals = [str(c.value) for row in ws.iter_rows(min_row=2) for c in row]
    assert any("CSV export" in v for v in vals) and "MATCHED (CSV)" in vals
