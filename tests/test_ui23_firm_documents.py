"""UI-23 (drop zone that sorts the firm's documents by content) and UI-21
(drawn samples under the manual pickers).

Everything is synthetic: a "PDF" here is a small text file, and the parsers
and the page-text reader are replaced by fakes that read that text. No real
name, PAN or account number appears anywhere.
"""
import csv
import re
import sys
import zipfile
from pathlib import Path

import pytest
import yaml

ROOT = Path(__file__).resolve().parent.parent
for _p in (ROOT, ROOT / "src"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

from agents.registry import discover  # noqa: E402
from agents.skill_partner_comp_recon import agent as agent_module  # noqa: E402
from agents.skill_partner_comp_recon import intake  # noqa: E402
from agents.skill_partner_comp_recon.parsers import advisory as advisory_parser  # noqa: E402
from agents.skill_partner_comp_recon.parsers import llp_statement as llp_parser  # noqa: E402
from agents.skill_partner_comp_recon.parsers import payment_schedule as schedule_parser  # noqa: E402
from agents.skill_partner_comp_recon.parsers import payout_advice as payout_parser  # noqa: E402

FY = "2025-26"

ADVISORY_26 = "Compensation summary : Year ended 31 March 2026\nDate : 10 April 2026\n"
ADVISORY_25 = "Compensation summary : Year ended 31 March 2025\nDate : 10 April 2025\n"
ADVISORY_27 = "Compensation summary : Year ended 31 March 2027\nDate : 10 April 2027\n"
TARGET = "Target compensation advice\nDate : 01 April 2026\n"
FORWARD = "Compensation summary FY2026-2027 (forward looking)\n"


def payout_text(month):
    return f"PAYOUT {month}\n"


def _fake_read(path, password=None):
    t = Path(path).read_text(encoding="utf-8", errors="replace")
    return None if t.startswith("BROKEN") else t


def _rec_payout(month):
    return {"month": month, "source_name": f"cert-{month}.pdf", "total_paid": 480000.0,
            "remuneration": 200000.0, "share_of_profit_gross": 300000.0,
            "additional_share_of_profit": 0.0, "tds": -20000.0}


def _fake_payout(path, password=None):
    t = _fake_read(path)
    if t.startswith("EXPLODE"):
        raise RuntimeError("synthetic parser failure")
    if t.startswith("PAYOUT "):
        return _rec_payout(t.split()[1])
    raise payout_parser.NotAnL1DocumentError("not a payout")


def _fake_advisory(path, password=None):
    t = _fake_read(path)
    m = re.search(r"Year ended 31 March (\d{4})", t)
    if not m:
        raise advisory_parser.NotAnL3DocumentError("not an advisory")
    y = int(m.group(1))
    return {"financial_year": f"{y - 1}-{str(y)[-2:]}"}


def _fake_llp(path, password=None):
    t = _fake_read(path)
    if t.startswith("LLPSTATEMENT"):
        return {"financial_year": t.split()[1]}
    raise llp_parser.NotAnL5DocumentError("not an LLP statement")


def _fake_schedule(path, password=None):
    t = _fake_read(path)
    if t.startswith("SCHEDULE"):
        return {"financial_year": t.split()[1]}
    raise schedule_parser.NotAPaymentScheduleError("not a schedule")


@pytest.fixture(autouse=True)
def fakes(monkeypatch):
    monkeypatch.setattr(intake._precheck, "read_first_page_text", _fake_read)
    monkeypatch.setattr(payout_parser, "parse", _fake_payout)
    monkeypatch.setattr(advisory_parser, "parse", _fake_advisory)
    monkeypatch.setattr(llp_parser, "parse", _fake_llp)
    monkeypatch.setattr(schedule_parser, "parse", _fake_schedule)


def folder(tmp_path, files, name="firm"):
    d = tmp_path / name
    d.mkdir()
    for fname, text in files.items():
        p = d / fname
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding="utf-8")
    return d


def sort(value, manual=None, password=None):
    res = intake.sort_documents(value, FY, password, manual)
    return res


def row(res, name_part):
    return next(r for r in res.rows if name_part in r.name)


def used(res, name_part):
    return row(res, name_part).outcome.startswith("used for")


def full_set():
    files = {f"p{m}.pdf": payout_text(m) for m in
             ("2025-04", "2025-05", "2025-06")}
    files["adv.pdf"] = ADVISORY_26
    return files


# ---------------------------------------------------------------- sorting

def test_documents_are_sorted_by_content_not_by_file_name(tmp_path):
    # Every file is named for a DIFFERENT kind of document than it is.
    d = folder(tmp_path, {"advisory.pdf": payout_text("2025-04"),
                          "payout.pdf": ADVISORY_26,
                          "schedule.pdf": "LLPSTATEMENT 2025-26\n",
                          "statement.pdf": "SCHEDULE 2025-26\n"})
    res = sort(str(d))
    try:
        assert used(res, "advisory.pdf") and "payout" in row(res, "advisory.pdf").outcome
        assert res.advisory.endswith("payout.pdf")
        assert res.llp.endswith("schedule.pdf")
        assert res.schedule.endswith("statement.pdf")
        assert [Path(p).name for p in res.advices] == ["advisory.pdf"]
    finally:
        intake.cleanup(res)


def test_target_letter_is_never_used_as_the_advisory(tmp_path):            # NEGATIVE
    d = folder(tmp_path, {**full_set(), "adv.pdf": TARGET, "target.pdf": TARGET})
    res = sort(str(d))
    try:
        assert res.advisory == ""
        assert not used(res, "adv.pdf") and not used(res, "target.pdf")
        assert "Target compensation letter" in row(res, "target.pdf").recognised
        assert "Still missing: this year's (FY 2025-26) Advisory." in res.missing
    finally:
        intake.cleanup(res)


def test_forward_summary_is_never_used_as_this_years_advisory(tmp_path):   # NEGATIVE
    d = folder(tmp_path, {"fwd.pdf": FORWARD, "next.pdf": ADVISORY_27})
    res = sort(str(d))
    try:
        assert res.advisory == "" and res.award == []
        assert not used(res, "fwd.pdf") and not used(res, "next.pdf")
        assert "forward-looking" in row(res, "fwd.pdf").why
        assert "after the reporting year" in row(res, "next.pdf").why
    finally:
        intake.cleanup(res)


def test_earlier_year_advisory_is_never_this_years_advisory(tmp_path):     # NEGATIVE
    d = folder(tmp_path, {"old.pdf": ADVISORY_25})
    res = sort(str(d))
    try:
        assert res.advisory == ""
        assert [Path(p).name for p in res.award] == ["old.pdf"]
        assert "award-year documents (FY 2024-25)" in row(res, "old.pdf").outcome
    finally:
        intake.cleanup(res)


def test_superseded_revision_is_never_used_by_letter_date(tmp_path):       # NEGATIVE
    d = folder(tmp_path, {"a.pdf": ADVISORY_26,
                          "b.pdf": "Compensation summary : Year ended 31 March 2026\nDate : 20 May 2026\n"})
    res = sort(str(d))
    try:
        assert res.advisory.endswith("b.pdf")
        assert not used(res, "a.pdf")
        assert "superseded by b.pdf" in row(res, "a.pdf").why
    finally:
        intake.cleanup(res)


def test_superseded_revision_is_never_used_by_revision_number(tmp_path):   # NEGATIVE
    base = "Compensation summary : Year ended 31 March 2026\nDate : 10 April 2026\n"
    d = folder(tmp_path, {"a.pdf": base + "Revision 1\n", "b.pdf": base + "Revision 2\n"})
    res = sort(str(d))
    try:
        assert res.advisory.endswith("b.pdf") and not used(res, "a.pdf")
        assert "highest revision number" in row(res, "b.pdf").why
    finally:
        intake.cleanup(res)


def test_unorderable_revisions_stop_the_run_naming_both_files(tmp_path):   # NEGATIVE
    d = folder(tmp_path, {"a.pdf": ADVISORY_26 + "one\n", "b.pdf": ADVISORY_26 + "two\n"})
    res = sort(str(d))
    try:
        assert "a.pdf" in res.stop and "b.pdf" in res.stop
        assert "by hand" in res.stop
        assert res.advisory == ""
    finally:
        intake.cleanup(res)


def test_unorderable_revisions_return_an_error_and_run_nothing(tmp_path, monkeypatch):   # NEGATIVE
    d = folder(tmp_path, {**full_set(), "b.pdf": ADVISORY_26 + "two\n"})
    called = []
    monkeypatch.setattr(agent_module, "_run_from_documents", lambda **kw: called.append(kw))
    out = agent_module.run(entity="X", firm_documents=str(d), financial_year=FY,
                           output_path=str(tmp_path / "o.xlsx"))
    assert "ERROR" in out and "adv.pdf" in out and "b.pdf" in out
    assert called == []


def test_identical_copies_are_not_a_conflict(tmp_path):
    d = folder(tmp_path, {"a.pdf": ADVISORY_26, "copy of a.pdf": ADVISORY_26})
    res = sort(str(d))
    try:
        assert res.stop == "" and res.advisory
        assert sum(used(res, n) for n in ("a.pdf", "copy of a.pdf")) == 1
    finally:
        intake.cleanup(res)


def test_unrecognised_file_is_listed_as_skipped_and_never_used(tmp_path):  # NEGATIVE
    d = folder(tmp_path, {**full_set(), "mystery.pdf": "just a menu\n", "notes.txt": "hello"})
    res = sort(str(d))
    try:
        assert not used(res, "mystery.pdf") and not used(res, "notes.txt")
        assert row(res, "mystery.pdf").outcome == "skipped"
        assert "not recognised" in row(res, "mystery.pdf").why
        assert "not a PDF or .eml" in row(res, "notes.txt").why
        assert all("mystery" not in p for p in res.advices + res.award + [res.advisory])
    finally:
        intake.cleanup(res)


def test_same_file_is_never_assigned_to_two_roles(tmp_path):               # NEGATIVE
    d = folder(tmp_path, {**full_set(), "llp.pdf": "LLPSTATEMENT 2025-26\n",
                          "sch.pdf": "SCHEDULE 2025-26\n", "old.pdf": ADVISORY_25})
    res = sort(str(d))
    try:
        paths = res.advices + res.award + [res.advisory, res.llp, res.schedule]
        paths = [p for p in paths if p]
        assert len(paths) == len(set(paths)) == 7
        assert len({r.name for r in res.rows}) == len(res.rows)
        assert res.missing == []
    finally:
        intake.cleanup(res)


def test_payout_outside_the_year_is_skipped_using_the_fy_months(tmp_path):   # NEGATIVE
    d = folder(tmp_path, {"in.pdf": payout_text("2025-04"), "out.pdf": payout_text("2026-04")})
    res = sort(str(d))
    try:
        assert used(res, "in.pdf") and not used(res, "out.pdf")
        assert "outside FY 2025-26" in row(res, "out.pdf").why
        assert [Path(p).name for p in res.advices] == ["in.pdf"]
    finally:
        intake.cleanup(res)


def test_payout_parser_failing_unexpectedly_is_reported_not_swallowed(tmp_path):   # NEGATIVE
    d = folder(tmp_path, {"boom.pdf": "EXPLODE\n", "ok.pdf": payout_text("2025-04")})
    res = sort(str(d))
    try:
        r = row(res, "boom.pdf")
        assert r.outcome == "skipped" and "synthetic parser failure" in r.why
        assert "Could not be read" in r.recognised
        assert res.advices and not any("boom" in p for p in res.advices)
    finally:
        intake.cleanup(res)


def test_unopenable_file_is_skipped_with_its_reason(tmp_path):             # NEGATIVE
    d = folder(tmp_path, {"bad.pdf": "BROKEN\n"})
    res = sort(str(d))
    try:
        assert row(res, "bad.pdf").outcome == "skipped"
        assert row(res, "bad.pdf").why
    finally:
        intake.cleanup(res)


def test_eml_is_read_as_the_llp_statement(tmp_path, monkeypatch):
    d = folder(tmp_path, {"mail.eml": "LLPSTATEMENT 2025-26\n"})
    monkeypatch.setattr(intake._precheck, "extract_pdf_from_eml", lambda p: p)
    res = sort(str(d))
    try:
        assert res.llp.endswith("mail.eml") and used(res, "mail.eml")
    finally:
        intake.cleanup(res)


def test_llp_for_another_year_is_not_used(tmp_path):                       # NEGATIVE
    d = folder(tmp_path, {"llp.pdf": "LLPSTATEMENT 2024-25\n"})
    res = sort(str(d))
    try:
        assert res.llp == "" and not used(res, "llp.pdf")
        assert "not FY 2025-26" in row(res, "llp.pdf").why
    finally:
        intake.cleanup(res)


def test_missing_roles_are_listed_in_the_table_text(tmp_path):
    d = folder(tmp_path, {"p.pdf": payout_text("2025-04")})
    res = sort(str(d))
    try:
        text = intake.inputs_text(res)
        assert text.splitlines()[0].startswith("**Inputs")
        for h in ("File", "Recognised as", "Year read from inside", "Used for / skipped and why"):
            assert h in text
        assert "Still missing: this year's (FY 2025-26) Advisory." in text
        assert "Still missing: LLP Statement of Account." in text
        assert "Still missing: monthly payout documents." not in text
    finally:
        intake.cleanup(res)


# ------------------------------------------------------- manual pick wins

def test_a_manual_pick_is_never_overridden(tmp_path):                      # NEGATIVE
    d = folder(tmp_path, full_set())
    res = sort(str(d), manual={"advisory": True, "payout": True})
    try:
        assert res.advisory == "" and res.advices == []
        assert row(res, "adv.pdf").outcome == "skipped"
        assert "a manual pick always wins" in row(res, "adv.pdf").why
        assert "Still missing: this year's" not in " ".join(res.missing)
    finally:
        intake.cleanup(res)


def test_run_hands_the_picker_file_to_the_report_not_the_dropped_one(tmp_path, monkeypatch):   # NEGATIVE
    d = folder(tmp_path, full_set())
    picked = tmp_path / "picked.pdf"
    picked.write_text(ADVISORY_26, encoding="utf-8")
    seen = {}
    monkeypatch.setattr(agent_module, "_run_from_documents", lambda **kw: seen.update(kw) or "ok")
    out = agent_module.run(entity="X", firm_documents=str(d), advisory_path=str(picked),
                           financial_year=FY, output_path=str(tmp_path / "o.xlsx"))
    assert seen["advisory_path"] == str(picked)
    assert "a manual pick always wins" in out
    assert seen["inputs_rows"]


def test_run_without_the_drop_zone_is_the_old_path(tmp_path, monkeypatch):  # NEGATIVE
    seen = {}
    monkeypatch.setattr(agent_module, "_run_from_documents", lambda **kw: seen.update(kw) or "ok")
    out = agent_module.run(entity="X", advices_dir="a", advisory_path="b",
                           financial_year=FY, output_path=str(tmp_path / "o.xlsx"))
    assert out == "ok" and seen["inputs_rows"] is None


def test_required_role_missing_from_the_folder_is_named(tmp_path, monkeypatch):
    d = folder(tmp_path, {"p.pdf": payout_text("2025-04")})
    called = []
    monkeypatch.setattr(agent_module, "_run_from_documents", lambda **kw: called.append(kw))
    out = agent_module.run(entity="X", firm_documents=str(d), financial_year=FY,
                           output_path=str(tmp_path / "o.xlsx"))
    assert "ERROR" in out and "Advisory" in out and called == []
    assert "| p.pdf |" in out


def test_drop_zone_needs_the_financial_year(tmp_path):
    d = folder(tmp_path, full_set())
    res = intake.sort_documents(str(d), "", None, {})
    try:
        assert "Financial year" in res.stop
    finally:
        intake.cleanup(res)


# ------------------------------------------------------------------- zips

def _zip(tmp_path, members, name="docs.zip"):
    p = tmp_path / name
    with zipfile.ZipFile(p, "w") as zf:
        for n, data in members.items():
            if isinstance(data, zipfile.ZipInfo):
                zf.writestr(data, "x")
            else:
                zf.writestr(n, data)
    return p


def test_zip_is_unpacked_and_sorted(tmp_path):
    z = _zip(tmp_path, {"inner/p.pdf": payout_text("2025-04"), "adv.pdf": ADVISORY_26})
    res = sort(str(z))
    try:
        assert res.advisory and len(res.advices) == 1
    finally:
        intake.cleanup(res)


def test_zip_member_with_parent_path_is_rejected(tmp_path):                # NEGATIVE
    z = _zip(tmp_path, {"../evil.pdf": ADVISORY_26, "ok.pdf": payout_text("2025-04")})
    res = sort(str(z))
    try:
        r = row(res, "evil.pdf")
        assert r.outcome == "skipped" and "leaves the archive folder" in r.why
        assert res.advisory == ""
        assert not (tmp_path.parent / "evil.pdf").exists()
        assert not any("evil" in p for p in res.advices)
    finally:
        intake.cleanup(res)


@pytest.mark.parametrize("member", ["/abs/evil.pdf", "C:/evil.pdf", "a\\..\\evil.pdf"])
def test_zip_absolute_and_backslash_paths_are_rejected(tmp_path, member):  # NEGATIVE
    z = _zip(tmp_path, {member: ADVISORY_26})
    res = sort(str(z))
    try:
        assert res.advisory == ""
        assert "leaves the archive folder" in row(res, "evil.pdf").why
    finally:
        intake.cleanup(res)


def test_zip_link_member_is_rejected(tmp_path):                            # NEGATIVE
    info = zipfile.ZipInfo("link.pdf")
    info.external_attr = (0o120777 << 16)
    z = _zip(tmp_path, {"link.pdf": info})
    res = sort(str(z))
    try:
        assert "a link inside the zip" in row(res, "link.pdf").why
        assert res.advices == [] and res.advisory == ""
    finally:
        intake.cleanup(res)


def test_zip_total_size_cap_stops_the_run(tmp_path, monkeypatch):          # NEGATIVE
    monkeypatch.setattr(intake, "MAX_TOTAL_BYTES", 100)
    z = _zip(tmp_path, {"a.pdf": "x" * 200})
    res = sort(str(z))
    try:
        assert "larger than the" in res.stop
        assert res.advisory == "" and res.advices == []
    finally:
        intake.cleanup(res)


def test_zip_declared_size_lie_is_caught_while_streaming(tmp_path, monkeypatch):   # NEGATIVE
    monkeypatch.setattr(intake, "MAX_TOTAL_BYTES", 100)
    z = _zip(tmp_path, {"a.pdf": "x" * 50, "b.pdf": "y" * 60})
    res = sort(str(z))
    try:
        assert "larger than the" in res.stop
    finally:
        intake.cleanup(res)


def test_file_count_cap_stops_the_run(tmp_path, monkeypatch):              # NEGATIVE
    monkeypatch.setattr(intake, "MAX_FILES", 3)
    z = _zip(tmp_path, {f"p{i}.pdf": payout_text("2025-04") for i in range(5)})
    res = sort(str(z))
    try:
        assert "more than 3 files" in res.stop
    finally:
        intake.cleanup(res)


def test_nested_archive_and_other_members_are_skipped_not_opened(tmp_path):   # NEGATIVE
    z = _zip(tmp_path, {"inner.zip": "PK", "notes.docx": "x", "adv.pdf": ADVISORY_26})
    res = sort(str(z))
    try:
        assert "nested archive" in row(res, "inner.zip").why
        assert "not a PDF or .eml" in row(res, "notes.docx").why
        assert res.advisory
    finally:
        intake.cleanup(res)


def test_unreadable_zip_is_listed_as_skipped(tmp_path):                    # NEGATIVE
    z = tmp_path / "broken.zip"
    z.write_bytes(b"not a zip")
    res = sort(str(z))
    try:
        assert "not a readable zip" in row(res, "broken.zip").why
    finally:
        intake.cleanup(res)


def test_cleanup_removes_the_working_folder(tmp_path):
    z = _zip(tmp_path, {"adv.pdf": ADVISORY_26})
    res = sort(str(z))
    assert Path(res.workdir).exists()
    intake.cleanup(res)
    assert not Path(res.workdir).exists()


# ------------------------------------------------ journals are identical

ACCTS = {"bank": "Assets:Bank:Current Account", "tds_expense": "Expenses:Tax:TDS",
         "remuneration_income": "Income:PGBP:Remuneration",
         "share_of_profit_income": "Income:PGBP:Share of Profit"}


def _entities(tmp_path):
    p = tmp_path / "entities.yaml"
    p.write_text(yaml.safe_dump({"SYN-LLP": {"name": "Synthetic Partner One", "pan": "AAAAA0000A",
                                             "status": "Individual",
                                             "partner_comp_accounts": dict(ACCTS)}}),
                 encoding="utf-8")
    return p


def _book(tmp_path):
    from tests.test_skill_partner_comp_recon import (
        _bank_deposit_txn, _gc_document_xml, _gc_tree_with_clearing, _write_gnucash_book)
    accounts, guids = _gc_tree_with_clearing()
    txns = [_bank_deposit_txn("dep-1", "2025-04-30", 480000.0,
                              guids["Current Account"], guids["Other Clearing"])]
    return _write_gnucash_book(tmp_path / "book.gnucash", _gc_document_xml(accounts, txns))


def _full_run(tmp_path, tag, **kw):
    out = tmp_path / f"{tag}.xlsx"
    jr = tmp_path / f"{tag}.csv"
    res = agent_module.run(entity="SYN-LLP", financial_year=FY, output_path=str(out),
                           journal_path=str(jr), entities_path=str(_entities(tmp_path)),
                           gnucash_path=_book(tmp_path), **kw)
    return str(res), out, jr


def test_journals_are_byte_identical_via_drop_zone_or_pickers(tmp_path):
    d = folder(tmp_path, {"one.pdf": payout_text("2025-04"), "two.pdf": ADVISORY_26}, "drop")
    pk = folder(tmp_path, {"one.pdf": payout_text("2025-04")}, "pick")
    adv = tmp_path / "pickadv.pdf"
    adv.write_text(ADVISORY_26, encoding="utf-8")

    res_a, out_a, jr_a = _full_run(tmp_path, "pickers", advices_dir=str(pk), advisory_path=str(adv))
    res_b, out_b, jr_b = _full_run(tmp_path, "dropzone", firm_documents=str(d))

    assert jr_a.exists() and jr_b.exists(), (res_a, res_b)
    assert jr_a.read_bytes() == jr_b.read_bytes()
    assert list(csv.reader(jr_a.open(encoding="utf-8")))  # not an empty comparison

    import openpyxl
    assert "Inputs" not in openpyxl.load_workbook(out_a).sheetnames   # NEGATIVE: pickers add no sheet
    wb = openpyxl.load_workbook(out_b)
    assert wb.sheetnames[0] == "Inputs"
    text = " ".join(str(c.value) for r in wb["Inputs"].iter_rows() for c in r if c.value)
    assert "one.pdf" in text and "two.pdf" in text
    assert res_b.startswith("**Inputs")


# ------------------------------------------------------------ the manifest

def _skill():
    return next(s for s in discover(refresh=True)
                if Path(s.manifest_path).parent.name == "skill_partner_comp_recon")


def test_manifest_declares_the_drop_zone_and_makes_the_pickers_optional():
    sk = _skill()
    by = {i.name: i for i in sk.inputs}
    assert by["firm_documents"].type == "files" and not by["firm_documents"].required
    assert {".pdf", ".eml", ".zip"} <= set(by["firm_documents"].file_types)
    assert not by["advices_dir"].required and not by["advisory_path"].required
    raw = yaml.safe_load(Path(sk.manifest_path).read_text(encoding="utf-8"))
    assert raw["run_args"]["firm_documents"] == "{inputs.firm_documents}"
    names = [i.name for i in sk.inputs]
    assert names.index("firm_documents") < names.index("advices_dir")
    h = next(x for x in sk.help.inputs if x.name == "firm_documents")
    assert "overriding" in h.tooltip and "CONTENT" in h.tooltip


def test_agent_still_enforces_the_required_roles_without_the_drop_zone():   # NEGATIVE
    out = agent_module.run(entity="X", advices_dir="", advisory_path="", financial_year=FY)
    assert out.startswith("ERROR") and "advices_dir" in out


# ------------------------------------------------- UI-21 drawn samples

from ui import _help  # noqa: E402

PAN_LIKE = re.compile(r"[A-Z]{5}[0-9]{4}[A-Z]")
LONG_NUMBER = re.compile(r"\b\d{9,18}\b")
SAMPLE_INPUTS = ("advices_dir", "advisory_path", "award_year_documents",
                 "llp_statement", "payment_schedule")


def _panel(name):
    hi = next(x for x in _skill().help.inputs if x.name == name)
    return hi, _help.sample_panel_markdown(hi)


@pytest.mark.parametrize("name", SAMPLE_INPUTS)
def test_every_document_picker_has_a_sample_with_the_required_parts(name):
    hi, body = _panel(name)
    assert hi.sample_html and "SAMPLE" in body and "made-up" in body
    assert "&larr;" in body                       # a tell-tale part is marked
    assert hi.looks_like and hi.not_these
    assert "Target compensation" in " ".join(hi.not_these) or name in (
        "advices_dir", "payment_schedule")
    assert "START year" in hi.filename_note


def test_advisory_panel_names_the_four_look_alikes():
    hi, body = _panel("advisory_path")
    joined = " ".join(hi.not_these).lower()
    for part in ("target compensation", "forward-looking", "older certificate", "superseded"):
        assert part in joined
    assert "(2)" in joined


@pytest.mark.parametrize("name", SAMPLE_INPUTS)
def test_no_sample_carries_a_pan_a_real_looking_account_or_a_real_name(name):   # NEGATIVE
    hi, body = _panel(name)
    text = re.sub(r"<[^>]+>", " ", body)
    assert not PAN_LIKE.search(body)
    assert not LONG_NUMBER.search(text)
    assert not re.search(r"https?://|<img|<script|data:image", body, re.I)   # drawn, never a scan
    names = re.findall(r"(?:Partner|Firm): ([A-Z][a-z]+ [A-Z][a-z]+(?: [A-Z][A-Za-z]+)*)", text)
    assert names and all(n.startswith("Example") for n in names)


def test_pan_guard_would_catch_a_pan():                                    # NEGATIVE (guard works)
    assert PAN_LIKE.search("Partner: ABCDE1234F")


YEAR = re.compile(r"\b20\d\d\b")


def test_no_skill_sample_or_its_help_text_carries_a_literal_year():        # NEGATIVE
    seen = 0
    for sk in discover():
        if not sk.help:
            continue
        for hi in sk.help.inputs:
            for field in (hi.sample_html, hi.looks_like, hi.filename_note, *hi.not_these):
                if field:
                    seen += 1
                    assert not YEAR.search(field), (sk.name, hi.name, YEAR.search(field).group(0))
    assert seen >= 15


@pytest.mark.parametrize("name", SAMPLE_INPUTS)
def test_partner_sample_keeps_the_sample_line_and_the_cy_py_legend(name):
    hi, _ = _panel(name)
    assert "SAMPLE" in hi.sample_html
    assert "CY = the financial year you are reconciling; PY = the year before it." in hi.sample_html
    assert hi.sample_html.index("SAMPLE") < hi.sample_html.index("CY = ")


def test_year_guard_would_catch_a_literal_year():                           # NEGATIVE (guard works)
    assert YEAR.search("Year ended 31 March 2026")
    assert not YEAR.search("Year ended 31 March [end of CY]")


def test_input_without_a_sample_still_renders():                           # NEGATIVE
    import gradio as gr
    sk = _skill()
    with gr.Blocks():
        assert _help.mount_sample_panel(sk, "entity") is False       # no sample: draws nothing
        assert _help.mount_sample_panel(sk, "no_such_input") is False
        assert _help.mount_sample_panel(sk, "advisory_path") is True


def test_samples_stay_under_the_manual_pickers_in_the_form():
    src = (ROOT / "ui" / "tabs" / "_generic.py").read_text(encoding="utf-8")
    loop_end = src.index("input_by_name[inp.name] = comp")
    assert "_help.mount_sample_panel(skill, inp.name)" in src[loop_end:loop_end + 400]
    # the drop zone itself carries no sample: its help says the pickers are overrides
    hi = next(x for x in _skill().help.inputs if x.name == "firm_documents")
    assert not hi.sample_html


def test_skill_without_a_help_block_renders_no_panel():                    # NEGATIVE
    import copy
    sk = copy.copy(_skill())
    object.__setattr__(sk, "help", None) if getattr(type(sk), "__dataclass_params__", None) and \
        type(sk).__dataclass_params__.frozen else setattr(sk, "help", None)
    assert _help.mount_sample_panel(sk, "advisory_path") is False


# ------------------------------------------------- "Check my files" button

from ui.tabs import _generic  # noqa: E402


def _check_inputs(d, **kw):
    base = {"firm_documents": [str(p) for p in sorted(Path(d).iterdir())], "fy": FY}
    base.update(kw)
    return base


def _tree(root):
    return sorted(str(p) for p in Path(root).rglob("*"))


def _table(text):
    return [ln for ln in text.splitlines() if ln.startswith("|")]


def test_check_shows_the_inputs_table_and_the_missing_lines(tmp_path):
    d = folder(tmp_path, {"p.pdf": payout_text("2025-04"), "t.pdf": TARGET})
    out = agent_module.check_documents(_check_inputs(d))
    assert out.startswith("**Inputs")
    assert "| p.pdf |" in out and "Target compensation letter" in out
    assert "Still missing: this year's (FY 2025-26) Advisory." in out


def test_check_writes_nothing_and_never_runs_the_reconciliation(tmp_path, monkeypatch):   # NEGATIVE
    d = folder(tmp_path, full_set())
    out_dir = tmp_path / "outputs"
    out_dir.mkdir()
    monkeypatch.setattr(_generic._config, "output_dir", lambda: out_dir, raising=False)
    boom = lambda *a, **k: (_ for _ in ()).throw(AssertionError("must not run"))  # noqa: E731
    monkeypatch.setattr(agent_module, "_run_from_documents", boom)
    monkeypatch.setattr(agent_module, "write_report_workbook", boom, raising=False)
    monkeypatch.setattr(agent_module, "write_journal_csv", boom, raising=False)
    before = _tree(tmp_path)
    out = agent_module.check_documents(_check_inputs(d))
    assert "ERROR" not in out
    assert _tree(tmp_path) == before                      # no workbook, no journal, nothing in outputs
    assert list(out_dir.iterdir()) == []
    assert list(tmp_path.glob("**/*.csv")) == [] and list(tmp_path.glob("**/*.xlsx")) == []


def test_check_cleans_up_its_temp_folder(tmp_path, monkeypatch):
    d = folder(tmp_path, full_set())
    made = []
    real = intake.cleanup
    monkeypatch.setattr(intake, "cleanup", lambda r: (made.append(r.workdir), real(r)))
    agent_module.check_documents(_check_inputs(d))
    assert made and not Path(made[0]).exists()


def test_check_never_opens_the_gnucash_book(tmp_path, monkeypatch):        # NEGATIVE
    import builtins
    import gzip
    book = tmp_path / "book.gnucash"
    book.write_bytes(b"synthetic")
    d = folder(tmp_path, full_set())
    opened = []
    real_open, real_gz = builtins.open, gzip.open
    monkeypatch.setattr(builtins, "open", lambda f, *a, **k: (opened.append(str(f)), real_open(f, *a, **k))[1])
    monkeypatch.setattr(gzip, "open", lambda f, *a, **k: (opened.append(str(f)), real_gz(f, *a, **k))[1])
    agent_module.check_documents(_check_inputs(d, gnucash_path=str(book)))
    assert not any(".gnucash" in o for o in opened)


def test_check_with_no_documents_says_so_plainly(tmp_path):
    out = agent_module.check_documents({"firm_documents": "", "fy": FY})
    assert out.startswith("Nothing to check")


def test_check_with_a_wrong_or_missing_password_gives_the_plain_message(tmp_path, monkeypatch):   # NEGATIVE
    d = folder(tmp_path, {"locked.pdf": "x"})

    def locked(path, password=None):
        return None
    monkeypatch.setattr(intake._precheck, "read_first_page_text", locked)
    monkeypatch.setattr(intake._precheck, "pdf_open_problem",
                        lambda p, pw: (intake._precheck.PASSWORD_WRONG_MESSAGE if pw
                                       else intake._precheck.PASSWORD_MISSING_MESSAGE))
    for pw, msg in (("", intake._precheck.PASSWORD_MISSING_MESSAGE),
                    ("wrong", intake._precheck.PASSWORD_WRONG_MESSAGE)):
        out = agent_module.check_documents(_check_inputs(d, doc_password=pw))
        assert msg in out and "Traceback" not in out
        assert "wrong" not in out.replace("did not open it", "") or pw == ""   # the password is never echoed


def test_check_survives_an_unexpected_failure_without_a_traceback(tmp_path, monkeypatch):   # NEGATIVE
    d = folder(tmp_path, full_set())
    monkeypatch.setattr(intake, "classify", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("synthetic")))
    out = agent_module.check_documents(_check_inputs(d))
    assert out.startswith("ERROR") and "Traceback" not in out


def test_check_table_and_run_table_agree_for_the_same_inputs(tmp_path, monkeypatch):
    d = folder(tmp_path, {**full_set(), "t.pdf": TARGET, "m.pdf": "menu\n"})
    monkeypatch.setattr(agent_module, "_run_from_documents", lambda **kw: "ok")
    run_out = agent_module.run(entity="X", firm_documents=str(d), financial_year=FY,
                               output_path=str(tmp_path / "o.xlsx"))
    check_out = agent_module.check_documents({"firm_documents": str(d), "fy": FY})
    assert _table(run_out) == _table(check_out) and _table(check_out)
    assert [ln for ln in run_out.splitlines() if "Still missing" in ln] == \
        [ln for ln in check_out.splitlines() if "Still missing" in ln]


def test_check_honours_manual_picks_like_the_run_does(tmp_path):
    d = folder(tmp_path, full_set())
    out = agent_module.check_documents(_check_inputs(d, advisory_path=str(tmp_path / "x.pdf")))
    assert "a manual pick always wins" in out


def test_run_check_handler_returns_text_and_names_no_other_output(tmp_path):
    d = folder(tmp_path, full_set())
    sk = _skill()
    vals = []
    for inp in sk.inputs:
        if inp.name == "firm_documents":
            vals.append([str(p) for p in sorted(d.iterdir())])
        elif inp.name == "fy":
            vals.append(FY)
        else:
            vals.append(None)
    out = _generic.run_check(sk, vals)
    assert out.startswith("**Inputs")


def _button_labels(skill):
    import gradio as gr
    with gr.Blocks() as demo:
        _generic.render(skill)
    return [getattr(b, "value", None) for b in demo.blocks.values() if isinstance(b, gr.Button)], demo


def test_the_partner_form_has_the_button_and_wires_only_the_result_text():
    labels, demo = _button_labels(_skill())
    assert "Check my files" in labels
    btn = next(b for b in demo.blocks.values() if getattr(b, "value", None) == "Check my files")
    dep = next(d for d in demo.fns.values() if any(t[0] == btn._id for t in d.targets))
    out_ids = [c._id for c in dep.outputs]
    assert len(out_ids) == 1                       # the result text only: no download, no path box


def test_a_skill_with_no_check_handler_shows_no_button():                  # NEGATIVE
    others = [s for s in discover(refresh=True) if s.check is None]
    assert others
    for s in others[:3]:
        labels, _ = _button_labels(s)
        assert "Check my files" not in labels
