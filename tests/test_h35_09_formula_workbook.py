"""
H35-09 -- the partner reconciliation workbook is inputs plus live formulas.

Recalculation approach: the repo has no spreadsheet engine (no LibreOffice, no
formula library). Like the ITR workbook tests (tests/skill_itr_workbook/
test_presentation.py, `_mini_eval`), this file carries a small pure-Python
evaluator, here covering exactly the grammar the writer emits: + - * /,
unary minus, cell and range references (same-sheet and cross-sheet) and the
functions SUM, ROUND, ABS and MAX. It recalculates EVERY formula cell in the
written workbook and compares it with the figure Python computed.

Synthetic data only.
"""
from __future__ import annotations

import copy
import dataclasses
import datetime as _dt
import json
import math
import re
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

ROOT = Path(__file__).resolve().parent.parent
for _p in (str(ROOT), str(ROOT / "src")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from agents.skill_partner_comp_recon import writer  # noqa: E402
from agents.skill_partner_comp_recon.engine import build_report  # noqa: E402
from agents.skill_partner_comp_recon.jv_emitter import build_journals  # noqa: E402

FIXTURE = ROOT / "tests" / "fixtures" / "partner_comp_recon_fy2025_26.yaml"
SNAPSHOT = ROOT / "tests" / "fixtures" / "partner_comp_recon_verdict_sheets_before_h35_09.json"

COMPUTED_SHEETS = ["Monthly grid", "Payroll stream", "One-offs", "Cohorts", "Capital",
                   "Interest on capital", "CTC check"]
VERDICT_SHEETS = ["Reconciliation", "Exceptions", "Open items", "Logic"]


# ----------------------------------------------------------------- evaluator

class _Book:
    def __init__(self, wb):
        self.wb = wb
        self.cache = {}

    def cell(self, sheet, ref):
        key = (sheet, ref.replace("$", ""))
        if key in self.cache:
            return self.cache[key]
        v = self.wb[sheet][key[1]].value
        if isinstance(v, str) and v.startswith("="):
            v = self._eval(sheet, v[1:])
        elif isinstance(v, (_dt.datetime, _dt.date)):
            v = (v.date() if isinstance(v, _dt.datetime) else v).toordinal()
        self.cache[key] = v
        return v

    def num(self, sheet, ref):
        v = self.cell(sheet, ref)
        if v is None:
            return 0.0
        if isinstance(v, str):
            raise ValueError(f"text in arithmetic: {sheet}!{ref} = {v!r}")
        return float(v)

    def rng(self, sheet, a, b):
        a, b = a.replace("$", ""), b.replace("$", "")
        ma, mb = re.match(r"([A-Z]+)(\d+)", a), re.match(r"([A-Z]+)(\d+)", b)
        out = []
        for r in range(int(ma.group(2)), int(mb.group(2)) + 1):
            for c in range(column_index_from_string(ma.group(1)),
                           column_index_from_string(mb.group(1)) + 1):
                v = self.cell(sheet, f"{get_column_letter(c)}{r}")
                if isinstance(v, (int, float)):       # SUM ignores text and blanks
                    out.append(float(v))
        return out

    def _eval(self, sheet, expr):
        sh = r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_]*))!"
        ref = r"\$?[A-Z]{1,3}\$?\d+"

        def rng_sub(m):
            s = m.group(1) or m.group(2) or sheet
            return f"_R({s!r},{m.group(3)!r},{m.group(4)!r})"
        expr = re.sub(rf"(?:{sh})?({ref}):({ref})", rng_sub, expr)

        def ref_sub(m):
            s = m.group(1) or m.group(2) or sheet
            return f"_C({s!r},{m.group(3)!r})"
        expr = re.sub(rf"(?<![A-Za-z_'\"])(?:{sh})?({ref})(?![\d(A-Za-z'])", ref_sub, expr)

        def xl_round(x, n=0):
            q = 10 ** int(n)
            return math.copysign(math.floor(abs(x) * q + 0.5) / q, x)

        env = {"_R": self.rng, "_C": self.num, "SUM": lambda *a: sum(
                   sum(x) if isinstance(x, list) else x for x in a),
               "ROUND": xl_round, "ABS": abs,
               "MAX": lambda *a: max(v for x in a for v in (x if isinstance(x, list) else [x])),
               "__builtins__": {}}
        return float(eval(expr, env))   # noqa: S307 -- own generated formulas only


def _formula_cells(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    yield ws.title, c


def _data(**over):
    d = yaml.safe_load(FIXTURE.read_text(encoding="utf-8"))
    d.setdefault("drivers", {})
    d["drivers"].setdefault("capital_interest_rate", 0.06)
    d["ctc_structuring"] = {"total": 1500000.0}
    d["monthly"][1]["capital_transferred"] = -200000.0     # two capital tranches
    d["monthly"][3]["capital_transferred"] = -100000.0
    d["drivers"]["capital_interest_from_date_overrides"] = {d["monthly"][3]["month"]: "2025-07-15"}
    for k, v in over.items():
        d[k] = v
    return d


def _write(tmp_path, data, name="w.xlsx"):
    report = build_report(data)
    out = tmp_path / name
    writer.write_report_workbook(report, str(out))
    return report, out


def _approx(a, b):
    return abs(a - b) <= 0.005


# ----------------------------------------------------- the required recalc test

def test_every_formula_cell_recalculates_to_the_python_figure(tmp_path):
    report, out = _write(tmp_path, _data())
    wb = load_workbook(out)
    book = _Book(wb)
    cells = list(_formula_cells(wb))
    assert len(cells) > 50, "workbook should be formula-driven"
    expected = {}                                    # (sheet, coord) -> python figure

    ws = wb["Monthly grid"]
    labels = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    attrs = [a for a, _l in writer._MONTHLY_LINE_LABELS]
    for r, attr in enumerate(attrs, start=2):
        total = 0.0
        for c, m in enumerate(report.monthly, start=2):
            v = getattr(m, attr)
            expected[("Monthly grid", f"{get_column_letter(c)}{r}")] = v
            total += v
        expected[("Monthly grid", f"{get_column_letter(2 + len(report.monthly))}{r}")] = total
    assert len(labels) == len(attrs)

    ws = wb["One-offs"]
    assert report.one_offs, "fixture must exercise a one-off"
    for i, o in enumerate(report.one_offs, start=2):
        expected[("One-offs", f"A{i}")] = o.net
        expected[("One-offs", f"B{i}")] = o.firms_tax_rate
        expected[("One-offs", f"C{i}")] = o.gross
        expected[("One-offs", f"D{i}")] = o.roundness

    ws = wb["Cohorts"]
    assert report.cohort_instalments
    for i, inst in enumerate(report.cohort_instalments, start=2):
        for col, v in zip("EFGH", (inst.gross, inst.firms_tax, inst.capital, inst.net)):
            if v is not None:
                expected[("Cohorts", f"{col}{i}")] = v

    sched = report.capital_interest_schedule
    ws = wb["Interest on capital"]
    assert sched.rows
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value in [x.month for x in sched.rows] and \
                isinstance(ws.cell(row=r, column=6).value, str):
            row = sched.rows[[x.month for x in sched.rows].index(ws.cell(row=r, column=1).value)]
            expected[("Interest on capital", f"B{r}")] = row.principal
            expected[("Interest on capital", f"E{r}")] = row.days
            expected[("Interest on capital", f"F{r}")] = row.interest
        if ws.cell(row=r, column=1).value == "Total computed interest":
            expected[("Interest on capital", f"B{r}")] = sched.total_interest

    ctc = report.ctc_check
    ws = wb["CTC check"]
    by_label = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
    for label, v in [("Target Compensation", ctc.target_compensation),
                     ("Remuneration (total)", ctc.remuneration_total),
                     ("Gross share of profit (total)", ctc.gross_sop_total),
                     ("CTC structuring (total)", ctc.ctc_structuring_total),
                     ("Arrears (additional share of profit, total)", ctc.arrears_total),
                     ("Cash pool actually paid", ctc.cash_pool),
                     ("Gap (Target Compensation - cash pool)", ctc.gap),
                     ("Firm's tax on pool (with arrears in scope)", ctc.firms_tax_on_pool)]:
        expected[("CTC check", f"B{by_label[label]}")] = v

    ws = wb["Capital"]
    assert report.capital_rule.status == "OK"
    expected[("Capital", "B2")] = report.capital_rule.required_cumulative_capital

    # every formula must evaluate, and wherever Python has a figure it must match
    checked = 0
    for sheet, c in cells:
        got = book.num(sheet, c.coordinate)            # raises if it cannot evaluate
        key = (sheet, c.coordinate)
        if key in expected:
            assert expected[key] is not None
            assert _approx(got, expected[key]), (key, got, expected[key], c.value)
            checked += 1
    # every figure we expected to be live really is a formula cell
    formula_keys = {(s, c.coordinate) for s, c in cells}
    missing = [k for k in expected if expected[k] is not None and k not in formula_keys]
    assert not missing, f"expected live formulas, found values: {missing[:5]}"
    assert checked >= 100


def test_payroll_totals_are_live_and_match(tmp_path):
    data = _data()
    data["payroll"] = [
        {"month": "2025-04", "gross_salary": 100000, "tds": -10000, "net_paid": 90000},
        {"month": "2025-05", "gross_salary": 120000, "tds": -12000, "net_paid": 108000},
    ]
    report, out = _write(tmp_path, data)
    wb = load_workbook(out)
    ws = wb["Payroll stream"]
    assert str(ws["B4"].value).startswith("=SUM(")
    book = _Book(wb)
    assert book.num("Payroll stream", "B4") == 220000
    assert book.num("Payroll stream", "D4") == 198000


# ------------------------------------------------------------- flow-through

def test_changed_input_flows_through_to_every_dependent_sheet(tmp_path):
    report, out = _write(tmp_path, _data())
    wb = load_workbook(out)
    before = _Book(wb)
    inp = wb["Inputs - monthly"]
    r_rem = writer._MONTHLY_INPUT_ROW["remuneration"]
    r_paid = writer._MONTHLY_INPUT_ROW["total_paid"]
    r_cap = writer._MONTHLY_INPUT_ROW["capital_transferred"]
    tot = get_column_letter(2 + len(report.monthly))
    grid_total_rem = before.num("Monthly grid", f"{tot}2")
    cash0 = None
    ws = wb["CTC check"]
    by_label = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
    cash_cell = f"B{by_label['Cash pool actually paid']}"
    cash0 = before.num("CTC check", cash_cell)
    misc0 = before.num("Monthly grid", "B10")

    tranche_month = report.capital_interest_schedule.rows[0].month
    tcol = get_column_letter(2 + [m.month for m in report.monthly].index(tranche_month))
    i_ws = wb["Interest on capital"]
    row = next(r for r in range(1, i_ws.max_row + 1) if i_ws.cell(row=r, column=1).value == tranche_month)
    princ0 = before.num("Interest on capital", f"B{row}")
    int0 = before.num("Interest on capital", f"F{row}")
    other0 = before.num("Monthly grid", "C2")

    inp[f"B{r_rem}"].value = inp[f"B{r_rem}"].value + 1000        # a changed input
    inp[f"{tcol}{r_cap}"].value = inp[f"{tcol}{r_cap}"].value - 50000   # a bigger capital tranche
    after = _Book(wb)                                              # fresh cache
    assert after.num("Monthly grid", f"{tot}2") == grid_total_rem + 1000
    assert after.num("CTC check", cash_cell) == cash0 + 1000
    assert after.num("Monthly grid", "B10") == misc0 - 1000        # misc derives live
    assert after.num("Interest on capital", f"B{row}") == princ0 + 50000
    assert after.num("Interest on capital", f"F{row}") > int0
    assert after.num("Monthly grid", "C2") == other0           # an untouched input is unchanged
    _ = r_paid


def test_changed_driver_rate_flows_into_the_one_off_gross_up(tmp_path):
    report, out = _write(tmp_path, _data())
    wb = load_workbook(out)
    d = wb["Drivers"]
    rate_row = next(r for r in range(1, d.max_row + 1) if d.cell(row=r, column=1).value == "Firm's tax rate")
    old = _Book(wb).num("One-offs", "C2")
    d.cell(row=rate_row, column=2).value = d.cell(row=rate_row, column=2).value / 2
    assert _Book(wb).num("One-offs", "C2") < old


# ------------------------------------------------- no hard-coded numbers

def test_computed_sheets_hold_no_hard_coded_figure(tmp_path):
    """A number typed into a computed sheet is a figure that would not move
    when its input changes. Only the cohort-less / payroll leaf rows, and
    explicitly unsupplied placeholders, may be literal."""
    report, out = _write(tmp_path, _data())
    wb = load_workbook(out)
    offenders = []
    for name in COMPUTED_SHEETS:
        if name not in wb.sheetnames or name == "Payroll stream":
            continue
        stop = None
        if name == "Capital":
            # below this header the sheet only mirrors the Reconciliation
            # sheet's own source figures (verdict side, kept in Python)
            stop = next(r for r in range(1, wb[name].max_row + 1)
                        if str(wb[name].cell(row=r, column=1).value).startswith("Cross-check"))
        for row in wb[name].iter_rows():
            for c in row:
                if stop is not None and c.row >= stop:
                    continue
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    offenders.append((name, c.coordinate, c.value))
    assert offenders == [], offenders[:10]


def test_input_sheets_carry_a_source_for_every_row(tmp_path):
    _report, out = _write(tmp_path, _data())
    wb = load_workbook(out)
    ws = wb["Inputs - monthly"]
    src_col = ws.max_column
    assert ws.cell(row=1, column=src_col).value == "Source"
    for r in range(2, ws.max_row + 1):
        assert str(ws.cell(row=r, column=src_col).value or "").strip(), r


# ------------------------------------------------- verdicts / journals unchanged

def test_verdict_sheets_equal_the_pre_h35_09_output(tmp_path):
    snap = json.loads(SNAPSHOT.read_text(encoding="utf-8"))
    _report, out = _write(tmp_path, yaml.safe_load(FIXTURE.read_text(encoding="utf-8")))
    wb = load_workbook(out)
    for name in VERDICT_SHEETS:
        rows = [[("" if v is None else v) for v in row] for row in wb[name].iter_rows(values_only=True)]
        want = [[("" if v is None else v) for v in row] for row in snap[name]]
        assert rows == want, name


def test_writing_the_workbook_never_changes_the_report_or_journals(tmp_path):
    data = yaml.safe_load(FIXTURE.read_text(encoding="utf-8"))
    report = build_report(copy.deepcopy(data))
    before_report = repr(dataclasses.asdict(report))
    before_j = repr([dataclasses.asdict(j) if dataclasses.is_dataclass(j) else j
                     for j in build_journals(report, data["accounts"])])
    writer.write_report_workbook(report, str(tmp_path / "x.xlsx"))
    assert repr(dataclasses.asdict(report)) == before_report
    after_j = repr([dataclasses.asdict(j) if dataclasses.is_dataclass(j) else j
                    for j in build_journals(report, data["accounts"])])
    assert after_j == before_j


# ------------------------------------------------- the old defect, directly

def test_text_labels_starting_with_equals_are_still_neutralised():              # NEGATIVE
    assert writer._text("= Gross incentive") == " = Gross incentive"
    assert not isinstance(writer._text("=SUM(A1)"), writer._F)
    assert writer._text(writer._F("=SUM(A1:A2)")) == "=SUM(A1:A2)"


def test_formulas_are_stored_as_formulas_not_text(tmp_path):                    # NEGATIVE
    """Before H35-09 every formula was stored as ' =SUM(...)' text."""
    _report, out = _write(tmp_path, _data())
    wb = load_workbook(out)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    assert not c.value.startswith(" ="), (ws.title, c.coordinate)
    tot = get_column_letter(2 + len(_report.monthly))
    assert wb["Monthly grid"][f"{tot}2"].value == f"=SUM(B2:{get_column_letter(1 + len(_report.monthly))}2)"


def test_missing_rate_still_shows_cannot_reconcile_not_a_formula(tmp_path):     # NEGATIVE
    data = _data()
    data["drivers"].pop("capital_interest_rate", None)
    data["drivers"]["firms_tax_rate"] = None
    report, out = _write(tmp_path, data)
    wb = load_workbook(out)
    # no formula may reference an unsupplied rate
    for sheet, c in _formula_cells(wb):
        if sheet == "One-offs" and c.column > 1:        # col A is the live net link
            pytest.fail(f"one-off formula written with no firm's tax rate: {c.coordinate}")
    assert "CANNOT RECONCILE" in " ".join(str(c.value) for row in wb["One-offs"].iter_rows() for c in row)
    ws = wb["Interest on capital"]
    assert "not supplied" in " ".join(str(c.value) for row in ws.iter_rows() for c in row)
