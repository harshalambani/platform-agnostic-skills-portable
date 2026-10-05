"""
TDS bundle (TDS-01..08) -- 26AS journal builder.

TDS-01 pins the behaviour that is correct today (it passes on the code before
any TDS fix), so the fixes below cannot quietly regress it. Each later section
carries its own negative tests.

Synthetic accounts and deductors only.
"""
from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
SCRIPT = SRC / "agents" / "skill_26as_journal" / "scripts" / "build_tds_journals.py"


def _load():
    if str(SRC) not in sys.path:
        sys.path.insert(0, str(SRC))
    spec = importlib.util.spec_from_file_location("build_tds_journals", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


m = _load()

BOB = "Income:Interest Income:Interest on BOB - FD"
ICICI = "Income:Interest Income:Interest on ICICI Bank - FD"
GENERIC_FD = "Income:Interest Income:Interest on FD"


def acc(path, typ="INCOME", blocked=False):
    return m.Account(path=path, leaf=path.split(":")[-1], type=typ,
                     special=blocked, blocked=blocked)


def chart(*extra):
    base = [acc(BOB), acc(ICICI), acc(GENERIC_FD),
            acc("Expense:TDS on Interest", "EXPENSE"),
            acc("Expense:TDS on Dividend", "EXPENSE"),
            acc("Liabilities:Suspense", "LIABILITY")]
    return base + list(extra)


def ded(sr, name, section, paid, tax):
    return m.Deductor(sr=sr, name=name, sections=(section,), amount_paid=paid,
                      tax_deducted=tax, tds_deposited=tax)


# ---- TDS-01: pinned, currently correct --------------------------------------------------

def test_tds01_b2a_bank_of_baroda_matches_the_bob_account():
    acct, conf, _b, _c, tied = m.match_credit_account("BANK OF BARODA", "A", chart(), GENERIC_FD)
    assert acct == BOB and tied == []


def test_tds01_b2b_a_bank_with_its_own_account_does_not_land_on_bob():
    acct, _conf, _b, _c, _t = m.match_credit_account("ICICI BANK LIMITED", "A", chart(), GENERIC_FD)
    assert acct == ICICI
    assert acct != BOB


def test_tds01_b3a_an_override_changes_only_its_own_row():
    ds = [ded(1, "BANK OF BARODA", "194A", 1000.0, 100.0),
          ded(2, "ICICI BANK LIMITED", "194A", 2000.0, 200.0)]
    js = m.build_journals(ds, chart(), overrides={2: "Liabilities:Suspense"})
    assert js[0].credit_account == BOB and js[0].credit_confidence != "Override"
    assert js[1].credit_account == "Liabilities:Suspense" and js[1].credit_confidence == "Override"


def test_tds01_b3c_an_override_always_stays_for_review_and_balanced():
    j = m.build_journals([ded(1, "BANK OF BARODA", "194A", 1000.0, 100.0)], chart(),
                         overrides={1: ICICI})[0]
    assert j.needs_review and j.balanced
    assert j.credit_basis == "Model pick - confirm"


def test_tds01_b4a_the_generic_fd_account_is_found_and_used_as_the_second_debit():
    assert m.find_generic_fd_account(chart()) == GENERIC_FD
    j = m.build_journals([ded(1, "BANK OF BARODA", "194A", 1000.0, 100.0)], chart())[0]
    assert {s.account for s in j.splits if s.debit} == {"Expense:TDS on Interest", GENERIC_FD}
    assert j.balanced


def test_tds01_b4d_a_bank_specific_fd_account_is_never_the_generic_one():
    only_specific = [acc(BOB), acc(ICICI)]
    assert m.find_generic_fd_account(only_specific) is None
    assert m.find_generic_fd_account(chart()) not in (BOB, ICICI)


# ---- TDS-02: BANK alone never ties a payer to BoB ---------------------------------------

import pytest  # noqa: E402


@pytest.mark.parametrize("payer", ["HDFC BANK LIMITED", "ICICI BANK", "STATE BANK OF INDIA",
                                   "SOME OTHER BANK LTD", "AXIS BANK"])
def test_tds02_other_banks_do_not_resolve_to_bob(payer):
    two = [acc(BOB), acc("Income:Interest Income:Interest on Chola Bond")]
    acct, conf, _b, _c, _t = m.match_credit_account(payer, "A", two, "")
    assert acct != BOB                    # NEGATIVE: the word BANK is not a match
    assert acct is None and conf == "Suspense"


def test_tds02_bank_of_baroda_still_resolves_to_bob():
    acct, _c, _b, _cs, _t = m.match_credit_account(
        "BANK OF BARODA", "A", [acc(BOB), acc(ICICI)], "")
    assert acct == BOB


def test_tds02_the_alias_table_does_not_carry_the_bare_word_bank():
    assert "BANK" not in m.ALIASES["BOB"]


# ---- TDS-03: the generic FD account is recognised under its real names ------------------

@pytest.mark.parametrize("leaf", ["Interest on Fixed Deposits", "Interest on FDs",
                                  "Term Deposit", "Deposits", "Interest on Term Deposits",
                                  "Interest on Fixed Deposit", "Interest on FD", "FD Interest"])
def test_tds03_fd_account_name_variants_are_recognised(leaf):
    path = "Income:Interest Income:" + leaf
    got = m.find_generic_fd_account([acc(path), acc(BOB)])
    assert got == path


@pytest.mark.parametrize("leaf", ["Interest on Savings", "Interest from Savings Account",
                                  "Interest on Bonds", "Interest on Recurring Deposit",
                                  "Interest on Income Tax Refund"])
def test_tds03_a_non_fd_interest_account_is_not_picked(leaf):
    path = "Income:Interest Income:" + leaf
    assert m.find_generic_fd_account([acc(path), acc(BOB)]) is None   # NEGATIVE


def test_tds03_a_bank_specific_fd_account_is_still_not_generic():
    assert m.find_generic_fd_account([acc("Income:Interest Income:Interest on HDFC Fixed Deposits")]) is None


def test_tds03_a_blocked_fd_account_is_still_not_picked():
    assert m.find_generic_fd_account(
        [acc("Income:Interest Income:Interest on Fixed Deposits", blocked=True)]) is None


# ---- shared helpers: a tiny synthetic book + 26AS workbook for run()-level tests --------

import csv  # noqa: E402

_NS = ('xmlns:gnc="http://www.gnucash.org/XML/gnc" xmlns:act="http://www.gnucash.org/XML/act" '
       'xmlns:slot="http://www.gnucash.org/XML/slot"')


def write_book(path, specs):
    """specs: [(full path, type, flags)]; flags is a set of 'placeholder'/'hidden'.
    Parents are created implicitly (as plain accounts) when not listed."""
    nodes = {}
    for full, typ, flags in specs:
        parts = full.split(":")
        for i in range(1, len(parts) + 1):
            p = ":".join(parts[:i])
            if p not in nodes:
                nodes[p] = [parts[i - 1], "id%d" % (len(nodes) + 1), "INCOME" if i < len(parts) else typ,
                            ":".join(parts[:i - 1]) or None, set()]
        nodes[full][2], nodes[full][4] = typ, set(flags)
    out = ["<gnc-v2 %s><gnc:book>" % _NS]
    out.append('<gnc:account><act:name>Root Account</act:name><act:id type="guid">root</act:id>'
               '<act:type>ROOT</act:type></gnc:account>')
    for full, (name, aid, typ, parent, flags) in nodes.items():
        par = nodes[parent][1] if parent else "root"
        slots = "".join("<slot><slot:key>%s</slot:key><slot:value type=\"string\">true</slot:value></slot>" % f
                        for f in sorted(flags))
        out.append('<gnc:account><act:name>%s</act:name><act:id type="guid">%s</act:id>'
                   '<act:type>%s</act:type><act:parent type="guid">%s</act:parent>%s</gnc:account>'
                   % (name, aid, typ, par, "<act:slots>%s</act:slots>" % slots if slots else ""))
    out.append("</gnc:book></gnc-v2>")
    Path(path).write_text("".join(out), encoding="utf-8")
    return Path(path)


def write_26as(path, parts):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in parts.items():
        ws = wb.create_sheet(title=title)
        ws.cell(1, 1, f"{title} - Details")
        ws.cell(2, 1, "Assessee Name: X  |  PAN: AAAAA1111A  |  Financial Year: 2025-26")
        ws.cell(3, 1, "Sr.No.")
        for r, (sr, name, section, amt, tax) in enumerate(rows, start=4):
            ws.cell(r, 1, sr)
            ws.cell(r, 2, name)
            ws.cell(r, 4, amt)
            ws.cell(r, 5, tax)
            ws.cell(r, 6, tax)
            ws.cell(r, 8, section)
    wb.save(path)
    return Path(path)


STD = [("Expense:TDS on Interest", "EXPENSE", ()), ("Expense:TDS on Dividend", "EXPENSE", ()),
       ("Liabilities:Suspense", "LIABILITY", ())]


def read_csv(path):
    with Path(path).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# ---- TDS-04: a placeholder/hidden FD account never silently absorbs the gap -------------

def test_tds04_placeholder_canonical_fd_account_is_never_posted_to():
    placeholder = chart()
    placeholder = [a for a in placeholder if a.path != GENERIC_FD] + \
        [acc(m.ACC_INTEREST_ON_FD, blocked=True)]
    j = m.build_journals([ded(1, "BANK OF BARODA", "194A", 1000.0, 100.0)], placeholder,
                         )[0]
    assert m.ACC_INTEREST_ON_FD not in {s.account for s in j.splits}   # NEGATIVE
    assert j.balanced
    assert j.needs_review                                              # loud, not silent
    assert j.blocked_targets == [m.ACC_INTEREST_ON_FD]
    assert m.BLOCKED_MARK in j.credit_basis
    assert "Liabilities:Suspense" in {s.account for s in j.splits if s.debit}


def test_tds04_hidden_fd_account_is_never_posted_to():
    accts = [a for a in chart() if a.path != GENERIC_FD] + \
        [acc("Income:Interest Income:Interest on Fixed Deposit", blocked=True)]
    j = m.build_journals([ded(1, "BANK OF BARODA", "194A", 1000.0, 100.0)], accts)[0]
    assert "Income:Interest Income:Interest on Fixed Deposit" not in {s.account for s in j.splits}
    assert j.balanced
    # the hidden account is not the canonical name, so the canonical (absent)
    # account is used and run() lists it under accounts to create (existing behaviour)
    assert m.ACC_INTEREST_ON_FD in {s.account for s in j.splits}


def test_tds04_a_postable_fd_account_is_still_used_and_nothing_is_flagged():
    j = m.build_journals([ded(1, "BANK OF BARODA", "194A", 1000.0, 100.0)], chart())[0]
    assert GENERIC_FD in {s.account for s in j.splits}
    assert j.blocked_targets == [] and m.BLOCKED_MARK not in j.credit_basis   # NEGATIVE


def test_tds04_no_journal_of_any_part_posts_to_a_blocked_account():
    blocked_tds = acc("Expense:TDS on Interest", "EXPENSE", blocked=True)
    accts = [a for a in chart() if a.path != "Expense:TDS on Interest"] + [blocked_tds]
    js = m.build_journals([ded(1, "BANK OF BARODA", "194A", 1000.0, 100.0)], accts)
    js += m.build_15g_journals([ded(1, "ICICI BANK LIMITED", "194A", 1000.0, 100.0)], accts)
    for j in js:
        assert "Expense:TDS on Interest" not in {s.account for s in j.splits}
        assert j.needs_review


def test_tds04_run_reports_the_blocked_account_as_missing_and_review_says_so(tmp_path):
    specs = STD + [("Income:Interest Income:Interest on BOB - FD", "INCOME", ()),
                   ("Income:Interest Income:Interest on ICICI Bank - FD", "INCOME", ()),
                   (m.ACC_INTEREST_ON_FD, "INCOME", ("placeholder",))]
    book = write_book(tmp_path / "b.gnucash", specs)
    x = write_26as(tmp_path / "x.xlsx", {"Part I": [(1, "BANK OF BARODA", "194A", 1000.0, 100.0)]})
    stats = m.run(x, book, tmp_path / "out.csv")
    assert m.ACC_INTEREST_ON_FD in stats["missing_accounts"]            # NEGATIVE: was []
    assert stats["blocked_accounts"] == [m.ACC_INTEREST_ON_FD]
    assert stats["needs_review"]
    assert all(r["Account"] != m.ACC_INTEREST_ON_FD for r in read_csv(tmp_path / "out.csv"))


def test_tds04_final_summary_names_the_blocked_rows(tmp_path):
    from agents.skill_26as_journal import tools as tl
    specs = STD + [("Income:Interest Income:Interest on BOB - FD", "INCOME", ()),
                   ("Income:Interest Income:Interest on ICICI Bank - FD", "INCOME", ()),
                   (m.ACC_INTEREST_ON_FD, "INCOME", ("placeholder",))]
    book = write_book(tmp_path / "b.gnucash", specs)
    x = write_26as(tmp_path / "x.xlsx", {"Part I": [(1, "BANK OF BARODA", "194A", 1000.0, 100.0)]})
    m.run(x, book, tmp_path / "out.csv")
    text = tl.final_summary(str(tmp_path / "out.csv"), str(book))
    assert "NOT POSTED" in text and "BANK OF BARODA" in text


def test_tds04_final_summary_is_quiet_when_nothing_was_blocked(tmp_path):
    from agents.skill_26as_journal import tools as tl
    specs = STD + [("Income:Interest Income:Interest on BOB - FD", "INCOME", ()),
                   ("Income:Interest Income:Interest on ICICI Bank - FD", "INCOME", ()),
                   (GENERIC_FD, "INCOME", ())]
    book = write_book(tmp_path / "b.gnucash", specs)
    x = write_26as(tmp_path / "x.xlsx", {"Part I": [(1, "BANK OF BARODA", "194A", 1000.0, 100.0)]})
    m.run(x, book, tmp_path / "out.csv")
    assert "NOT POSTED" not in tl.final_summary(str(tmp_path / "out.csv"), str(book))   # NEGATIVE


# ---- TDS-05: Sr restarts per Part, so overrides are keyed series + Sr ----------------------

_TARGET = "Income:Interest Income:Interest on ICICI Bank - FD"


def _two_part_run(tmp_path, overrides):
    import json
    specs = STD + [(BOB, "INCOME", ()), (_TARGET, "INCOME", ()),
                   (m.ACC_INTEREST_ON_FD, "INCOME", ())]
    book = write_book(tmp_path / "b.gnucash", specs)
    x = write_26as(tmp_path / "x.xlsx", {
        "Part I": [(1, "BANK OF BARODA", "194A", 1000.0, 100.0)],
        "Part II": [(1, "BANK OF BARODA", "194A", 5000.0, 0.0)],
    })
    ov = tmp_path / "ov.json"
    ov.write_text(json.dumps(overrides), encoding="utf-8")
    out = tmp_path / "out.csv"
    assert m.main(["build", str(x), str(book), str(out), str(ov)]) == 0
    rows = {}
    for r in read_csv(tmp_path / "out-review.csv"):
        rows[m.series_for_category(r["Category"]) + r["Sr"]] = r["Credit Account"]
    return rows


def test_tds05_parse_override_key_reads_series_and_sr():
    assert m.parse_override_key("15GJ2") == ("15GJ", 2)
    assert m.parse_override_key("tdsj7") == ("TDSJ", 7)
    assert m.parse_override_key("TCSJ 3") == ("TCSJ", 3)
    assert m.parse_override_key("Sr 7") == ("TDSJ", 7)
    assert m.parse_override_key("7") == ("TDSJ", 7)
    assert m.parse_override_key("15GJ2") != ("TDSJ", 15)      # NEGATIVE: not read as Sr 15
    assert m.parse_override_key("no number") is None


def test_tds05_split_overrides_sends_each_key_to_its_own_part():
    a, b, c = m.split_overrides({"3": "x", "15GJ3": "y", "TCSJ3": "z", "junk": "w"})
    assert (a, b, c) == ({3: "x"}, {3: "y"}, {3: "z"})


def test_tds05_a_part_ii_override_does_not_change_the_part_i_row(tmp_path):
    base = _two_part_run_dir(tmp_path, "a", {})
    got = _two_part_run_dir(tmp_path, "b", {"15GJ1": _TARGET})
    assert got["TDSJ1"] == base["TDSJ1"]                      # NEGATIVE: Part I untouched
    assert got["15GJ1"] == _TARGET


def test_tds05_a_bare_number_never_changes_a_part_ii_row(tmp_path):
    base = _two_part_run_dir(tmp_path, "a", {})
    got = _two_part_run_dir(tmp_path, "b", {"1": _TARGET})
    assert got["15GJ1"] == base["15GJ1"]                      # NEGATIVE: Part II untouched
    assert got["TDSJ1"] == _TARGET                            # bare = Part I, as before


def _two_part_run_dir(tmp_path, name, overrides):
    d = tmp_path / name
    d.mkdir(exist_ok=True)
    return _two_part_run(d, overrides)


def test_tds05_tools_normalize_keeps_the_series_and_bare_stays_bare():
    from agents.skill_26as_journal import tools as tl
    assert tl._normalize_overrides({"15GJ2": "a", "Sr 7": "b", 3: "c", "x": "d"}) == \
        {"15GJ2": "a", "7": "b", "3": "c"}
    assert "15" not in tl._normalize_overrides({"15GJ2": "a"})   # NEGATIVE: not Sr 15


def test_tds05_gate_matches_review_rows_by_series_not_by_sr_alone(tmp_path):
    from agents.skill_26as_journal import tools as tl
    out = tmp_path / "o.csv"
    (tmp_path / "o-review.csv").write_text(
        "Sr,Deductor,Category,Confidence,Tied Candidates\n"
        "1,PART ONE PAYER,A,High,\n"
        "1,PART TWO PAYER,G,Ambiguous,Acc:One;Acc:Two\n", encoding="utf-8")
    # a bare "1" is Part I (High, not gated) -- it must NOT be judged against the Part II row
    assert tl._gate_ambiguous_overrides({"1": "Acc:Other"}, str(out)) == ({"1": "Acc:Other"}, [])
    # the Part II key IS gated against the Part II row's tied candidates
    rej = tl._gate_ambiguous_overrides({"15GJ1": "Acc:Other"}, str(out))
    assert isinstance(rej, str) and rej.startswith("REJECTED") and "15GJ1" in rej
    ok = tl._gate_ambiguous_overrides({"15GJ1": "Acc:Two"}, str(out))
    assert ok == ({"15GJ1": "Acc:Two"}, [])
