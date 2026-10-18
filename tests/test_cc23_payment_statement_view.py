"""tests/test_cc23_payment_statement_view.py -- CC-23: the read-only "Payments and statements by card" view.
Synthetic fixtures only. The view books nothing and decides no open item."""
from __future__ import annotations

import re
from datetime import date

from openpyxl import load_workbook

from agents.skill_cc_spend_booking import report as R
from agents.skill_cc_spend_booking.bookio import BookTxn, CCP_PATH
from tests import test_cc22_icici_csv as t22
from tests.test_cc24_csv_booking import _res, _src
from tests.test_cc27_csv_payment_match import PAY, _blocks
from tests.test_cc_spend_booking import (APR, FOOD, MAY, april_rows, book_pay, go, may_rows, stmt, tie)

TITLE = "Payments and statements by card:"


def _view(reply):
    assert TITLE in reply
    return reply.split(TITLE)[1].split("\n\n")[0]


def _pdf(tmp_path, monkeypatch, txns, start="2025-04-01", end="2025-05-31", rows_apr=None):
    tmp_path.mkdir(parents=True, exist_ok=True)
    s1, s2 = stmt(APR, rows_apr or april_rows()), stmt(MAY, may_rows())
    return go(tmp_path, monkeypatch, [s1, s2], [tie(s1, printed=1900.0), tie(s2, printed=300.0)], txns,
              start=start, end=end)


def _csv(tmp_path, monkeypatch, blocks, txns):
    d = _src(tmp_path, blocks=blocks)
    (tmp_path / "w").mkdir(parents=True, exist_ok=True)
    return go(tmp_path / "w", monkeypatch, [], [], txns, result_override=_res(d))


def test_a_payment_matched_to_a_csv_row_is_not_a_gap(tmp_path, monkeypatch):            # NEGATIVE
    reply, _r, _c, _o = _csv(tmp_path, monkeypatch, _blocks(("07-Apr-25", "5,000.00")),
                             [book_pay("p1", 5000, date(2025, 4, 8))])
    v = _view(reply)
    assert "5,000.00" in v and "CSV export" in v
    assert "payment line with no bank payment" not in v
    assert "Card not known" not in v


def test_a_fully_booked_statement_shows_all_rows_not_a_gap(tmp_path, monkeypatch):      # NEGATIVE
    _r, rows1, _c, _o = _pdf(tmp_path / "a", monkeypatch, [book_pay("p1", 1900, date(2025, 5, 12))])
    booked = [BookTxn(f"b{i}", date(2025, 5, 12), "x", r["Transaction ID"], ((FOOD, 1.0), (CCP_PATH, -1.0)))
              for i, r in enumerate(rows1)]
    reply, rows2, _c, _o = _pdf(tmp_path / "b", monkeypatch, [book_pay("p1", 1900, date(2025, 5, 12))] + booked)
    assert rows2 == []
    apr = [ln for ln in _view(reply).splitlines() if "Apr 2025" in ln and "PDF" in ln]
    assert apr and all("in the book: all rows" in ln for ln in apr)
    assert all("in the book: none" not in ln for ln in apr)


def test_a_bank_payment_with_no_statement_line_is_card_not_known(tmp_path, monkeypatch):    # NEGATIVE
    reply, _r, _c, _o = _csv(tmp_path, monkeypatch, _blocks(), [book_pay("p1", 5000, date(2025, 4, 8))])
    v = _view(reply)
    tail = v.split("Card not known")[1]
    assert "Apr 2025" in tail and "5,000.00" in tail
    assert "5,000.00" not in v.split("Card not known")[0]          # never under a card


def test_a_statement_with_no_payment_is_flagged(tmp_path, monkeypatch):
    reply, _r, _c, _o = _pdf(tmp_path, monkeypatch, [])
    assert "statement with no payment" in _view(reply)


def test_a_paid_statement_is_not_flagged_as_unpaid(tmp_path, monkeypatch):               # NEGATIVE
    reply, _r, _c, _o = _pdf(tmp_path, monkeypatch, [book_pay("p1", 1900, date(2025, 5, 12))])
    apr = [ln for ln in _view(reply).splitlines() if "Apr 2025" in ln and "PDF" in ln]
    assert apr and "statement with no payment" not in apr[0] and "1,900.00" in apr[0]


def test_partial_range_prints_the_line_and_a_full_year_does_not(tmp_path, monkeypatch):    # NEGATIVE
    part, _r, _c, _o = _pdf(tmp_path / "a", monkeypatch, [])
    assert "not the whole financial year" in part
    full, _r, _c, _o = _pdf(tmp_path / "b", monkeypatch, [], start="2025-04-01", end="2026-03-31")
    assert "not the whole financial year" not in full


def test_the_view_changes_no_journal_and_no_open_item(tmp_path, monkeypatch):            # NEGATIVE
    with_view, rows_a, _c, _o = _pdf(tmp_path / "a", monkeypatch, [book_pay("p1", 1900, date(2025, 5, 12))])
    monkeypatch.setattr(R, "payment_statement_block", lambda rep: "")
    without, rows_b, _c, _o = _pdf(tmp_path / "b", monkeypatch, [book_pay("p1", 1900, date(2025, 5, 12))])
    assert rows_a == rows_b
    pat = r"(NOT complete: \d+ open item|completed successfully)"
    assert re.search(pat, with_view).group(1) == re.search(pat, without).group(1)


def test_a_payment_outside_the_range_is_not_counted(tmp_path, monkeypatch):               # NEGATIVE
    reply, _r, _c, _o = _csv(tmp_path, monkeypatch, _blocks(), [book_pay("p1", 7777, date(2025, 6, 2))])
    assert "7,777.00" not in _view(reply)


def test_the_workbook_has_the_view_sheet(tmp_path, monkeypatch):
    _r, _rows, _c, out = _pdf(tmp_path, monkeypatch, [book_pay("p1", 1900, date(2025, 5, 12))])
    ws = load_workbook(out.with_name(out.stem + "_report.xlsx"))["Payments and statements by card"]
    assert "not the whole financial year" in str(ws["A2"].value)
    assert [c.value for c in ws[1]][0] == "Card"
