"""
tests/skill_itr_workbook/test_itr01_ai_pass.py -- ITR-01: the ITR Workbook's
AI tag-suggestion pass is declared (llm: true), capped, stops on the first
failure, says why in plain words without ever exposing a key, and an empty
mapping never reads as a normal success. Fully offline; synthetic fixtures.
"""
from __future__ import annotations

import logging
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

ROOT = Path(__file__).resolve().parent.parent.parent
SRC = ROOT / "src"
SCRIPTS = SRC / "agents" / "skill_itr_workbook" / "scripts"
AGENT_DIR = SRC / "agents" / "skill_itr_workbook"
FIXTURES = Path(__file__).resolve().parent / "fixtures"

for p in (str(SRC), str(SCRIPTS), str(AGENT_DIR), str(Path(__file__).resolve().parent)):
    if p not in sys.path:
        sys.path.insert(0, p)

import agent  # noqa: E402
import fixture_gen  # noqa: E402
import presentation  # noqa: E402
import suggest  # noqa: E402
from agents import llm_errors  # noqa: E402

ENTITIES_EXAMPLE = ROOT / "bundling" / "canonical" / "itr" / "entities.example.yaml"
SCRIPS_EXAMPLE = ROOT / "bundling" / "canonical" / "itr" / "scrips.example.yaml"

FAKE_KEY = "sk-FAKE0000000000000000KEYabcdef1234"


class _Leaf:
    def __init__(self, i):
        self.guid, self.path, self.total, self.tag = f"g{i}", f"Assets/Leaf {i}", 10.0 * i, "X"


class _HttpErr(Exception):
    def __init__(self, status, message):
        super().__init__(message)
        self.status_code = status
        self.body = {"error": {"message": message}}


def _leaves(n):
    return [_Leaf(i) for i in range(n)]


def _write_config(tmp_path):
    cfg = tmp_path / "config.yaml"
    cfg.write_text(yaml.safe_dump({
        "provider": "openai_compatible",
        "openai_compatible": {"base_url": "http://127.0.0.1:9/v1", "api_key": FAKE_KEY, "default_model": "m"},
    }), encoding="utf-8")
    return str(cfg)


# --- the manifest -----------------------------------------------------------

def test_skill_yaml_declares_llm_true():
    data = yaml.safe_load((AGENT_DIR / "skill.yaml").read_text(encoding="utf-8"))
    assert data["requires"]["llm"] is True


# --- stop on first failure --------------------------------------------------

def test_dead_endpoint_401_makes_exactly_one_call():
    calls = []

    def ask(*a):
        calls.append(a)
        raise _HttpErr(401, f"Incorrect API key provided: {FAKE_KEY}")

    run = suggest.run_suggestions(_leaves(8), {}, ask=ask, cancelled=lambda: False)
    assert len(calls) == 1 and run.attempted == 1
    assert run.failure_kind == "auth"
    assert "refused the key" in run.stopped_reason and "401" in run.stopped_reason
    assert run.left_without == 8


def test_key_never_appears_in_reason():
    def ask(*a):
        raise _HttpErr(401, f"Incorrect API key provided: {FAKE_KEY}. Bearer {FAKE_KEY}")

    run = suggest.run_suggestions(_leaves(2), {}, ask=ask, cancelled=lambda: False)
    assert FAKE_KEY not in run.stopped_reason
    assert "FAKE0000" not in run.stopped_reason


@pytest.mark.parametrize("exc,kind,words", [
    (_HttpErr(403, "forbidden"), "auth", "refused the key"),
    (_HttpErr(404, "The model `zz` does not exist"), "model_not_found", "does not have the model"),
    (TimeoutError("Request timed out."), "timeout", "did not answer in time"),
    (_HttpErr(429, "slow down"), "rate_limit", "rate-limiting"),
    (FileNotFoundError("Config not found: x"), "no_endpoint", "no AI endpoint"),
])
def test_failure_wording(exc, kind, words):
    f = llm_errors.classify(exc, model="zz")
    assert f.kind == kind and words in f.reason


def test_timeout_is_not_called_a_bad_key():
    f = llm_errors.classify(TimeoutError("Request timed out."))
    assert "refused the key" not in f.reason


def test_redact_removes_explicit_secret_and_key_shapes():
    out = llm_errors.redact(f"bad {FAKE_KEY} and Bearer abc123def456", secrets=("plainsecret1",))
    assert FAKE_KEY not in out and "abc123def456" not in out
    assert "plainsecret1" not in llm_errors.redact("x plainsecret1 y", secrets=("plainsecret1",))


# --- cap --------------------------------------------------------------------

def test_cap_is_never_exceeded():
    calls = []

    def ask(*a):
        calls.append(1)
        return None

    run = suggest.run_suggestions(_leaves(suggest.MAX_SUGGESTIONS_PER_RUN + 10), {}, ask=ask,
                                  cancelled=lambda: False)
    assert len(calls) == suggest.MAX_SUGGESTIONS_PER_RUN
    assert run.failure_kind == "cap" and "limit" in run.stopped_reason
    assert run.left_without == suggest.MAX_SUGGESTIONS_PER_RUN + 10


def test_nothing_unmapped_makes_no_call():
    run = suggest.run_suggestions([], {}, ask=lambda *a: pytest.fail("AI called"), cancelled=lambda: False)
    assert run.attempted == 0 and run.stopped_reason is None


# --- progress ----------------------------------------------------------------

def test_progress_says_n_of_total():
    import queue
    from agents.base_agent import set_progress_queue
    q = queue.Queue()
    set_progress_queue(q)
    try:
        suggest.run_suggestions(_leaves(3), {}, ask=lambda *a: None, cancelled=lambda: False)
    finally:
        set_progress_queue(None)
    snippets = []
    while not q.empty():
        snippets.append(q.get()["snippet"])
    assert "Suggesting tags: 1 of 3" in snippets and "Suggesting tags: 3 of 3" in snippets


# --- end to end through agent.run -------------------------------------------

def _html(tmp_path):
    p = tmp_path / "bs.html"
    p.write_text(fixture_gen.build_syn_ind_html(), encoding="utf-8")
    return str(p)


def _run(tmp_path, mapping=None, entity_key=None, config=None):
    out = tmp_path / "out.xlsx"
    summary = agent.run(
        _html(tmp_path), str(out), mapping_file=mapping, entities_path=str(ENTITIES_EXAMPLE),
        entity_key=entity_key, scrips_path=str(SCRIPS_EXAMPLE), config_path=config or "no-such-config.yaml",
    )
    return summary, out


def test_run_with_dead_key_reports_stop_at_top_and_hides_key(tmp_path, monkeypatch, caplog):
    calls = []

    def ask(*a):
        calls.append(1)
        raise _HttpErr(401, f"Incorrect API key provided: {FAKE_KEY}")

    monkeypatch.setattr(suggest, "_ask", ask)
    monkeypatch.setattr(suggest, "_is_cancelled", lambda: False)
    caplog.set_level(logging.DEBUG)
    summary, out = _run(tmp_path, str(FIXTURES / "syn_ind_partial.mapping.yaml"), "SYN-IND",
                        _write_config(tmp_path))
    assert len(calls) == 1
    first = summary.splitlines()[0]
    assert first.startswith("WARNING: AI tag suggestions stopped after 1 call(s)")
    assert "refused the key" in first and "3 of 3" in first
    snippet = next(out.parent.glob("out.xlsx*-proposed-mappings.yaml")).read_text(encoding="utf-8")
    for blob in (summary, snippet, caplog.text):
        assert FAKE_KEY not in blob


def test_fully_mapped_run_makes_no_ai_call(tmp_path, monkeypatch):
    monkeypatch.setattr(suggest, "_ask", lambda *a: pytest.fail("AI called on a fully mapped run"))
    summary, _ = _run(tmp_path, str(FIXTURES / "syn_ind.mapping.yaml"), "SYN-IND")
    assert "Mapping: OK" in summary
    assert "AI tag suggestions" not in summary


def test_empty_mapping_never_reads_as_normal_success(tmp_path, monkeypatch):
    monkeypatch.setattr(suggest, "_ask", lambda *a: None)
    monkeypatch.setattr(suggest, "_is_cancelled", lambda: False)
    summary, out = _run(tmp_path)   # cold start: no mapping anywhere
    assert summary.splitlines()[0].startswith(presentation.EMPTY_MAPPING_ERROR_MARKER)
    wb = openpyxl.load_workbook(out)
    soi = wb["Statement of Income"]
    assert presentation.EMPTY_MAPPING_ERROR_MARKER in str(soi["A1"].value)


def test_partially_mapped_run_has_no_empty_mapping_banner(tmp_path, monkeypatch):
    monkeypatch.setattr(suggest, "_ask", lambda *a: None)
    monkeypatch.setattr(suggest, "_is_cancelled", lambda: False)
    summary, out = _run(tmp_path, str(FIXTURES / "syn_ind_partial.mapping.yaml"), "SYN-IND")
    assert presentation.EMPTY_MAPPING_ERROR_MARKER not in summary
    assert presentation.EMPTY_MAPPING_ERROR_MARKER not in str(openpyxl.load_workbook(out)["Statement of Income"]["A1"].value)


def test_empty_mapping_exit_code_is_nonzero(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(suggest, "_is_cancelled", lambda: False)
    monkeypatch.setattr(suggest, "_ask", lambda *a: None)
    rc = agent.main([_html(tmp_path), str(tmp_path / "o.xlsx"), "--entities-path", str(ENTITIES_EXAMPLE),
                     "--scrips-path", str(SCRIPS_EXAMPLE)])
    assert rc == 1


def test_stop_never_leaves_a_workbook_that_looks_complete(tmp_path, monkeypatch):
    monkeypatch.setattr(suggest, "_ask", lambda *a: pytest.fail("AI called after Stop"))
    monkeypatch.setattr(suggest, "_is_cancelled", lambda: True)
    summary, out = _run(tmp_path, str(FIXTURES / "syn_ind_partial.mapping.yaml"), "SYN-IND")
    assert summary.splitlines()[0].startswith("ERROR: the run was stopped by the user")
    assert "NO workbook was built" in summary.splitlines()[0]
    assert openpyxl.load_workbook(out).sheetnames == ["Reconciliation"]
    assert not list(out.parent.glob("out.xlsx*-proposed-mappings.yaml"))
