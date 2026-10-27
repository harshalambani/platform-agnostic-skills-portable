"""
ITR-19 -- a gift of shares out of the book is not a capital-gains sale.

A disposal whose only non-stock legs are EXPENSE / EQUITY accounts brought in
no money and no receivable: not a transfer (s.47(iii)). It must stay out of
Schedule CG and the CG totals, still use up its buy lot (FIFO), and be listed
on the CapitalGains sheet with its reason. Any money / receivable leg keeps it
a sale.

Synthetic fixtures only (hand-built Book objects, made-up scrips/amounts).
"""
from __future__ import annotations

import sys
from datetime import date
from fractions import Fraction
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent.parent
SCRIPTS = ROOT / "src" / "agents" / "skill_itr_workbook" / "scripts"
RULES_DIR = ROOT / "bundling" / "canonical" / "itr" / "rules"
for p in (str(SCRIPTS), str(Path(__file__).resolve().parent)):
    if p not in sys.path:
        sys.path.insert(0, p)

import parse_gnucash as pg  # noqa: E402
import lots as lots_engine  # noqa: E402
import rules as rules_engine  # noqa: E402
import schedules as sch  # noqa: E402
import write_workbook as ww  # noqa: E402

YEAR_KEY = "2024-25"


def _accounts():
    accts = {
        "root": pg.Account(guid="root", name="Root", type="ROOT", parent_guid=None),
        "stock": pg.Account(guid="stock", name="FakeScrip", type="STOCK", parent_guid="root"),
        "bank": pg.Account(guid="bank", name="FakeBank", type="BANK", parent_guid="root"),
        "recv": pg.Account(guid="recv", name="FakeReceivable", type="RECEIVABLE", parent_guid="root"),
        "asset": pg.Account(guid="asset", name="FakeBrokerAsset", type="ASSET", parent_guid="root"),
        "gift": pg.Account(guid="gift", name="FakeGiftExpense", type="EXPENSE", parent_guid="root"),
        "eq": pg.Account(guid="eq", name="FakeEquity", type="EQUITY", parent_guid="root"),
        "gain": pg.Account(guid="gain", name="FakeCapGain", type="INCOME", parent_guid="root"),
    }
    pg._build_paths(accts)
    return accts


def _buy(guid, d, qty, cost):
    return pg.Transaction(guid=guid, date_posted=d, description="Buy", splits=[
        pg.Split(guid=guid + "-s", account_guid="stock", value=Fraction(cost), quantity=Fraction(qty)),
        pg.Split(guid=guid + "-b", account_guid="bank", value=Fraction(-cost), quantity=Fraction(-cost)),
    ])


def _sell(guid, d, qty, cost, proceeds, money="bank"):
    gain = proceeds - cost
    splits = [
        pg.Split(guid=guid + "-s", account_guid="stock", value=Fraction(-cost), quantity=Fraction(-qty)),
        pg.Split(guid=guid + "-m", account_guid=money, value=Fraction(proceeds), quantity=Fraction(proceeds)),
    ]
    if gain:
        splits.append(pg.Split(guid=guid + "-g", account_guid="gain", value=Fraction(-gain), quantity=Fraction(-gain)))
    return pg.Transaction(guid=guid, date_posted=d, description="Sell", splits=splits)


def _gift(guid, d, qty, cost, other="gift"):
    return pg.Transaction(guid=guid, date_posted=d, description="Gift of shares", splits=[
        pg.Split(guid=guid + "-s", account_guid="stock", value=Fraction(-cost), quantity=Fraction(-qty)),
        pg.Split(guid=guid + "-o", account_guid=other, value=Fraction(cost), quantity=Fraction(cost)),
    ])


def _book(*txns):
    return pg.Book(accounts=_accounts(), transactions=list(txns))


def _cg(book):
    rules = rules_engine.load_rules(RULES_DIR, YEAR_KEY)
    return sch.build_capital_gains({}, {}, book, YEAR_KEY, rules, {}, sch.FmvTables(nse={}, mf={}))


class _Cell:
    coordinate = "A1"


_LAYOUT = {k: _Cell() for k in ("ltcg_rate_before", "ltcg_rate_after", "stcg_rate_before", "stcg_rate_after")}


def test_real_sale_still_reported_with_right_gain():
    book = _book(_buy("b1", date(2023, 1, 1), 100, 10000),
                 _sell("s1", date(2024, 6, 1), 100, 10000, 13000))
    cg = _cg(book)
    assert len(cg.lot_rows) == 1
    assert cg.lot_rows[0].booked_gain == pytest.approx(3000.0)
    assert cg.lot_rows[0].proceeds == pytest.approx(13000.0)
    assert cg.non_sale_rows == []
    assert cg.reconciliation_ok


def test_gift_against_expense_is_not_in_cg_lots_or_totals():
    book = _book(_buy("b1", date(2023, 1, 1), 100, 10000),
                 _gift("g1", date(2024, 6, 1), 100, 10000))
    recons = lots_engine.reconstruct_lots(book, YEAR_KEY)
    assert lots_engine.all_lots(recons) == []
    cg = _cg(book)
    assert cg.lot_rows == []
    assert cg.lt_taxable_gross == 0 and cg.st_taxable_gross == 0
    assert cg.lt_taxable_on_after_split == 0 and cg.st_taxable_on_after_split == 0
    assert len(cg.non_sale_rows) == 1
    row = cg.non_sale_rows[0]
    assert row.scrip == "FakeScrip" and row.qty == pytest.approx(100.0)
    assert "s.47(iii)" in row.reason


def test_gift_then_sale_of_rest_matches_remaining_lot():
    # Lot A (older, 100 @ 10000) is gifted; lot B (200 @ 30000) is then sold.
    book = _book(_buy("bA", date(2022, 1, 1), 100, 10000),
                 _buy("bB", date(2023, 3, 1), 200, 30000),
                 _gift("g1", date(2024, 5, 1), 100, 10000),
                 _sell("s1", date(2024, 8, 1), 200, 30000, 36000))
    recons = lots_engine.reconstruct_lots(book, YEAR_KEY)
    sale_lots = lots_engine.all_lots(recons)
    assert len(sale_lots) == 1
    lot = sale_lots[0]
    assert lot.buy_txn_guid == "bB"
    assert lot.buy_date == date(2023, 3, 1)
    assert lot.attribution == "matched"
    assert lot.fifo_flag is None  # gifted lot A was consumed, so not "earlier available"
    assert lot.gain == pytest.approx(6000.0)
    gifted = lots_engine.non_sale_lots(recons)
    assert len(gifted) == 1 and gifted[0][0].buy_txn_guid == "bA"


def test_sale_into_asset_receivable_is_still_a_sale():
    for acct in ("recv", "asset", "bank"):
        book = _book(_buy("b1", date(2023, 1, 1), 100, 10000),
                     _sell("s1", date(2024, 6, 1), 100, 10000, 12500, money=acct))
        cg = _cg(book)
        assert len(cg.lot_rows) == 1, acct
        assert cg.lot_rows[0].proceeds == pytest.approx(12500.0)
        assert cg.non_sale_rows == [], acct


def test_disposal_against_equity_is_treated_like_a_gift():
    book = _book(_buy("b1", date(2023, 1, 1), 100, 10000),
                 _gift("g1", date(2024, 6, 1), 100, 10000, other="eq"))
    cg = _cg(book)
    assert cg.lot_rows == []
    assert len(cg.non_sale_rows) == 1


def test_gift_outside_the_year_is_not_listed_but_still_consumes_its_lot():
    book = _book(_buy("bA", date(2020, 1, 1), 100, 10000),
                 _buy("bB", date(2021, 1, 1), 100, 20000),
                 _gift("g1", date(2022, 5, 1), 100, 10000),
                 _sell("s1", date(2024, 8, 1), 100, 20000, 25000))
    cg = _cg(book)
    assert cg.non_sale_rows == []
    assert len(cg.lot_rows) == 1
    assert cg.lot_rows[0].buy_date == date(2021, 1, 1)


def test_disposal_with_a_money_leg_mixed_in_is_a_sale():
    # Expense leg AND a bank leg: money came in, so it is a sale.
    t = pg.Transaction(guid="m1", date_posted=date(2024, 6, 1), description="Part-paid", splits=[
        pg.Split(guid="m1-s", account_guid="stock", value=Fraction(-10000), quantity=Fraction(-100)),
        pg.Split(guid="m1-b", account_guid="bank", value=Fraction(4000), quantity=Fraction(4000)),
        pg.Split(guid="m1-e", account_guid="gift", value=Fraction(6000), quantity=Fraction(6000)),
    ])
    book = _book(_buy("b1", date(2023, 1, 1), 100, 10000), t)
    cg = _cg(book)
    assert len(cg.lot_rows) == 1
    assert cg.non_sale_rows == []


def test_not_a_sale_row_appears_on_capitalgains_sheet_with_reason():
    book = _book(_buy("b1", date(2023, 1, 1), 100, 10000),
                 _buy("b2", date(2023, 2, 1), 50, 6000),
                 _gift("g1", date(2024, 6, 1), 100, 10000),
                 _sell("s1", date(2024, 9, 1), 50, 6000, 7000))
    cg = _cg(book)
    wb = openpyxl.Workbook()
    ww.write_capital_gains_sheet(wb, cg, _LAYOUT)
    rows = [[c.value for c in r] for r in wb["CapitalGains"].iter_rows()]
    flat = [v for r in rows for v in r]
    assert "Disposals that are not a sale (not in Schedule CG)" in flat
    reason_rows = [r for r in rows if r[0] == "FakeScrip" and len(r) > 5 and r[5] and "s.47(iii)" in str(r[5])]
    assert len(reason_rows) == 1
    # The real sale is still in Lot detail, and only once.
    sale_rows = [r for r in rows if r[0] == "FakeScrip" and r[3] in ("ST", "LT")]
    assert len(sale_rows) == 1 and sale_rows[0][6] == pytest.approx(7000.0)


def test_sheet_has_no_not_a_sale_section_when_there_are_none():
    book = _book(_buy("b1", date(2023, 1, 1), 100, 10000),
                 _sell("s1", date(2024, 6, 1), 100, 10000, 13000))
    cg = _cg(book)
    wb = openpyxl.Workbook()
    ww.write_capital_gains_sheet(wb, cg, _LAYOUT)
    flat = [c.value for r in wb["CapitalGains"].iter_rows() for c in r]
    assert not any("not a sale" in str(v) for v in flat if v)
