"""
tests/skill_itr_workbook/test_itr07_tag_picker.py -- ITR-07: the tag picker.

1. An explicit HOLD choice: stored as unmapped with a reason, listed at the top
   of the run and on the Unclassified sheet, never fileable.
2. The picker offers only tags valid for the row's side; Apply refuses others.
3. Plain-word synonyms find tags, but only among tags the row's side allows.

Synthetic data only.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
import yaml

ROOT = Path(__file__).resolve().parent.parent.parent
SRC = ROOT / "src"
SCRIPTS = SRC / "agents" / "skill_itr_workbook" / "scripts"
for p in (str(ROOT), str(SRC), str(SCRIPTS), str(SRC / "agents" / "skill_itr_workbook"),
          str(Path(__file__).resolve().parent), str(ROOT / "tests")):
    if p not in sys.path:
        sys.path.insert(0, p)

import _review_js as js  # noqa: E402
import agent  # noqa: E402
import apply_mapping_corrections as amc  # noqa: E402
import fixture_gen  # noqa: E402
import tags  # noqa: E402
from ui.tabs import itr_mapping_review as ui_mod  # noqa: E402

ENTITIES_EXAMPLE = ROOT / "bundling" / "canonical" / "itr" / "entities.example.yaml"
SCRIPS_EXAMPLE = ROOT / "bundling" / "canonical" / "itr" / "scrips.example.yaml"
FIXTURES = Path(__file__).resolve().parent / "fixtures"
MISC = "ca85f6240bb66d5d38a0e4f40d908525"           # Assets/Misc Holding (unmapped)

needs_node = pytest.mark.skipif(not js.have_node(), reason="node not installed")


# --- helpers ----------------------------------------------------------------------

def _mapping_with_hold(tmp_path, note="ask the CA first"):
    entries = yaml.safe_load((FIXTURES / "syn_ind.mapping.yaml").read_text(encoding="utf-8"))
    for e in entries:
        if e["guid"] == MISC:
            e["tag"] = "HOLD"
            if note is not None:
                e["note"] = note
            else:
                e.pop("note", None)
    f = tmp_path / "held.mapping.yaml"
    f.write_text(yaml.safe_dump(entries), encoding="utf-8")
    return f


def _run(tmp_path, mapping_file):
    html = tmp_path / "bs.html"
    html.write_text(fixture_gen.build_syn_ind_html(), encoding="utf-8")
    out = tmp_path / "out.xlsx"
    summary = agent.run(str(html), str(out), mapping_file=str(mapping_file),
                        entities_path=str(ENTITIES_EXAMPLE), entity_key="SYN-IND",
                        scrips_path=str(SCRIPS_EXAMPLE), config_path="no-such-config.yaml")
    return summary, out


def _sheet_text(xlsx, name_part):
    wb = openpyxl.load_workbook(xlsx, data_only=False)
    out = []
    for ws in wb.worksheets:
        if name_part.lower() in ws.title.lower():
            for row in ws.iter_rows(values_only=True):
                out.append(" | ".join(str(c) for c in row if c is not None))
    return "\n".join(out)


def _all_text(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=False)
    return {ws.title: "\n".join(" | ".join(str(c) for c in r if c is not None)
                                for r in ws.iter_rows(values_only=True)) for ws in wb.worksheets}


# --- 1. HOLD: backend ----------------------------------------------------------------

def test_hold_is_a_valid_entry_tag_but_not_a_real_tag():
    assert tags.is_valid_entry_tag("HOLD")
    assert not tags.is_valid_tag("HOLD")
    assert "HOLD" not in tags.all_tags()
    for kind in tags.ROW_KINDS:                    # in no side's set
        assert "HOLD" not in tags.allowed_tags_for_kind(kind)


def test_held_leaf_is_unmapped_with_reason_and_never_resolved(tmp_path):
    import configs
    import mapping
    import parse_eguile as pe
    tree = pe.parse_html(fixture_gen.build_syn_ind_html())
    res = mapping.resolve_tree(tree, configs.load_mapping(_mapping_with_hold(tmp_path)))
    held = [u for u in res.unmapped if u.held]
    assert [u.guid for u in held] == [MISC]
    assert held[0].reason == "ask the CA first"
    assert MISC not in res.resolved                  # never counted as mapped


def test_hold_without_note_gets_default_reason(tmp_path):
    summary, _ = _run(tmp_path, _mapping_with_hold(tmp_path, note=None))
    assert "HELD" in summary and "held for review" in summary


def test_ordinary_unmapped_leaf_is_not_marked_held(tmp_path):
    summary, _ = _run(tmp_path, FIXTURES / "syn_ind_partial.mapping.yaml")
    assert "HELD:" not in summary                   # NEGATIVE: unmapped is not held


def test_run_with_hold_lists_it_first_and_is_not_filing_ready(tmp_path):
    summary, out = _run(tmp_path, _mapping_with_hold(tmp_path))
    head = summary.splitlines()
    held_line = next(i for i, l in enumerate(head) if l.startswith("HELD:"))
    status_line = next(i for i, l in enumerate(head) if l.startswith("STATUS:"))
    assert "NOT filing-ready" in head[held_line]
    assert "ask the CA first" in head[held_line]
    assert "REVIEW ITEM" in head[status_line]        # never plain OK
    assert "STATUS: OK" not in summary


def test_held_account_is_on_unclassified_sheet_with_its_reason(tmp_path):
    _, out = _run(tmp_path, _mapping_with_hold(tmp_path))
    text = _sheet_text(out, "Unclass")
    assert "Misc Holding" in text and "ask the CA first" in text
    assert "HELD ON PURPOSE" in text


def test_held_account_is_on_unclassified_and_a_mapped_one_is_not(tmp_path):
    """NEGATIVE: the same account, mapped, is NOT on the Unclassified sheet;
    held, it is - so a held row is never treated as mapped."""
    mapped_dir = tmp_path / "a"
    mapped_dir.mkdir()
    _, out_mapped = _run(mapped_dir, FIXTURES / "syn_ind.mapping.yaml")
    held_dir = tmp_path / "b"
    held_dir.mkdir()
    _, out_held = _run(held_dir, _mapping_with_hold(held_dir))
    assert "Misc Holding" not in _all_text(out_mapped).get("Unclassified", "")
    assert "Misc Holding" in _all_text(out_held)["Unclassified"]


def test_run_without_hold_is_unchanged_layout(tmp_path):
    """NEGATIVE: no held rows -> no HELD wording anywhere."""
    summary, out = _run(tmp_path, FIXTURES / "syn_ind.mapping.yaml")
    assert "HELD" not in summary
    assert "HELD ON PURPOSE" not in _sheet_text(out, "Unclass")


def test_snippet_marks_held_leaf_as_hold(tmp_path):
    _, out = _run(tmp_path, _mapping_with_hold(tmp_path))
    items = yaml.safe_load(fixture_gen.snippet_for(out).read_text(encoding="utf-8"))
    item = next(i for i in items if i["guid"] == MISC)
    assert item["tag"] == "HOLD" and item["note"] == "ask the CA first"
    assert item.get("section") == "Assets"


# --- 2. side-restricted tags: backend ----------------------------------------------------

def _base_mapping(tmp_path):
    f = tmp_path / "m.yaml"
    f.write_text(yaml.safe_dump([{"guid": "g0", "path": "Assets/X", "tag": "AL_CASH_BANK"}]),
                 encoding="utf-8")
    return str(f)


def test_bs_only_tag_refused_on_income_expense_row(tmp_path):
    out = tmp_path / "o.yaml"
    n, invalid = amc.apply_corrections_map(
        _base_mapping(tmp_path), {"g1": "AL_SECURITIES"}, str(out),
        paths={"g1": "Income/Dividends"}, row_kinds={"g1": tags.KIND_INCOME_EXPENSE})
    assert n == 0 and [i[1] for i in invalid] == ["g1"]
    assert "g1" not in {e["guid"] for e in yaml.safe_load(out.read_text(encoding="utf-8"))}


def test_income_tag_refused_on_balance_sheet_row(tmp_path):
    out = tmp_path / "o.yaml"
    n, invalid = amc.apply_corrections_map(
        _base_mapping(tmp_path), {"g1": "OS_DIVIDEND"}, str(out),
        paths={"g1": "Assets/Shares"}, row_kinds={"g1": tags.KIND_ASSET})
    assert n == 0 and len(invalid) == 1


def test_asset_tag_refused_on_liability_row(tmp_path):
    out = tmp_path / "o.yaml"
    n, invalid = amc.apply_corrections_map(
        _base_mapping(tmp_path), {"g1": "AL_CASH_BANK"}, str(out),
        paths={"g1": "Liabilities/Loan"}, row_kinds={"g1": tags.KIND_LIABILITY})
    assert n == 0 and len(invalid) == 1


def test_right_side_tag_is_applied(tmp_path):
    out = tmp_path / "o.yaml"
    n, invalid = amc.apply_corrections_map(
        _base_mapping(tmp_path), {"g1": "OS_DIVIDEND"}, str(out),
        paths={"g1": "Income/Dividends"}, row_kinds={"g1": tags.KIND_INCOME_EXPENSE})
    assert n == 1 and not invalid


def test_unknown_side_accepts_only_what_every_side_accepts(tmp_path):
    out = tmp_path / "o.yaml"
    n, invalid = amc.apply_corrections_map(
        _base_mapping(tmp_path), {"a": "AL_SECURITIES", "b": "OS_DIVIDEND", "c": "PERSONAL"},
        str(out), paths={"a": "x", "b": "y", "c": "z"},
        row_kinds={k: tags.KIND_UNKNOWN for k in "abc"})
    assert {i[1] for i in invalid} == {"a", "b"} and n == 1


def test_hold_is_accepted_on_every_side_with_reason(tmp_path):
    out = tmp_path / "o.yaml"
    guids = {f"k{i}": k for i, k in enumerate(tags.ROW_KINDS)}
    n, invalid = amc.apply_corrections_map(
        _base_mapping(tmp_path), {g: "HOLD" for g in guids}, str(out),
        paths={g: "p" for g in guids}, row_kinds=guids,
        hold_reasons={g: "waiting on statement" for g in guids})
    assert n == len(guids) and not invalid
    saved = {e["guid"]: e for e in yaml.safe_load(out.read_text(encoding="utf-8"))}
    assert all(saved[g]["tag"] == "HOLD" and saved[g]["note"] == "waiting on statement" for g in guids)


def test_non_hold_entries_do_not_keep_a_hold_reason(tmp_path):
    out = tmp_path / "o.yaml"
    amc.apply_corrections_map(_base_mapping(tmp_path), {"g1": "OS_DIVIDEND"}, str(out),
                              paths={"g1": "Income/D"}, row_kinds={"g1": tags.KIND_INCOME_EXPENSE},
                              hold_reasons={"g1": "stale reason"})
    saved = {e["guid"]: e for e in yaml.safe_load(out.read_text(encoding="utf-8"))}
    assert "stale reason" not in str(saved["g1"].get("note"))


def test_hold_can_be_replaced_by_a_real_tag_and_back(tmp_path):
    held = tmp_path / "h.yaml"
    held.write_text(yaml.safe_dump([{"guid": "g1", "path": "Assets/X", "tag": "HOLD", "note": "why"}]),
                    encoding="utf-8")
    out = tmp_path / "o.yaml"
    n, _ = amc.apply_corrections_map(str(held), {"g1": "AL_CASH_BANK"}, str(out),
                                     row_kinds={"g1": tags.KIND_ASSET})
    assert n == 1
    assert yaml.safe_load(out.read_text(encoding="utf-8"))[0]["tag"] == "AL_CASH_BANK"


# --- the kind table itself ------------------------------------------------------------------

def test_kinds_allowing_never_includes_a_forbidden_kind():
    for tag in tags.all_tags():
        for kind in tags.ROW_KINDS:
            assert (kind in tags.kinds_allowing(tag)) == (tag in tags.allowed_tags_for_kind(kind))


def test_balance_sheet_only_and_income_only_tags_do_not_overlap_on_kinds():
    assert tags.KIND_INCOME_EXPENSE not in tags.kinds_allowing("AL_SECURITIES")
    assert tags.KIND_ASSET not in tags.kinds_allowing("OS_DIVIDEND")


# --- 3. synonyms ---------------------------------------------------------------------------

def test_synonyms_find_the_expected_tags():
    def hits(word):
        return {t for t in tags.all_tags()
                if any(word in s for s in tags.synonyms_for(t))}
    assert "DED_80D_CANDIDATE" in hits("medical") and "DED_80D_CANDIDATE" in hits("health")
    for w in ("school", "tuition", "life cover"):
        assert "DED_80C_CANDIDATE" in hits(w)


def test_synonym_for_partner_points_at_the_partner_tds_tag():
    assert "partner" in tags.synonyms_for("TAXPAID_TDS_PARTNER")


def test_every_synonym_belongs_to_a_known_tag_or_hold():
    pending = {"TAXPAID_TDS_PARTNER"}           # arrives with ITR-02
    for t in tags.TAG_SYNONYMS:
        assert t in tags.all_tags() or t == tags.HOLD or t in pending


# --- Review Mapping tab ----------------------------------------------------------------------

def _root(tmp_path):
    data_root = tmp_path / "Data"
    (data_root / "itr" / "mappings").mkdir(parents=True)
    (data_root / "outputs").mkdir()
    (data_root / "itr" / "entities.yaml").write_text(yaml.safe_dump(
        {"ONE": {"name": "One", "pan": "AAAAA0000A", "status": "Individual", "default_regime": "new"}}),
        encoding="utf-8")
    return data_root


def _write(root, rows, held=None):
    snip = [{"guid": g, "path": p, "tag": "REPLACE_ME", "note": "n"} for g, p in rows]
    (root / "outputs" / "s-proposed-mappings.yaml").write_text(
        "# entity: ONE\n" + yaml.safe_dump(snip), encoding="utf-8")
    if held:
        (root / "itr" / "mappings" / "ONE.mapping.yaml").write_text(yaml.safe_dump(
            [{"guid": g, "path": p, "tag": "HOLD", "note": r} for g, p, r in held]), encoding="utf-8")


def _html(root):
    with patch("ui._config.data_root_dir", return_value=root):
        return ui_mod._load_review_data("ONE")


def _rows(root):
    with patch("ui._config.data_root_dir", return_value=root):
        rows, _ = ui_mod._collect_review("ONE")
    return {r["guid"]: r for r in rows}


def test_row_kind_comes_from_the_path_root(tmp_path):
    root = _root(tmp_path)
    _write(root, [("i1", "Income/Dividends"), ("a1", "Assets/Bank"), ("l1", "Liabilities/Loan"),
                  ("q1", "Weird/Thing")])
    by = _rows(root)
    assert by["i1"]["kind"] == tags.KIND_INCOME_EXPENSE
    assert by["a1"]["kind"] == tags.KIND_ASSET
    assert by["l1"]["kind"] == tags.KIND_LIABILITY
    assert by["q1"]["kind"] == tags.KIND_UNKNOWN


def test_held_row_is_never_shown_as_mapped(tmp_path):
    root = _root(tmp_path)
    _write(root, [], held=[("h1", "Assets/Hold Me", "waiting on bank")])
    r = _rows(root)["h1"]
    assert r["held"] and r["hold_reason"] == "waiting on bank"
    ui_mod._row_presentation(r, {})
    assert r["_badges"][ui_mod.TARGET_COL]["text"] == "ON HOLD"
    assert "unmapped" in r["_tags"] and "held" in r["_tags"]
    assert "confirmed" not in r["_tags"] and "mapped" not in r["_tags"]


def test_page_explains_what_unmapped_and_hold_do(tmp_path):
    root = _root(tmp_path)
    _write(root, [("a1", "Assets/Bank")])
    html = _html(root)
    assert "NOT filing-ready" in html and "HOLD" in html and "Unclassified" in html


def test_save_refuses_wrong_side_tag_and_says_why(tmp_path):
    root = _root(tmp_path)
    _write(root, [("i1", "Income/Dividends")])
    payload = json.dumps({"context": {"entity_key": "ONE"}, "changes": [
        {"guid": "i1", "path": "Income/Dividends", "tag": "AL_SECURITIES", "kind": tags.KIND_INCOME_EXPENSE}]})
    with patch("ui._config.data_root_dir", return_value=root):
        msg = ui_mod._save_changes(payload)
    assert "NOT applied" in msg and "not valid for a Income/expense account" in msg
    mf = root / "itr" / "mappings" / "ONE.mapping.yaml"
    assert (not mf.exists()) or "AL_SECURITIES" not in mf.read_text(encoding="utf-8")


def test_save_with_missing_side_fails_closed(tmp_path):
    root = _root(tmp_path)
    _write(root, [("q1", "Weird/Thing")])
    payload = json.dumps({"context": {"entity_key": "ONE"}, "changes": [
        {"guid": "q1", "path": "Weird/Thing", "tag": "AL_SECURITIES", "kind": ""}]})
    with patch("ui._config.data_root_dir", return_value=root):
        msg = ui_mod._save_changes(payload)
    assert "NOT applied" in msg


def test_save_hold_keeps_the_reason(tmp_path):
    root = _root(tmp_path)
    _write(root, [("a1", "Assets/Bank")])
    payload = json.dumps({"context": {"entity_key": "ONE"}, "changes": [
        {"guid": "a1", "path": "Assets/Bank", "tag": "HOLD", "kind": tags.KIND_ASSET,
         "hold_reason": "need the passbook"}]})
    with patch("ui._config.data_root_dir", return_value=root):
        ui_mod._save_changes(payload)
    saved = yaml.safe_load((root / "itr" / "mappings" / "ONE.mapping.yaml").read_text(encoding="utf-8"))
    assert saved[0]["tag"] == "HOLD" and saved[0]["note"] == "need the passbook"


# --- picker in the browser engine (node) ---------------------------------------------------------

PRE = """
globalThis.__alerts = [];
globalThis.alert = function (m) { __alerts.push(String(m)); };
globalThis.__promptAnswer = 'because';
globalThis.prompt = function (t, d) { globalThis.__promptText = t; return globalThis.__promptAnswer; };
"""

# list the dropdown after typing a search, return [primary names]
LIST = """
function listed(q) {
  const s = $id(APP + '-picker-search');
  s.onfocus();
  s.value = q || '';
  if (s.oninput) s.oninput();
  return $id(APP + '-picker-dropdown').children.map(c => c._h || c.innerHTML || c.textContent);
}
"""


def _page(tmp_path):
    root = _root(tmp_path)
    _write(root, [("i1", "Income/Dividends"), ("a1", "Assets/Bank")])
    return _html(root)


def _idx(out_view, path_part):
    for r in out_view:
        if any(path_part in c for c in r["cells"]):
            return r["idx"]
    raise AssertionError(path_part)


@needs_node
def test_picker_for_income_row_lists_no_balance_sheet_tag(tmp_path):
    html = _page(tmp_path)
    out = js.run_js(html, "itrmap", LIST + """
        const v = view(APP); const i = v.find(r => r.cells.some(c => c.includes('Dividends'))).idx;
        select(APP, i);
        return listed('');
    """, tmp_path, PRE)
    joined = " ".join(out)
    assert "OS_DIVIDEND" in joined and "HOLD" in joined
    assert "AL_SECURITIES" not in joined and "AL_CASH_BANK" not in joined


@needs_node
def test_picker_for_asset_row_lists_no_income_tag(tmp_path):
    html = _page(tmp_path)
    out = js.run_js(html, "itrmap", LIST + """
        const v = view(APP); const i = v.find(r => r.cells.some(c => c.includes('Bank'))).idx;
        select(APP, i);
        return listed('');
    """, tmp_path, PRE)
    joined = " ".join(out)
    assert "AL_CASH_BANK" in joined and "HOLD" in joined
    assert "OS_DIVIDEND" not in joined and "SALARY_GROSS" not in joined


@needs_node
def test_synonym_never_surfaces_a_forbidden_tag(tmp_path):
    """NEGATIVE: 'shares' is a synonym of AL_SECURITIES; on an income row the
    search must not list it. 'dividends' on an asset row must not list OS_DIVIDEND."""
    html = _page(tmp_path)
    out = js.run_js(html, "itrmap", LIST + """
        const v = view(APP);
        const inc = v.find(r => r.cells.some(c => c.includes('Dividends'))).idx;
        const ast = v.find(r => r.cells.some(c => c.includes('Bank'))).idx;
        select(APP, inc);
        const a = listed('shares');
        select(APP, ast);
        const b = listed('dividends');
        return {a, b};
    """, tmp_path, PRE)
    assert "AL_SECURITIES" not in " ".join(out["a"])
    assert "OS_DIVIDEND" not in " ".join(out["b"])


@needs_node
def test_synonym_finds_the_right_tag_on_a_valid_row(tmp_path):
    html = _page(tmp_path)
    out = js.run_js(html, "itrmap", LIST + """
        const v = view(APP); const i = v.find(r => r.cells.some(c => c.includes('Bank'))).idx;
        select(APP, i);
        return listed('demat');
    """, tmp_path, PRE)
    assert any("AL_SECURITIES" in c for c in out)


@needs_node
def test_apply_of_wrong_side_tag_is_refused_and_changes_nothing(tmp_path):
    """NEGATIVE: pick a valid asset tag, then add an income row to the selection
    (the stale pick is now wrong for it). Apply must refuse the WHOLE apply with
    a message, change neither row, and record nothing."""
    html = _page(tmp_path)
    out = js.run_js(html, "itrmap", """
        const before = view(APP);
        const ast = before.find(r => r.cells.some(c => c.includes('Bank'))).idx;
        const inc = before.find(r => r.cells.some(c => c.includes('Dividends'))).idx;
        select(APP, ast);
        const sb = $id(APP + '-picker-search'); sb.value = ''; sb.onfocus();
        const k = $id(APP + '-picker-dropdown').children.findIndex(c => (c._h || '').includes('AL_CASH_BANK'));
        pick(APP, k);
        const tr = $id(APP + '-tbody').children.find(t => Number(t.dataset.idx) === inc);
        tr.onclick({ shiftKey: false, ctrlKey: true, metaKey: false });
        const mid = view(APP).map(r => r.cls.includes('selected'));
        applySel(APP);
        return {mid, alerts: __alerts, before, after: view(APP), p: payload(%s)};
    """ % json.dumps(ui_mod.PAYLOAD_VAR), tmp_path, PRE)
    assert all(out["mid"]), "both rows should be selected for this scenario"
    assert any("Not applied" in a for a in out["alerts"])
    assert out["p"].get("changes", []) == []
    strip = lambda v: [[c for c in r["cells"]] for r in v]
    assert strip(out["before"]) == strip(out["after"])


@needs_node
def test_hold_asks_for_a_reason_and_stores_it_then_clears_on_real_tag(tmp_path):
    html = _page(tmp_path)
    out = js.run_js(html, "itrmap", """
        const v = view(APP); const i = v.find(r => r.cells.some(c => c.includes('Bank'))).idx;
        select(APP, i);
        pick(APP, 0);                 // HOLD is listed first
        applySel(APP);
        const p1 = payload(%s);
        const asked = globalThis.__promptText;
        // the row stays selected after Apply
        const sb = $id(APP + '-picker-search'); sb.value = ''; sb.onfocus();
        const dd = $id(APP + '-picker-dropdown').children.map(c => c._h || c.innerHTML || '');
        const k = dd.findIndex(t => t.includes('AL_CASH_BANK'));
        pick(APP, k); applySel(APP);
        const p2 = payload(%s);
        return {p1, p2, asked};
    """ % (json.dumps(ui_mod.PAYLOAD_VAR), json.dumps(ui_mod.PAYLOAD_VAR)), tmp_path, PRE)
    c1 = [c for c in out["p1"]["changes"] if c["tag"] == "HOLD"]
    assert c1 and c1[0]["hold_reason"] == "because"
    assert "on hold" in out["asked"]
    c2 = [c for c in out["p2"]["changes"] if c["tag"] == "AL_CASH_BANK"]
    assert c2 and not c2[0].get("hold_reason")


@needs_node
def test_cancelling_the_hold_prompt_assigns_nothing(tmp_path):
    html = _page(tmp_path)
    out = js.run_js(html, "itrmap", """
        globalThis.__promptAnswer = null;
        const v = view(APP); const i = v.find(r => r.cells.some(c => c.includes('Bank'))).idx;
        select(APP, i);
        pick(APP, 0); applySel(APP);
        return {p: payload(%s), alerts: __alerts};
    """ % json.dumps(ui_mod.PAYLOAD_VAR), tmp_path, PRE)
    assert not [c for c in out["p"].get("changes", []) if c.get("tag") == "HOLD"]
    assert not any("cannot be reassigned" in a for a in out["alerts"])
