"""
tests/skill_itr_workbook/test_itr06_foreign_tax_credit.py -- ITR-06: foreign
tax credit. Foreign tax is NOT TDS: it never enters the TDS credit or the 26AS
tie-out; relief per country is the lower of the foreign tax and the Indian tax
on that income; the same dividend is never counted twice; relief is never
claimed without the Form 67 line; a run with no foreign activity is unchanged.

Synthetic data only: the SYN-IND fixture, a fictitious country and in-memory
broker reports.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

ROOT = Path(__file__).resolve().parent.parent.parent
SCRIPTS = ROOT / "src" / "agents" / "skill_itr_workbook" / "scripts"
FIXTURES = Path(__file__).resolve().parent / "fixtures"
RULES_DIR = ROOT / "bundling" / "canonical" / "itr" / "rules"

for p in (str(SCRIPTS), str(Path(__file__).resolve().parent)):
    if p not in sys.path:
        sys.path.insert(0, p)

import configs  # noqa: E402
import fixture_gen  # noqa: E402
import foreign_tax as ft  # noqa: E402
import mapping as mapping_engine  # noqa: E402
import parse_eguile as pe  # noqa: E402
import parse_foreign as pf  # noqa: E402
import parse_form16  # noqa: E402
import parse_gnucash as pg  # noqa: E402
import rules as rules_engine  # noqa: E402
import schedules as sch  # noqa: E402
import tags as tag_vocab  # noqa: E402
import write_workbook as ww  # noqa: E402

YEAR_KEY = "2024-25"
REGIME = "new"
TDS_INT_GUID = "3b305b86a4a5d6c626c6234b129ea174"      # Expense/TDS on Interest
BANK_INT_GUID = "11fc5a724efb9eb161ac039825b3e6dd"     # Income/Bank Interest
COUNTRY = "Exampleland"


# ---------------------------------------------------------------------------
# Builders
# ---------------------------------------------------------------------------

def _mapping(tmp_path, *, foreign_tax_tag=False, foreign_div=None):
    """syn_ind mapping, optionally with the TDS-on-interest account re-tagged
    TAXPAID_FOREIGN and the bank-interest account re-tagged as a foreign
    OS_DIVIDEND (flag name given by `foreign_div`)."""
    entries = yaml.safe_load((FIXTURES / "syn_ind.mapping.yaml").read_text(encoding="utf-8"))
    for e in entries:
        if foreign_tax_tag and e["guid"] == TDS_INT_GUID:
            e["tag"] = "TAXPAID_FOREIGN"
        if foreign_div and e["guid"] == BANK_INT_GUID:
            e["tag"] = "OS_DIVIDEND"
            e["flags"] = [foreign_div]
    path = tmp_path / "m.mapping.yaml"
    path.write_text(yaml.safe_dump(entries), encoding="utf-8")
    return configs.load_mapping(path)


def _report(*, fsi_rows=(), ordinary_dividend=None):
    dividend = None
    if ordinary_dividend is not None:
        dividend = pf.DividendData(
            total=ordinary_dividend, financial_year="2024-25",
            quarterly=[pf.QuarterlyBucket("1-Apr to 15-Jun", ordinary_dividend, series="Dividend Income")],
        )
    return pf.ForeignIncomeReport(
        source_path="<in-memory>", financial_year="2024-25", report_as_on=None,
        schedule_fa=pf.ScheduleFAData(),
        schedule_fsi=pf.ScheduleFSIData(rows=list(fsi_rows), financial_year="2024-25"),
        form_67=pf.Form67Data(), schedule_tr=pf.ScheduleTRData(),
        dividend=dividend, interest=pf.InterestData(),
    )


def _fsi(country=COUNTRY, income=100000.0, paid=15000.0):
    return pf.ScheduleFSIRow(
        head_of_income="Other Sources", country_code="999", country=country, tin=None,
        income=income, tax_paid_outside=paid, tax_payable_india=None, tax_relief=None,
        dtaa_article=None,
    )


def _model(tmp_path, *, foreign_report=None, foreign_dividends_in_book=False, **mapping_kw):
    tree = pe.parse_html(fixture_gen.build_syn_ind_html())
    book = pg.parse_book(FIXTURES / "syn_ind.gnucash")
    loaded = _mapping(tmp_path, **mapping_kw)
    result = mapping_engine.resolve_tree(tree, loaded)
    form16 = parse_form16.parse_form16(FIXTURES / "syn_ind_form16.pdf")
    rules = rules_engine.load_rules(RULES_DIR, YEAR_KEY)
    entity = configs.load_entities(ROOT / "bundling" / "canonical" / "itr" / "entities.example.yaml")["SYN-IND"]
    scrips = configs.load_scrips(ROOT / "bundling" / "canonical" / "itr" / "scrips.example.yaml")
    model = sch.build_all_schedules(
        tree, result.resolved, book, form16, YEAR_KEY, rules, REGIME, entity.status, entity.dob,
        scrips, sch.load_fmv_tables(), foreign_report=foreign_report,
        foreign_dividends_in_book=foreign_dividends_in_book,
    )
    return model, tree, result, loaded, rules, entity, form16


def _workbook(tmp_path, built):
    model, tree, result, loaded, rules, entity, form16 = built
    out = tmp_path / "wb.xlsx"
    ww.write_workbook(
        str(out), tree, model, rules, rules_engine.load_user_rules(RULES_DIR / "user_rules.yaml"),
        entity, REGIME, YEAR_KEY, form16.opted_out_115bac, [], [], [], result.unmapped, "v1",
        "2026-01-01T00:00:00", {}, result.resolved, loaded.entries,
    )
    return openpyxl.load_workbook(str(out))


def _all_text(wb) -> str:
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    out.append(c.value)
    return "\n".join(out)


def _foreign_amount(tmp_path) -> float:
    tree = pe.parse_html(fixture_gen.build_syn_ind_html())
    return abs(sch._node_by_guid(tree)[TDS_INT_GUID].total)


# ---------------------------------------------------------------------------
# The tag and the flag exist
# ---------------------------------------------------------------------------

def test_tag_and_flag_are_in_the_vocabulary():
    assert tag_vocab.is_valid_tag("TAXPAID_FOREIGN")
    assert tag_vocab.is_valid_flag("FOREIGN_INCOME") and tag_vocab.is_valid_flag("AL_FOREIGN")
    assert "TAXPAID_FOREIGN" in tag_vocab.RE_TAGS and "TAXPAID_FOREIGN" not in tag_vocab.BS_TAGS


# ---------------------------------------------------------------------------
# NEGATIVE: foreign tax never enters the TDS credit or the 26AS tie-out
# ---------------------------------------------------------------------------

def test_foreign_tax_is_not_in_the_tds_credit(tmp_path):
    amount = _foreign_amount(tmp_path)
    assert amount > 0
    base = _model(tmp_path)[0]
    foreign = _model(tmp_path, foreign_tax_tag=True)[0]
    # The account left TDS-on-interest and went nowhere into taxes paid.
    assert foreign.taxes_paid.tds_interest == pytest.approx(base.taxes_paid.tds_interest - amount)
    assert foreign.taxes_paid.total == pytest.approx(base.taxes_paid.total - amount)
    assert foreign.foreign_tax.book_foreign_tax == pytest.approx(amount)
    # ... and is not part of any tie-out.
    assert not foreign.taxes_paid.tie_out_conflicts
    assert foreign.taxes_paid.as26_tds_interest == 0.0 and foreign.taxes_paid.as26_tds_dividend == 0.0


# ---------------------------------------------------------------------------
# Relief: the LOWER of foreign tax and Indian tax, per country
# ---------------------------------------------------------------------------

def test_relief_is_the_foreign_tax_when_that_is_lower(tmp_path):
    rep = _report(ordinary_dividend=300000.0, fsi_rows=[_fsi(income=100000.0, paid=1.0)])
    base = _model(tmp_path, foreign_report=_report(ordinary_dividend=300000.0))[0]
    m = _model(tmp_path, foreign_report=rep)[0]
    assert m.foreign_tax.relief_total == pytest.approx(1.0)
    assert m.computation.tax_block.tax_liability == pytest.approx(
        base.computation.tax_block.tax_liability - 1.0)
    assert m.computation.foreign_tax_relief == pytest.approx(1.0)


def test_relief_never_exceeds_the_indian_tax_on_that_income(tmp_path):
    rep = _report(ordinary_dividend=300000.0, fsi_rows=[_fsi(income=100000.0, paid=9_999_999.0)])
    m = _model(tmp_path, foreign_report=rep)[0]
    row = m.foreign_tax.rows[0]
    assert row.foreign_tax == pytest.approx(9_999_999.0)
    assert row.indian_tax < row.foreign_tax
    assert row.relief == pytest.approx(row.indian_tax)
    assert m.foreign_tax.relief_total <= m.computation.tax_block.tax_liability_before_foreign_relief
    assert m.computation.tax_block.tax_liability >= 0.0


def test_relief_is_per_country_not_pooled():
    sched = ft.ForeignTaxSchedule(rows=[
        ft.ForeignTaxRow("Aland", income=100000.0, foreign_tax=0.0, source="t"),
        ft.ForeignTaxRow("Bland", income=100000.0, foreign_tax=50000.0, source="t"),
    ])
    relief = ft.apply_relief(sched, tax_liability=20000.0, total_income=200000.0)
    a, b = sched.rows
    assert a.relief == 0.0                              # Aland paid nothing abroad
    assert b.relief == pytest.approx(10000.0)           # capped at Indian tax on Bland's income
    assert relief == pytest.approx(10000.0)             # Aland's spare headroom is not lent to Bland


def test_relief_never_exceeds_the_whole_tax_liability():
    sched = ft.ForeignTaxSchedule(rows=[ft.ForeignTaxRow("Aland", 500000.0, 900000.0, "t")])
    relief = ft.apply_relief(sched, tax_liability=1000.0, total_income=100000.0)
    assert relief <= 1000.0 and relief >= 0.0


def test_no_relief_without_a_country(tmp_path):
    """Book-only foreign tax: the country is unknown, so no relief is claimed
    and the reader is told so."""
    m = _model(tmp_path, foreign_tax_tag=True)[0]
    assert m.foreign_tax.relief_total == 0.0 and not m.foreign_tax.relief_claimed
    assert m.foreign_tax.unallocated_tax == pytest.approx(_foreign_amount(tmp_path))
    assert any("NO relief is claimed" in w for w in m.foreign_tax.warnings)


def test_relief_flows_into_interest_234(tmp_path):
    """234B/C measure against the tax AFTER relief, so the relief reaches them."""
    rep = _report(ordinary_dividend=300000.0, fsi_rows=[_fsi(income=100000.0, paid=5000.0)])
    base = _model(tmp_path, foreign_report=_report(ordinary_dividend=300000.0))[0]
    m = _model(tmp_path, foreign_report=rep)[0]
    relief = m.foreign_tax.relief_total
    assert relief > 0
    assert m.interest_234 is not None
    # The assessed tax the interest engine saw is the reduced one.
    assert m.computation.tax_block.tax_liability == pytest.approx(
        base.computation.tax_block.tax_liability - relief)


# ---------------------------------------------------------------------------
# NEGATIVE: the same dividend is never counted twice
# ---------------------------------------------------------------------------

def test_book_and_report_both_counting_dividends_is_a_loud_stop(tmp_path):
    rep = _report(ordinary_dividend=700.0)
    with pytest.raises(ft.ForeignDoubleCountError) as ei:
        _model(tmp_path, foreign_report=rep, foreign_dividends_in_book=False, foreign_div="FOREIGN_INCOME")
    assert ft.DOUBLE_COUNT_ERROR_MARKER in str(ei.value)


def test_al_foreign_flag_on_an_os_dividend_also_trips_the_stop(tmp_path):
    rep = _report(ordinary_dividend=700.0)
    with pytest.raises(ft.ForeignDoubleCountError):
        _model(tmp_path, foreign_report=rep, foreign_dividends_in_book=False, foreign_div="AL_FOREIGN")


def test_book_holds_them_report_is_kept_out_of_other_sources(tmp_path):
    """foreign_dividends_in_book=True: the report's dividends feed only
    FSI/TR/FA and Form 67 -- never Other Sources again."""
    base = _model(tmp_path, foreign_div="FOREIGN_INCOME", foreign_dividends_in_book=True)[0]
    rep = _report(ordinary_dividend=700.0)
    m = _model(tmp_path, foreign_report=rep, foreign_dividends_in_book=True, foreign_div="FOREIGN_INCOME")[0]
    assert m.other_sources.dividend_gross == pytest.approx(base.other_sources.dividend_gross)
    assert m.other_sources.taxable_total == pytest.approx(base.other_sources.taxable_total)
    assert m.other_sources.foreign_dividend_gross == pytest.approx(700.0)   # detail only


def test_no_stop_without_a_book_foreign_account(tmp_path):
    """A report alone (flag false) is the normal path and is not stopped."""
    base = _model(tmp_path)[0]
    m = _model(tmp_path, foreign_report=_report(ordinary_dividend=700.0), foreign_dividends_in_book=False)[0]
    assert m.other_sources.taxable_total == pytest.approx(base.other_sources.taxable_total + 700.0)


def test_book_figures_never_added_to_the_reports(tmp_path):
    """Tax and income exist in both the book and the report: the report's
    figure is used, the two are never added, and a mismatch is warned."""
    amount = _foreign_amount(tmp_path)
    rep = _report(ordinary_dividend=300000.0, fsi_rows=[_fsi(income=100000.0, paid=amount + 500.0)])
    m = _model(tmp_path, foreign_report=rep, foreign_tax_tag=True)[0]
    assert m.foreign_tax.rows[0].foreign_tax == pytest.approx(amount + 500.0)
    assert any("never added" in w for w in m.foreign_tax.warnings)


def test_agent_stops_loudly_and_builds_nothing(tmp_path, monkeypatch):
    sys.path.insert(0, str(SCRIPTS.parent))
    import agent
    rep = _report(ordinary_dividend=700.0)
    monkeypatch.setattr(agent.foreign_engine, "load_foreign_income_report", lambda p: rep)
    html = tmp_path / "bs.html"
    html.write_text(fixture_gen.build_syn_ind_html(), encoding="utf-8")
    mf = tmp_path / "SYN-IND.mapping.yaml"
    entries = yaml.safe_load((FIXTURES / "syn_ind.mapping.yaml").read_text(encoding="utf-8"))
    for e in entries:
        if e["guid"] == BANK_INT_GUID:
            e["tag"] = "OS_DIVIDEND"
            e["flags"] = ["FOREIGN_INCOME"]
    mf.write_text(yaml.safe_dump(entries), encoding="utf-8")
    out = tmp_path / "out.xlsx"
    summary = agent.run(
        str(html), str(out), mapping_file=str(mf), entity_key="SYN-IND",
        entities_path=str(ROOT / "bundling" / "canonical" / "itr" / "entities.example.yaml"),
        config_path="no-such-config.yaml", foreign_report_xlsx="in-memory.xlsx",
    )
    assert ft.DOUBLE_COUNT_ERROR_MARKER in summary
    assert "STATUS: BUILT" not in summary
    wb = openpyxl.load_workbook(str(out))
    assert "ForeignTax" not in wb.sheetnames and "Statement of Income" not in wb.sheetnames


# ---------------------------------------------------------------------------
# NEGATIVE: relief is never claimed without the Form 67 line
# ---------------------------------------------------------------------------

def test_form_67_line_appears_wherever_relief_is_claimed(tmp_path):
    rep = _report(ordinary_dividend=300000.0, fsi_rows=[_fsi(income=100000.0, paid=5000.0)])
    built = _model(tmp_path, foreign_report=rep)
    assert built[0].foreign_tax.relief_claimed
    wb = _workbook(tmp_path, built)
    assert "ForeignTax" in wb.sheetnames
    text = _all_text(wb)
    assert "Form 67 must be filed" in text
    # On the Statement of Income ladder the relief line itself says so.
    soi_text = "\n".join(
        c.value for row in wb["Statement of Income"].iter_rows() for c in row
        if isinstance(c.value, str))
    assert "Foreign tax relief (s.90/91) -- Form 67 must be filed" in soi_text


def test_no_relief_means_no_relief_line_and_no_form_67_claim(tmp_path):
    built = _model(tmp_path, foreign_tax_tag=True)       # tax present, country unknown
    assert not built[0].foreign_tax.relief_claimed
    wb = _workbook(tmp_path, built)
    soi_text = "\n".join(
        c.value for row in wb["Statement of Income"].iter_rows() for c in row
        if isinstance(c.value, str))
    assert "Foreign tax relief" not in soi_text
    assert "FORM 67" not in _all_text(wb)               # no claim, so no claim line


def test_statement_of_income_total_reflects_relief(tmp_path):
    rep = _report(ordinary_dividend=300000.0, fsi_rows=[_fsi(income=100000.0, paid=5000.0)])
    built = _model(tmp_path, foreign_report=rep)
    wb = _workbook(tmp_path, built)
    ws = wb["Statement of Income"]
    found = [c for row in ws.iter_rows() for c in row if c.value == "Total tax liability"]
    assert found
    row = found[0].row
    formulas = [c.value for c in ws[row] if isinstance(c.value, str) and c.value.startswith("=")]
    assert formulas and "MAX(0," in formulas[0]
    assert wb["ForeignTax"].title == "ForeignTax"


# ---------------------------------------------------------------------------
# NEGATIVE: a run with no foreign activity is unchanged
# ---------------------------------------------------------------------------

def test_no_foreign_activity_changes_nothing(tmp_path):
    m = _model(tmp_path)[0]
    assert not m.foreign_tax.active and m.foreign_tax.rows == []
    assert m.computation.foreign_tax_relief == 0.0
    assert m.computation.tax_block.foreign_tax_relief == 0.0
    assert m.computation.tax_block.tax_liability_before_foreign_relief == 0.0
    wb = _workbook(tmp_path, _model(tmp_path))
    assert "ForeignTax" not in wb.sheetnames
    text = _all_text(wb)
    assert "Form 67" not in text and "Foreign tax relief" not in text


def test_a_report_with_no_foreign_tax_adds_no_sheet_and_no_relief(tmp_path):
    rep = _report(ordinary_dividend=300000.0, fsi_rows=[_fsi(income=100000.0, paid=0.0)])
    m = _model(tmp_path, foreign_report=rep)[0]
    assert not m.foreign_tax.active and m.foreign_tax.relief_total == 0.0
    base = _model(tmp_path, foreign_report=_report(ordinary_dividend=300000.0))[0]
    assert m.computation.tax_block.tax_liability == pytest.approx(base.computation.tax_block.tax_liability)
