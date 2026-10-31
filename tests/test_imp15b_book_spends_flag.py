"""IMP-15 4b (Book Spends): a look-alike hand entry no longer withholds a row; it is written and
flagged "check before import". Only an exact Num match leaves a row out. Synthetic only."""
from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from agents.skill_cc_spend_booking import agent as ccsb
from agents.skill_cc_spend_booking import journal as J
from agents.skill_cc_spend_booking.bookio import BookTxn
from tests.test_cc_spend_booking import BANK, FOOD, basic, book_pay


def _hand(desc="Swiggy dinner with friends", amount=1500.0, d=date(2025, 4, 5)):
    return BookTxn("hand1", d, desc, "", ((FOOD, amount), (BANK, -amount)))


def _journals(out):
    ws = load_workbook(out.with_name(out.stem + "_report.xlsx"))["Journals"]
    hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]


def _txns(*extra):
    return [book_pay("p1", 1900, date(2025, 5, 12)), *extra]


def test_a_look_alike_row_is_written_flagged_and_reported(tmp_path, monkeypatch):
    reply, rows, _c, out = basic(tmp_path, monkeypatch, txns=_txns(_hand()))
    assert "SWIGGY ORDER" in {r["Description"] for r in rows}            # NEGATIVE: not withheld
    assert "WITHHELD" not in reply
    assert "CHECK BEFORE IMPORT - possibly already in the book" in reply
    assert "NOT complete" in reply                                       # never reads as a clean run


def test_no_look_alike_means_no_flag(tmp_path, monkeypatch):
    reply, rows, _c, out = basic(tmp_path, monkeypatch, txns=_txns())
    assert "CHECK BEFORE IMPORT - possibly already in the book" not in reply     # NEGATIVE
    assert all(j["Status"] != J.ST_POSSIBLY for j in _journals(out))


def test_an_exact_num_match_is_still_left_out_and_not_also_flagged(tmp_path, monkeypatch):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    _r, _rows, _c, out = basic(tmp_path / "a", monkeypatch, txns=_txns())
    num = next(j["Num"] for j in _journals(out) if "SWIGGY" in str(j["Description"]).upper())
    own = BookTxn("own1", date(2025, 4, 5), "Swiggy", num, ((FOOD, 1500.0), (BANK, -1500.0)))
    _r2, rows, _c2, out2 = basic(tmp_path / "b", monkeypatch, txns=_txns(own))
    assert "SWIGGY ORDER" not in {r["Description"] for r in rows}        # definite: left out
    j = next(x for x in _journals(out2) if "SWIGGY" in str(x["Description"]).upper())
    assert j["Status"] == J.ST_BOOKED and j["Status"] != J.ST_POSSIBLY   # NEGATIVE: not flagged as well


def test_one_hand_entry_flags_one_row_only(tmp_path, monkeypatch):
    # pairing is one-to-one: a single hand entry can flag at most one row
    _r, _rows, _c, out = basic(tmp_path, monkeypatch, txns=_txns(_hand()))
    assert len([j for j in _journals(out) if j["Status"] == J.ST_POSSIBLY]) == 1


def test_an_unreadable_book_is_loud_and_never_a_silent_drop(tmp_path, monkeypatch):
    _r, _rows, _c, _o = basic(tmp_path, monkeypatch, txns=_txns())

    def boom(_p):
        raise OSError("cannot read")
    monkeypatch.setattr(ccsb, "read_book", boom)
    bad = tmp_path / "bad"
    bad.mkdir()
    text = ccsb.run(str(tmp_path / "pdfs"), "SYN", str(tmp_path / "book.gnucash"), str(bad / "o.csv"),
                    "Custom date range", "2025-04-01", "2025-05-31",
                    str(tmp_path / "entities.yaml"), str(tmp_path / "config.yaml"))
    text = str(getattr(text, "text", text))
    assert text.startswith("ERROR") and "Nothing was booked" in text     # loud, never a quiet partial file
    assert not (bad / "o.csv").exists()
